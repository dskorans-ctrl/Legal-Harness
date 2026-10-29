"""Render portable brief packages and check restored Word draft paragraphs."""

from dataclasses import dataclass, replace
from datetime import date
import os
from pathlib import Path
import re
import shutil

from docx import Document

from casefile import dates, extract, outputs, quotes, roles, stamps


CITE = re.compile(r'\[(?:cite|page) ([^\]\n]+)\]')
QUOTED = re.compile(r'"([^\"]*)"')
BRACKET = re.compile(r'\[[^\]]*\]')


def conflict_excerpts(casefile, pages, info, transform):
    """Keep every reading and both complete pages, prioritizing defense against ours."""
    by_id = {page.page_id: page for page in pages}
    flags = [flag for flag in casefile.flags if flag.kind == 'conflict']

    def priority(flag):
        pair = {roles.role(by_id[pid].binder, info) for pid in flag.page_ids if pid in by_id}
        return not {'defense', 'ours'} <= pair

    lines = []
    for flag in sorted(flags, key=priority):
        lines.extend("Model's reading: " + transform(about) for about in flag.abouts)
        for pid, cite in zip(flag.page_ids, flag.cites):
            if pid not in by_id:
                raise outputs.OutputError(f'conflict page {pid} is missing; restore the page before packaging')
            lines.append(f'From [cite {cite}]:')
            lines.append(transform(by_id[pid].text))
    return '\n\n'.join(lines) or 'none'


def package_text(packet, casefile, pages, info, *, issue, keep, rules_text, transform, pages_all=False):
    stored = {fact.id: fact for fact in casefile.facts}

    def chronology(fact):
        parsed = dates.parse(fact.date)
        if parsed is None and fact.id in stored:
            parsed = next((dates.parse(value) for value in stored[fact.id].dates if dates.parse(value)), None)
        return parsed or date.max

    lines = [rules_text.rstrip(), f'Matter {packet.matter}', 'Issue: ' + transform(issue),
             'Real names kept in this package (never placeholders): ' + (', '.join(keep) or 'none'),
             '## Facts']
    for fact in sorted(packet.facts, key=chronology):
        lines.append(f'{transform(fact.date) if fact.date else "Undated"}, '
                     f'{transform(fact.provider) or "Provider not stated"}, '
                     f'{fact.category.replace("_", " ")}: "{transform(fact.quote)}" [cite {fact.cite}]')
    lines += ['## Conflicts', conflict_excerpts(casefile, pages, info, transform), '## Gaps']
    gaps = [transform(about) for flag in casefile.flags if flag.kind == 'gap' for about in flag.abouts]
    lines.extend(gaps or ['none'])
    if pages_all:
        lines.append('## Pages')
        lines.extend(extract.page_block(replace(page, text=transform(page.text))) for page in pages)
    return '\n\n'.join(lines) + '\n'


class Draft(list):
    """Body paragraphs with a count of ignored tables and populated header/footer parts."""

    def __init__(self, paragraphs, *, skipped=0):
        super().__init__(paragraphs)
        self.skipped = skipped


def read_draft(path):
    document = Document(path)
    parts = {part.partname: part for section in document.sections
             for part in (section.header.part, section.first_page_header.part,
                          section.even_page_header.part, section.footer.part,
                          section.first_page_footer.part, section.even_page_footer.part)}
    skipped = len(document.tables) + sum(bool(part.element.xpath('.//w:t')) for part in parts.values())
    return Draft([paragraph.text for paragraph in document.paragraphs], skipped=skipped)


@dataclass(frozen=True)
class Unit:
    index: int
    text: str
    cites: tuple[str, ...]
    quotes: tuple[str, ...]


def parse_draft(paragraphs):
    result = []
    for index, text in enumerate(paragraphs):
        spans = tuple(QUOTED.finditer(text))
        outside = QUOTED.sub(lambda match: ' ' * len(match[0]), text)
        result.append(Unit(index, text, tuple(match[1] for match in CITE.finditer(outside)),
                           tuple(match[1] for match in spans)))
    return result


@dataclass(frozen=True)
class Checked:
    unit: Unit
    paragraph: outputs.Paragraph
    failures: tuple[str, ...]


@dataclass(frozen=True)
class Report:
    paragraphs: tuple[Checked, ...]
    uncited: tuple[int, ...]
    rule_hits: int
    skipped: int

    @property
    def total(self):
        return sum(bool(item.unit.cites) for item in self.paragraphs)

    @property
    def failed(self):
        return sum(bool(item.failures) for item in self.paragraphs)

    @property
    def checks(self):
        return sum(bool(item.unit.cites) and not item.failures for item in self.paragraphs)

    @property
    def ok(self):
        return self.failed == 0


def _resolve(cite, pages):
    matches = [page for page in pages if cite == page.page_id
               or stamps.normalize(cite) in {stamps.normalize(value) for value in page.stamps}]
    return matches[0] if len(matches) == 1 else None


def _checked_prose(runs, rules):
    masked = ''.join(' ' * len(run.text) if run.exempt else run.text for run in runs)
    findings = rules.check(masked, unit='text')
    result, offset = [], 0
    for run in runs:
        value = run.text
        if not run.exempt:
            spans = []
            for finding in sorted(findings, key=lambda f: (f.start, -f.end)):
                start, end = max(0, finding.start - offset), min(len(value), finding.end - offset)
                if end <= start:
                    continue
                if spans and start <= spans[-1][1]:
                    left, right, rule_id = spans[-1]
                    spans[-1] = (left, max(right, end), rule_id)
                else:
                    spans.append((start, end, finding.rule_id))
            for start, end, rule_id in reversed(spans):
                value = value[:start] + f'[withheld: {rule_id}]' + value[end:]
        result.append(outputs.Run(value, exempt=run.exempt))
        offset += len(run.text)
    return outputs.Paragraph(result), len(findings)


def fact_check(draft, pages, rules, *, restore_prose, restore_quote):
    """Check citations per paragraph; quote restoration never confers trust on prose."""
    skipped = getattr(draft, 'skipped', 0)
    units = parse_draft(draft) if not draft or isinstance(draft[0], str) else draft
    checked, uncited, rule_hits = [], [], 0
    for unit in units:
        failures = []
        resolved = []
        for cite in unit.cites:
            page = _resolve(cite, pages)
            if page is None:
                failures.append(f'paragraph {unit.index + 1}: unknown or ambiguous cite {cite}; correct the cite')
            else:
                resolved.append(page)
        if not unit.cites:
            uncited.append(unit.index + 1)
        restored_quotes = []
        for quote in unit.quotes:
            candidates = [restore_quote(quote, page) for page in resolved]
            good = bool(resolved) and all(value is not None and quotes.found_in(value, page.text)
                                          for value, page in zip(candidates, resolved))
            if not good and unit.cites:
                labels = ', '.join(page.page_id for page in resolved) or ', '.join(unit.cites)
                failures.append(f'paragraph {unit.index + 1}: quote not on {labels}; correct the quote or cite')
            restored_quotes.append((candidates[0] if good else restore_prose(quote), good))
        if unit.cites and not unit.quotes:
            prose_numbers = set(re.findall(r'\d+', BRACKET.sub('', unit.text)))
            for page in resolved:
                if not prose_numbers <= set(re.findall(r'\d+', page.text)):
                    failures.append(f'paragraph {unit.index + 1}: number not on {page.page_id}; correct the number or cite')
        runs, cursor, quote_index = [], 0, 0
        tokens = re.compile(r'"[^\"]*"|\[(?:cite|page) [^\]\n]+\]')
        restored_parts = []
        for match in tokens.finditer(unit.text):
            prose = restore_prose(unit.text[cursor:match.start()])
            restored_parts.append(prose)
            runs.append(outputs.Run(prose))
            if match[0].startswith('"'):
                value, verified = restored_quotes[quote_index]
                quote_index += 1
                restored_parts.append('"' + value + '"')
                runs.append(outputs.Run('"' + value + '"', exempt='quote' if verified else ''))
            else:
                restored_parts.append(match[0])
                if _resolve(CITE.fullmatch(match[0])[1], pages):
                    runs.append(outputs.Run(match[0], exempt='cite'))
                else:
                    runs.append(outputs.Run(match[0]))
            cursor = match.end()
        prose = restore_prose(unit.text[cursor:])
        restored_parts.append(prose)
        runs.append(outputs.Run(prose))
        paragraph, hits = _checked_prose(runs, rules)
        rule_hits += hits
        if hits:
            failures.append(f'paragraph {unit.index + 1}: firm rule hit; revise the withheld prose')
        restored = ''.join(restored_parts)
        if restored.count('\n') != unit.text.count('\n') or restored.count('\r') != unit.text.count('\r'):
            failures.append(f'paragraph {unit.index + 1}: paragraph count changed; correct the restoration')
        checked.append(Checked(unit, paragraph, tuple(dict.fromkeys(failures))))
    return Report(tuple(checked), tuple(uncited), rule_hits, skipped)


def render(report, info, *, playbook_version, model):
    paragraphs = []
    for item in report.paragraphs:
        paragraphs.append(item.paragraph)
        for reason in item.failures:
            paragraphs.append(outputs.Paragraph([outputs.Run('[FACT CHECK FAILED: ' + reason + ']', exempt='reason')]))
    summary = [outputs.p(f'Checks: {report.checks} of {report.total}; failures: {report.failed}; '
                         f'uncited: {len(report.uncited)}; rule hits: {report.rule_hits}; skipped: {report.skipped}'),
               outputs.p('Uncited paragraphs: ' + (', '.join(map(str, report.uncited)) or 'none')),
               outputs.p('Open items: ' + ('correct each marked paragraph' if report.failed else 'none')),
               outputs.p(f'Playbook {playbook_version}; model {model}')]
    return outputs.Doc('Restored brief draft', outputs.header(info, '', 'brief'),
                       [outputs.Section('Draft', paragraphs), outputs.Section('Fact check', summary)])


def publish_package(folder, packet_id, files):
    """Publish privately, refusing replacement and rolling back only owned paths."""
    folder = Path(folder)
    if not re.fullmatch(r'[a-f0-9]{12}', packet_id):
        raise outputs.OutputError('invalid packet id')
    if folder.is_symlink() or not folder.parent.is_dir() or (folder.exists() and not folder.is_dir()):
        raise outputs.OutputError('outbox must be a plain folder')
    if any(name not in {'package.md', 'package.json', 'check.txt'} for name in files):
        raise outputs.OutputError('invalid package filename')
    target, temporary = folder / packet_id, folder / ('.' + packet_id + '.tmp')
    if target.exists() or target.is_symlink():
        raise outputs.OutputError(f'package already exists: {packet_id}; use the existing package')
    created_parent = created_temp = reserved = False
    try:
        if not folder.exists():
            folder.mkdir(mode=0o700)
            created_parent = True
        temporary.mkdir(mode=0o700)
        created_temp = True
        os.chmod(temporary, 0o700)
        for name, data in files.items():
            with (temporary / name).open('xb') as stream:
                os.fchmod(stream.fileno(), 0o600)
                stream.write(data.encode('utf-8') if isinstance(data, str) else data)
        target.mkdir(mode=0o700)
        reserved = True
        os.replace(temporary, target)
        created_temp = False
    except BaseException:
        if created_temp:
            shutil.rmtree(temporary)
        if reserved:
            target.rmdir()
        if created_parent and not any(folder.iterdir()):
            folder.rmdir()
        raise
    return target

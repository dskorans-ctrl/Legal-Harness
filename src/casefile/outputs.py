"""Shared checked document primitives and fixed output job steps."""

from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import datetime, timezone
import fcntl
import os
import re
from pathlib import Path
import sqlite3
import sys

from docx import Document

from casefile import facts, gates, matterinfo, pages, playbook, quotes, roles, rules, stamps
from casefile.model import ModelClient, ModelError
from casefile.runlog import open_db


class OutputError(ValueError):
    """A task cannot produce a checked document."""


@dataclass(frozen=True)
class Run:
    text: str
    exempt: str = ''


@dataclass
class Paragraph:
    runs: list[Run]
    level: int | None = None
    warning_pages: tuple[str, ...] = ()
    unit: str = 'text'
    page_break: bool = False
    bills_question: bool = False


@dataclass
class Section:
    title: str
    paragraphs: list[Paragraph] = field(default_factory=list)


@dataclass
class Doc:
    title: str
    header: list[Paragraph]
    sections: list[Section]


GROUPS = {
    'Records referenced': ('records_referenced',),
    'History': ('mechanism_of_injury', 'prior_history', 'patient_statement', 'chief_complaint', 'pain_score',
                'family_history', 'comorbidity_social'),
    'Examination findings': ('exam_finding', 'diagnostic_test'),
    'Imaging': ('imaging_result',),
    'Diagnoses': ('diagnosis',),
    'Opinions': ('causation_statement', 'mmi_statement', 'impairment_rating', 'work_status', 'future_treatment'),
}
ELEMENTS = ('treating_relationship', 'mechanism', 'history_prior_conditions', 'exam_findings',
            'imaging_reviewed', 'diagnosis', 'causation', 'treatment_reasonable_necessary_related',
            'bills_reasonable_related', 'work_status_restrictions', 'mmi', 'permanency_rating',
            'future_treatment', 'records_reviewed')


def p(text, **kwargs):
    return Paragraph([Run(text)], **kwargs)


def cite(value):
    return Run(f'[page {value}]' if '/' in value else f'[cite {value}]', exempt='cite')


def evidence(item, prefix=''):
    return Paragraph([Run(prefix), Run('"' + item['quote'] + '"', exempt='quote'), Run(' '), cite(item['cite'])])


def note(text):
    return p("Model's note: " + text)


def header(info, matter, kind, *, when='', deponent=''):
    result = [p('Attorney Work Product | Privileged & Confidential')]
    for key, label in (('matter_type', 'Matter type'), ('client_name', 'Client'),
                       ('date_of_accident', 'Date of accident'), ('treating_physician', 'Treating physician'),
                       ('iab_hearing_number', 'IAB hearing number'), ('hearing_date', 'Hearing date')):
        result.append(p(f'{label}: {getattr(info, key) or "NEEDED"}'))
    if kind == 'depo':
        result.append(p(f'Deposition of {deponent or "NEEDED"}, {when or "NEEDED"}'))
    result.append(p(f'Matter {matter}'))
    return result


def build_line(casefile, playbook_version, model):
    return p(f'Built {datetime.now(timezone.utc).date().isoformat()} from case file entry '
             f'{casefile.last_seq}; playbook {playbook_version}; model {model}')


def checked_runs(paragraph, rules):
    """Withhold rule spans from untrusted failure details only in documents."""
    result = []
    for run in paragraph.runs:
        if run.exempt != 'reason':
            result.append(run)
            continue
        value = run.text
        findings = sorted(rules.check(value, unit='text'), key=lambda f: (f.start, -f.end))
        spans = []
        for finding in findings:
            if spans and finding.start <= spans[-1][1]:
                start, end, rule_id = spans[-1]
                spans[-1] = (start, max(end, finding.end), rule_id)
            else:
                spans.append((finding.start, finding.end, finding.rule_id))
        for start, end, rule_id in reversed(spans):
            value = value[:start] + f'[withheld: {rule_id}]' + value[end:]
        result.append(Run(value))
    return result


def page_warning(pages):
    ids = [page.page_id for page in pages if not page.stamps and page.text.strip()]
    return [Paragraph([Run('These pages carry no Bates stamp and are cited by page id: '),
                       Run(', '.join(ids), exempt='cite'), Run('.')])] if ids else []


def check(doc, rules, stamped=True, providers=()):
    findings = []
    texts = [(doc.title, 'text', False)]
    for paragraph in doc.header:
        texts.append((''.join(run.text if not run.exempt else '\n' for run in checked_runs(paragraph, rules)), paragraph.unit, paragraph.bills_question))
    for section in doc.sections:
        texts.append((section.title, 'text', False))
        for paragraph in section.paragraphs:
            texts.append((''.join(run.text if not run.exempt else '\n' for run in checked_runs(paragraph, rules)), paragraph.unit, paragraph.bills_question))
    for text, unit, bills_question in texts:
        findings.extend(rules.check(text, unit=unit, stamped=stamped, providers=providers, bills_question=bills_question))
    return findings


def render(doc, path, *, rules, stamped=True, providers=()):
    errors = [f for f in check(doc, rules, stamped, providers) if f.severity == 'error']
    if errors:
        raise OutputError('\n'.join(f'{f.rule_id}: {f.excerpt}' for f in errors))
    document = Document()
    document.add_heading(doc.title, level=0)
    for paragraph in doc.header:
        _write_paragraph(document, paragraph, rules)
    for section in doc.sections:
        document.add_heading(section.title, level=1)
        for paragraph in section.paragraphs or [p('none')]:
            _write_paragraph(document, paragraph, rules)
    document.save(path)


def _write_paragraph(document, paragraph, rules):
    target = document.add_heading(level=paragraph.level) if paragraph.level else document.add_paragraph()
    for run in checked_runs(paragraph, rules):
        target.add_run(run.text)
    if paragraph.page_break:
        target.add_run().add_break(__import__('docx.enum.text', fromlist=['WD_BREAK']).WD_BREAK.PAGE)


def providers(casefile):
    return tuple(dict.fromkeys(value for fact in casefile.facts for value in fact.providers if value))


def publish(folder, name, doc, rules, casefile, counts, *, failed_steps):
    folder = Path(folder)
    directory = folder / 'outputs'
    if not folder.is_dir() or directory.is_symlink() or (directory.exists() and not directory.is_dir()):
        raise OutputError(f'cannot write {name}: matter and outputs must be plain folders')
    suffix = '.partial.docx' if failed_steps else '.docx'
    target = directory / (name + suffix)
    temporary = target.with_name(target.name + '.tmp')
    if temporary.exists() or temporary.is_symlink() or target.is_symlink():
        raise OutputError(f'cannot write {name}: output or temporary path is a link or already exists: {temporary}')
    created = not directory.exists()
    owned = False
    try:
        directory.mkdir(exist_ok=True)
        with temporary.open('xb'):
            owned = True
        render(doc, temporary, rules=rules, providers=providers(casefile))
        os.replace(temporary, target)
    except BaseException:
        if owned:
            temporary.unlink(missing_ok=True)
        if created and directory.exists() and not any(directory.iterdir()):
            directory.rmdir()
        raise
    if not failed_steps:
        casefile = facts.append(folder / 'casefile.jsonl', [facts.Event('output', dict(name=name,
            path=str(target), built_from_seq=casefile.last_seq, facts=counts['facts'],
            conflicts=counts['conflicts'], gaps=counts['gaps'], open=counts['open']))])
        (directory / (name + '.partial.docx')).unlink(missing_ok=True)
    return casefile


def provider_pages(casefile, text, *, pages=(), info=None):
    folded = quotes.fold(text)
    names = {quotes.fold(name): name for name in providers(casefile)}
    exact = [name for name in names if name == folded]
    candidates = exact or [name for name in names if folded and re.search(
        r'(?<!\w)' + re.escape(folded) + r'(?!\w)', name)]
    if len(candidates) > 1:
        raise OutputError('ambiguous provider: ' + ', '.join(names[name] for name in candidates))
    ids = tuple(dict.fromkeys(pid for fact in casefile.facts for pid, provider in zip(fact.page_ids, fact.providers)
                             if quotes.fold(provider) in candidates))
    defense = {page.binder for page in pages if page.page_id in ids and roles.role(page.binder, info or matterinfo.MatterInfo()) == 'defense'}
    return tuple(dict.fromkeys((*ids, *(page.page_id for page in pages if page.binder in defense))))


def select_pages(casefile, pages, *, provider='', binder=(), selected=None, doctor=True, info=None):
    if selected:
        result = []
        for value in selected.split(','):
            matches = [page for page in pages if value.strip() == page.page_id
                       or stamps.normalize(value.strip()) in page.stamps]
            if len(matches) != 1:
                raise OutputError(f'unknown or ambiguous page: {value.strip()}')
            result.extend(matches)
        return tuple(dict.fromkeys(result))
    if binder:
        unknown = set(binder) - {page.binder for page in pages}
        if unknown:
            raise OutputError('unknown binder: ' + ', '.join(sorted(unknown)))
        return tuple(page for page in pages if page.binder in binder)
    if provider:
        ids = provider_pages(casefile, provider, pages=pages, info=info)
        if not ids:
            raise OutputError('no fact names that provider; pass --binder or --pages; providers on file: '
                              + ', '.join(providers(casefile)))
        return tuple(page for page in pages if page.page_id in ids)
    if doctor:
        raise OutputError('name the deponent with --provider, or the records with --binder or --pages')
    return tuple(pages)


@dataclass
class Started:
    config: object
    folder: Path
    book: object
    loaded: object
    casefile: object
    info: object
    log: object
    client: object
    counts: dict


@contextmanager
def start(config, number, job, *, opener=None, stdin=None, out=None):
    """Refuse incomplete matter inputs before creating a run log or a model client."""
    from casefile import status

    book = playbook.load(config.paths.playbook)
    loaded = pages.load(config.paths.matters, number, config.ocr.stamp_pattern)
    if not loaded.pages:
        raise OutputError(f'no pages to read in {number}')
    folder = config.paths.matters / number
    with (folder / 'casefile.jsonl.lock').open('a') as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)
        casefile = facts.load(folder / 'casefile.jsonl')
        if not casefile.facts:
            raise OutputError('run casefile memo first')
        rules.check_stamps(loaded.pages, book.rules)
        info = matterinfo.load(folder)
        stream = sys.stdin if stdin is None else stdin
        if not info.has_header and stream.isatty():
            info = matterinfo.ask(info, stream, sys.stdout if out is None else out)
        with open_db(config.paths.state / 'runlog.sqlite') as log:
            counts = status.snapshot(loaded, casefile, log.coverage(number, job='memo'), book, folder=folder)
            yield Started(config, folder, book, loaded, casefile, info, log,
                          ModelClient(config.model, opener), counts)


ERRORS = (OutputError, facts.CaseFileError, pages.PagesError, playbook.PlaybookError,
          matterinfo.MatterInfoError, rules.RulesError, ModelError, OSError, sqlite3.Error, ValueError)


def ask(packet, card, *, client, rules, log, job, playbook_version):
    return gates.ask(card, packet.pages, client=client, rules=rules, log=log, job=job,
                     matter=packet.matter, playbook_version=playbook_version, body=packet.body(),
                     providers=tuple(dict.fromkeys(f.provider for f in packet.facts if f.provider)))


def failure(section, reason=None):
    reason = reason or 'gate check; the gated answer is unavailable'
    gate, _, detail = reason.partition(':')
    return Paragraph([Run(f"{section}: the model's answer failed the {gate} check after one retry: "),
                      Run(detail.strip(), exempt='reason')])


def failure_reason(outcome):
    return next((v.message for v in reversed(outcome.verdicts) if v.verdict == 'fail'), 'gate: answer unavailable')


def needed(info, pages):
    result = [f'{key.replace("_", " ").capitalize()}: NEEDED' for key in
              matterinfo.HEADER[:-1] if not getattr(info, key)]
    for binder in dict.fromkeys(page.binder for page in pages):
        if roles.role(binder, info) not in ('ours', 'defense', 'other'):
            result.append(f'ROLE NEEDED: {binder}')
    return result


def open_items_section(items, extra=()):
    paragraphs = []
    for item in items:
        paragraphs.append(Paragraph([Run(f'{item.label}: {item.status}: ', exempt='cite'),
                                     Run(item.reason, exempt='reason')]))
    paragraphs.extend(p(value) for value in extra)
    return Section('Open items', paragraphs or [p('none')])


def fact_paragraphs(fact, pages):
    present = {page.page_id for page in pages}
    primary = next((i for i, pid in enumerate(fact.page_ids) if pid in present), 0)
    result = [evidence(dict(quote=fact.quote, cite=fact.cites[primary]), (fact.dates[primary] or 'Undated') + ': ')]
    if len(fact.cites) > 1:
        result.append(Paragraph([Run('Also cited: '), *[run for i,c in enumerate(fact.cites)
                         if i != primary for run in (cite(c), Run(' '))]]))
    for text, pid in zip(fact.readings, fact.reading_pages):
        prefix = [Run("Model's reading")]
        if pid != fact.page_ids[primary]:
            prefix += [Run(' of '), cite(fact.cites[fact.page_ids.index(pid)])]
        result.append(Paragraph([*prefix, Run(': ' + text)]))
    return result


def grouped(records, pages):
    result = []
    used = set()
    for title, categories in (*GROUPS.items(), ('Other statements', ())):
        selected = [fact for fact in records if (fact.category in categories if categories else fact.id not in used)]
        if selected:
            result.append(p(title, level=2))
            for fact in selected:
                result.extend(fact_paragraphs(fact, pages))
                used.add(fact.id)
    return result or [p('none')]


def conflicts(casefile, pages, info, *, predicate=lambda flag: True):
    by_id = {page.page_id: page for page in pages}
    selected = [flag for flag in casefile.flags if flag.kind == 'conflict' and predicate(flag)]
    def priority(flag):
        pair = {roles.role(by_id[pid].binder, info) for pid in flag.page_ids if pid in by_id}
        return not {'defense', 'ours'} <= pair
    selected.sort(key=priority)
    result = []
    for flag in selected:
        result.append(Paragraph([Run('Conflict between '), *[run for i,c in enumerate(flag.cites)
            for run in ((Run(' and ') if i else Run('')), cite(c))], Run('.')]))
        result.extend(p("Model's reading: " + text) for text in flag.abouts)
        for pid, value in zip(flag.page_ids, flag.cites):
            result.append(Paragraph([Run('From '), cite(value), Run(':')]))
            result.extend(evidence(dict(quote=fact.quote, cite=value)) for fact in casefile.facts if pid in fact.page_ids)
    return result or [p('none')]


def elements(answer, pages, *, title='Elements', reason=None):
    section, gaps, subjects = Section(title), Section('Gaps'), Section('Items to raise with the doctor (subjects only)')
    if answer is None:
        section.paragraphs.append(failure(title, reason))
        return section, gaps, subjects
    found = {}
    for item in answer['elements']:
        if item['element'] not in found or item['status'] == 'STATED':
            found[item['element']] = item
    pages = tuple(page for page in pages if page.page_id in answer.get('searched_pages', [p.page_id for p in pages]))
    labels = ', '.join(page.stamps[0] if page.stamps else page.page_id for page in pages)
    for name in ELEMENTS:
        item = found.get(name)
        if item is None:
            paragraph = Paragraph([Run(name + ': NOT FOUND on pages '), *[run for i,page in enumerate(pages)
                for run in (Run(', ' if i else ''), cite(page.stamps[0] if page.stamps else page.page_id))]])
        else:
            paragraph = evidence(item, name + ': ' + item['status'] + ': ')
        section.paragraphs.append(paragraph)
        if item and item['absent_note']:
            section.paragraphs.append(note(item['absent_note']))
        if item is None or item['status'] == 'PARTIAL':
            gaps.paragraphs.append(paragraph)
            subjects.paragraphs.append(p(name))
    if not labels:
        gaps.paragraphs.append(p('No pages sent'))
    return section, gaps, subjects


def labelled_facts(answer, casefile, uses):
    if answer is None:
        return []
    return [fact for fact in casefile.facts if any(item['use'] in uses and quotes.same(item['quote'], fact.quote)
            and item['cite'] in (*fact.cites, *fact.page_ids) for item in answer['facts'])]


def question_paragraphs(answer, title, *, reason=None, omitted=0, bills_question=False):
    if answer is None:
        return [failure(title, reason)]
    result = []
    if omitted:
        result.append(p(f'{omitted} facts were not sent to the model for this section (prompt budget); they are listed in the memo'))
    for n,item in enumerate(answer['questions'], 1):
        result += [p(f'{n}. {item["text"]}', unit='question', bills_question=bills_question), evidence(item, 'Rests on: ')]
        if item['note']:
            result.append(note(item['note']))
    return result


def finish(started, name, doc, *, failed_steps, out):
    casefile = publish(started.folder, name, doc, started.book.rules, started.casefile,
                       started.counts, failed_steps=failed_steps)
    opened = next(section for section in doc.sections if section.title == 'Open items')
    has_open = [paragraph for paragraph in opened.paragraphs if ''.join(r.text for r in paragraph.runs) != 'none']
    print(f'{name} {started.loaded.number} output={name}{".partial" if failed_steps else ""}.docx '
          f'open={len(has_open)} casefile_entry={casefile.last_seq}', file=out)
    return 1 if failed_steps or has_open else 0

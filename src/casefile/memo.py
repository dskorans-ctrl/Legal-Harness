"""Build the interim medical memo from case file evidence (D41)."""

from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

from casefile.rules import Finding, Rules
from casefile.outputs import Paragraph, Run, Section


class MemoError(ValueError):
    """The memo cannot be checked or written."""


@dataclass(frozen=True)
class OpenItem:
    page_id: str
    label: str
    status: str
    reason: str


@dataclass
class Memo:
    title: str
    built: str
    sections: list[Section]
    pages: tuple


def label(page):
    return f'{" ".join(page.stamps)} ({page.page_id})' if page.stamps else page.page_id


def _cite(cite, page_id):
    return f'[page {page_id}]' if cite == page_id else f'[cite {cite}]'


def _paragraph(text):
    return Paragraph([Run(text)])


def _added(count, noun):
    return f'{count} {noun}{"" if count == 1 else "s"} added'


def _fact_paragraphs(fact, primary, heading_page):
    paragraphs = [Paragraph([
        Run(f'"{fact.quote}"', exempt="quote"),
        Run(" " + _cite(fact.cites[primary], fact.page_ids[primary]))])]
    others = [(cite, page_id) for index, (cite, page_id) in
              enumerate(zip(fact.cites, fact.page_ids)) if index != primary]
    if others:
        paragraphs.append(_paragraph("Also cited: " + " and ".join(
            _cite(cite, page_id) for cite, page_id in others)))
    for reading, page_id in zip(fact.readings, fact.reading_pages):
        prefix = "Model's reading"
        if page_id != heading_page:
            cite = fact.cites[fact.page_ids.index(page_id)]
            prefix += " of " + _cite(cite, page_id)
        paragraphs.append(_paragraph(f"{prefix}: {reading}"))
    return paragraphs


def build(casefile, pages, open_items, *, matter, playbook_version, model) -> Memo:
    pages = tuple(pages)
    open_items = tuple(open_items)
    previous = next((event for event in reversed(casefile.events)
                     if event.kind == "output" and event.payload["name"] == "memo"), None)
    if previous is None:
        changed = "First memo for this matter"
    else:
        counts = casefile.changed_since(previous.seq)
        changed = ", ".join(_added(count, noun) for count, noun in (
            (counts["fact"], "fact"), (counts["cite"], "cite"),
            (counts["fact"] + counts["reading"], "reading"), (counts["flag"], "flag")))
        if counts["about"]:
            changed += ", " + _added(counts["about"], "flag reading")
    changes = Section("What changed since the last memo", [_paragraph(changed)])
    stale = Section("Outputs out of date", [_paragraph(
        ", ".join(name for name in casefile.stale_outputs() if name != "memo") or "none")])
    opened = Section("Open items", [Paragraph([
        Run(f'{item.label}: {item.status}: '), Run(item.reason, exempt="reason")])
        for item in open_items] or [_paragraph("none")])
    runs = {run["run_id"] for run in casefile.runs if run["job"] == "memo" and run["read"]}
    conflicts = Section("Conflicts", [_paragraph(
        f'Pages read in {len(runs)} {"run" if len(runs) == 1 else "runs"}; '
        'conflicts and gaps are checked only between pages sent to the model together.')])
    gaps = Section("Gaps and missing records")
    for flag in casefile.flags:
        citations = [_cite(cite, page_id) for cite, page_id in zip(flag.cites, flag.page_ids)]
        section = conflicts if flag.kind == "conflict" else gaps
        prefix = "Conflict between " if flag.kind == "conflict" else "Gap at "
        section.paragraphs.append(_paragraph(prefix + " and ".join(citations) + "."))
        for about in flag.abouts:
            section.paragraphs.append(_paragraph(f"Model's reading: {about}"))
    if not any(flag.kind == "conflict" for flag in casefile.flags):
        conflicts.paragraphs.append(_paragraph("none"))
    if not gaps.paragraphs:
        gaps.paragraphs.append(_paragraph("none"))
    verified = Section("Verified facts by record")
    unstamped = tuple(page.page_id for page in pages if not page.stamps and page.text.strip())
    if unstamped:
        verified.paragraphs.append(Paragraph([], warning_pages=unstamped))
    records = casefile.facts
    present = {page.page_id for page in pages}
    placement = {fact.id: next((index for index, page_id in enumerate(fact.page_ids)
                               if page_id in present), None) for fact in records}
    open_status = {item.page_id: item.status for item in open_items}
    for page in pages:
        verified.paragraphs.append(Paragraph([Run(f'{page.binder.capitalize()} page {label(page)}')], level=2))
        quoted = False
        for fact in records:
            primary = placement[fact.id]
            if primary is None or fact.page_ids[primary] != page.page_id:
                continue
            verified.paragraphs.extend(_fact_paragraphs(fact, primary, page.page_id))
            quoted = True
        if not quoted:
            state = open_status.get(page.page_id)
            verified.paragraphs.append(_paragraph(
                f'Open item: {state} (see Open items).' if state is not None
                else 'No fact was quoted from this page.'))
    sections = [changes, stale, opened, conflicts, gaps, verified]
    absent = [fact for fact in records if placement[fact.id] is None]
    if absent:
        missing = Section("Facts on pages not in this matter")
        for fact in absent:
            missing.paragraphs.extend(_fact_paragraphs(fact, 0, None))
        sections.append(missing)
    built = (f'Built {datetime.now(timezone.utc).date().isoformat()} from case file entry '
             f'{casefile.last_seq}; playbook {playbook_version}; model {model}')
    return Memo(f'Medical memo: matter {matter}', built,
                sections, pages)


def _runs(paragraph, rules):
    if paragraph.warning_pages:
        return [Run('These pages carry no Bates stamp and are cited by page id: '
                    f'{", ".join(paragraph.warning_pages)}.')]
    return paragraph.runs


def check_binders(pages, rules):
    """Called by the memo job before opening the run log or calling the model."""
    for binder in dict.fromkeys(page.binder for page in pages):
        for finding in rules.check(binder, unit="text", stamped=all(page.stamps for page in pages)):
            if finding.severity == "error":
                raise MemoError(f'binder {binder} fails the firm rule {finding.rule_id} '
                                f'({finding.excerpt}); remove the binder by hand and intake it '
                                'under another name (question 34)')


def check(memo, rules: Rules) -> list[Finding]:
    findings = []
    texts = [(memo.title, "title"), (memo.built, "build line")]
    for section in memo.sections:
        texts.append((section.title, section.title))
        for paragraph in section.paragraphs:
            texts.append(("".join(run.text if not run.exempt else "\n"
                                  for run in _runs(paragraph, rules)), section.title))
    for text, section in texts:
        for finding in rules.check(text, unit="text", stamped=all(page.stamps for page in memo.pages)):
            excerpt = " ".join(finding.excerpt.split())
            location = " ".join(section.split())
            findings.append(Finding(finding.rule_id, finding.severity, finding.start,
                                    finding.end, f'{excerpt} (in {location})'))
    return findings


def render(memo, path, *, rules: Rules):
    check_binders(memo.pages, rules)
    findings = check(memo, rules)
    errors = [finding for finding in findings if finding.severity == "error"]
    if errors:
        raise MemoError("\n".join(f'{finding.rule_id}: {finding.excerpt}' for finding in errors))
    path = Path(path)
    # Only the output directory may be created; an absent matter is a refusal.
    if not path.parent.is_dir() and (path.parent.name != "outputs" or not path.parent.parent.is_dir()):
        raise MemoError(f'cannot write memo {path}: matter folder must already exist')
    document = Document()
    document.add_heading(memo.title, level=0)
    document.add_paragraph(memo.built)
    for section in memo.sections:
        document.add_heading(section.title, level=1)
        for paragraph in section.paragraphs:
            target = document.add_heading(level=paragraph.level) if paragraph.level is not None else document.add_paragraph()
            for run in _runs(paragraph, rules):
                target.add_run(run.text)
    try:
        path.parent.mkdir(exist_ok=True)
        document.save(path)
    except OSError as exc:
        raise MemoError(f'cannot write memo {path}: {exc}') from exc

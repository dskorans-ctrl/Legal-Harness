"""Run unread pages through the gates and publish a checked medical memo."""

import fcntl
import os
import sqlite3

from docx import Document

from casefile import extract, facts, matterinfo, memo, outputs, pages, pipeline, playbook, rules, status
from casefile.model import ModelClient, ModelError
from casefile.matter import Page
from casefile.runlog import open_db


def _run_event(run_id, card, part, outcome):
    return facts.Event("run", dict(run_id=run_id, job="memo", card=card.name,
        card_version=card.version, pages=[page.page_id for page in part], ok=outcome is not None and outcome.ok,
        read=list(outcome.read) if outcome is not None else [],
        uncited=list(outcome.uncited) if outcome is not None else [],
        failed=[page_id for page_id, _ in outcome.failed] if outcome is not None
        else [page.page_id for page in part]))


def _run_locked(config, number, folder, path, book, loaded, open_pages, *, out, err, opener):
    try:
        with path.open("ab"):
            pass
    except OSError as exc:
        raise facts.CaseFileError(f"cannot append case file {path}: {exc}") from exc
    casefile = facts.load(path)
    if casefile.torn:
        print(status.torn_message(casefile), file=err)
    before = casefile.last_seq
    card = extract.card(book)
    client = ModelClient(config.model, opener=opener)
    with open_db(config.paths.state / "runlog.sqlite") as log:
        start = log.connection.execute("SELECT COUNT(*) FROM runs").fetchone()[0]
        coverage = status.committed_coverage(casefile, log.coverage(number, job="memo"))
        unread = tuple(page for page in loaded.pages
                       if coverage.get((page.page_id, card.name), (None,))[0] != "read")
        chunks, refused = pipeline.chunk(card, unread, client.config.num_ctx, number)
        # One pipeline call per chunk preserves earlier merges if a later server call fails.
        for part in (*((page,) for page, _ in refused), *chunks):
            result = pipeline.run(part, card=card, client=client, rules=book.rules,
                log=log, matter=number, playbook_version=book.version, job="memo", only_unread=False)
            for chunk, outcome in zip(result.chunks, result.outcomes):
                events = facts.merge(casefile, outcome.answer, chunk, outcome.run_id) if outcome.ok else []
                events.append(_run_event(outcome.run_id, card, chunk, outcome))
                casefile = facts.append(path, events)
            for page, _ in result.refused:
                run_id = log.coverage(number, job="memo")[(page.page_id, card.name)][2]
                casefile = facts.append(path, [_run_event(run_id, card, (page,), None)])
        counts = status.snapshot(loaded, casefile, log.coverage(number, job="memo"))
        runs = log.connection.execute("SELECT COUNT(*) FROM runs").fetchone()[0] - start
    all_pages = (*loaded.pages, *open_pages)
    binders = list(dict.fromkeys(page.binder for page in all_pages))
    if loaded.open_pages:
        binders.sort()
    all_pages = sorted(all_pages, key=lambda page: (binders.index(page.binder), page.number))
    info = matterinfo.load(folder)
    document = memo.build(casefile, all_pages, counts["open_items"], matter=number,
                          playbook_version=book.version, model=config.model.name, info=info)
    header = outputs.header(info, number, "memo")
    findings = outputs.check(outputs.Doc("", header, []), book.rules)
    errors = [f for f in findings if f.severity == "error"]
    if errors:
        raise memo.MemoError("\n".join(f'{f.rule_id}: {f.excerpt}' for f in errors))
    output = folder / "outputs" / "memo.docx"
    temporary = output.with_name("memo.docx.tmp")
    if temporary.is_symlink() or temporary.exists():
        raise memo.MemoError(f"cannot write memo: temporary output already exists: {temporary}")
    try:
        memo.render(document, temporary, rules=book.rules)
        rendered = Document(temporary)
        built_line = rendered.paragraphs[1]
        for paragraph in header:
            target = built_line.insert_paragraph_before()
            for run in paragraph.runs:
                target.add_run(run.text)
        rendered.save(temporary)
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)
    changes = casefile.changed_since(before)
    casefile = facts.append(path, [facts.Event("output", dict(name="memo", path=str(output),
        built_from_seq=casefile.last_seq, facts=counts["facts"], conflicts=counts["conflicts"],
        gaps=counts["gaps"], open=counts["open"]))])
    print(f'memo {number} chunks={len(chunks)} runs=+{runs} read={counts["read"]} '
          f'open={counts["open"]} facts={counts["facts"]} conflicts={counts["conflicts"]} '
          f'gaps={counts["gaps"]} new_facts={changes["fact"]} new_flags={changes["flag"]} '
          f'memo={output}', file=out)
    status.print_open(counts["open_items"], out)
    return 1 if counts["open"] else 0


def run(config, number, *, out, err, opener=None) -> int:
    try:
        book = playbook.load(config.paths.playbook)
        loaded = pages.load(config.paths.matters, number, config.ocr.stamp_pattern)
        if not loaded.pages and not loaded.open_pages:
            raise pages.PagesError(f"no pages to read in {number}")
        rules.check_stamps(loaded.pages, book.rules)
        # Open intake pages still contribute headings and labels to the memo.
        open_pages = tuple(Page(page_id.rsplit("/", 1)[0], int(page_id.rsplit("/", 1)[1]), "", stamps)
                           for page_id, stamps, _, _ in loaded.open_pages)
        folder = config.paths.matters / number
        path = folder / "casefile.jsonl"
        with (folder / "casefile.jsonl.lock").open("a") as lock:
            fcntl.flock(lock, fcntl.LOCK_EX)
            return _run_locked(config, number, folder, path, book, loaded, open_pages,
                               out=out, err=err, opener=opener)
    except ModelError as exc:
        print(f"memo: cannot run: {exc}; the pages of this run stay open until a later run reads them", file=err)
        return 2
    except (matterinfo.MatterInfoError, rules.RulesError, playbook.PlaybookError, pages.PagesError, facts.CaseFileError, memo.MemoError,
            OSError, sqlite3.Error) as exc:
        print(f"memo: cannot run: {exc}", file=err)
        return 2

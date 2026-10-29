"""Independent page and saved-document oracles for the brief package and draft checker."""

from dataclasses import replace
from pathlib import Path
import sys

from docx import Document
import pytest

from casefile import brief, cli, dates, facts, matterinfo, outputs, packet, quotes, roles
from conftest import ROOT
from test_outputs import setup_case

FIXTURES = ROOT / 'tests/fixtures/anonymizer/9001.001'
RULES_TEXT = 'Drafter rules supplied by the caller. Keep [DR-n] and [DATE D+n] unchanged.'


def contract_totals(results):
    return ' '.join(f'{kind}_ok={sum(rows)} of {len(rows)}' for kind, rows in results.items())


@pytest.fixture(scope='module')
def contract_results():
    return {'package': [], 'draft': []}


def draft_file(tmp_path, name='draft', *, table=False):
    paragraphs = (FIXTURES / (name + '.txt')).read_text(encoding='utf-8').splitlines()
    doc = Document()
    for value in paragraphs:
        doc.add_paragraph(value)
    if table:
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Ignored table'
    path = tmp_path / (name + '.docx')
    doc.save(path)
    return path, paragraphs


def quote_restore(value, page):
    # This table is independently transcribed from brief section 1b, not from the anonymizer.
    replacements = {'[DATE D+0]': '03/11/2025', '[DATE D+3]': '03/14/2025',
                    '[DATE D+13]': '03/24/2025', '[DATE D+27]': '04/07/2025',
                    '[DATE D+38]': '04/18/2025', '[DATE D+59]': '05/09/2025',
                    '[DATE D+101]': '06/20/2025', '[DR-2]': 'Ines Castellanos-Reyes'}
    for token, original in replacements.items():
        value = value.replace(token, original)
    return value if quotes.found_in(value, page.text) else None


def prose_restore(value):
    return value.replace('[CLAIMANT]', 'Morgan Ashby-Vale').replace('[DATE D+3]', '3/14/2025')


def check(draft, pages, rules, **kwargs):
    return brief.fact_check(draft, pages, rules,
                            restore_prose=kwargs.get('restore_prose', prose_restore),
                            restore_quote=kwargs.get('restore_quote', quote_restore))


def saved(tmp_path, report, info, rules):
    doc = brief.render(report, info, playbook_version='pb-invented', model='stand-in')
    path = tmp_path / 'checked.docx'
    outputs.render(doc, path, rules=rules)
    return [paragraph.text for paragraph in Document(path).paragraphs]


@pytest.mark.parametrize('row', range(9))
def test_package_contract(tmp_path, row, contract_results):
    contract_results['package'].append(False)
    _, _, folder, loaded, cf, book = setup_case(tmp_path)
    info = matterinfo.load(folder)
    value = packet.build(cf, loaded.pages, matter=loaded.number, task='brief', options={},
                         playbook_version=book.version)
    keep = ('Northgate Medical Evaluations', 'Brightwater Family Medicine')

    def package(**kwargs):
        return brief.package_text(value, cf, loaded.pages, info, issue='Review injury', keep=keep,
                                  rules_text=RULES_TEXT, transform=kwargs.pop('transform', lambda s: s), **kwargs)

    if row == 0:
        text = package()
        assert text.startswith(RULES_TEXT + '\n\nMatter 9001.001\n'), 'package rules parameter lost'
    elif row == 1:
        text = package(transform=lambda s: s.replace('Review injury', 'Review [CLAIMANT]'))
        assert 'Issue: Review [CLAIMANT]\n' in text, 'package issue not transformed'
        seen = []
        package(transform=lambda s: seen.append(s) or s)
        assert all(f.quote in seen and f.provider in seen for f in value.facts if f.provider)
        assert all(about in seen for flag in cf.flags for about in flag.abouts)
        assert all(p.text in seen for p in loaded.pages if any(p.page_id in flag.page_ids for flag in cf.flags if flag.kind == 'conflict'))
    elif row == 2:
        selected = (replace(value.facts[0], date=''), replace(value.facts[1], date='05/09/2025'),
                    replace(value.facts[2], date='03/14/2025'), replace(value.facts[3], date='03/14/2025'))
        value = replace(value, facts=selected)
        # Remove the original dates too, so this is a genuinely undated fact.
        events = [replace(event, payload={**event.payload, 'date': ''})
                  if event.kind == 'fact' and event.payload['id'] == selected[0].id else event for event in cf.events]
        cf = facts.CaseFile(events)

        def transform(text):
            parsed = dates.parse(text)
            return f'[DATE D{(parsed - dates.parse("03/11/2025")).days:+d}]' if parsed else text

        text = package(transform=transform)
        lines = text.split('## Facts\n\n', 1)[1].split('\n\n## Conflicts', 1)[0].split('\n\n')
        assert [line.rsplit('[cite ', 1)[1] for line in lines] == [f.cite + ']' for f in (selected[2], selected[3], selected[1], selected[0])]
        assert lines[0].startswith('[DATE D+3], '), 'fact date line shape lost'
        assert lines[-1].startswith('Undated, '), 'undated fact must be last'
    elif row == 3:
        body = package().split('## Facts\n\n', 1)[1].split('\n\n## Conflicts', 1)[0]
        assert body.count('[cite ') == len(value.facts)
        assert all(f'"{f.quote}" [cite {f.cite}]' in body for f in value.facts)
    elif row == 4:
        # Put an ours-only conflict first in storage; defense against ours must move ahead of it.
        cf = facts.CaseFile([*cf.events, facts.Event('flag', dict(id='G-test', kind='conflict',
            cites=['MED0001', 'MED0003'], page_ids=['medical/001', 'medical/003'],
            about='Local comparison', run_id=1))])
        cf.events.insert(0, cf.events.pop())
        text = package().split('## Conflicts\n\n', 1)[1].split('\n\n## Gaps', 1)[0]
        ordered = [flag for flag in cf.flags if flag.kind == 'conflict']
        by_id = {page.page_id: page for page in loaded.pages}
        ordered.sort(key=lambda flag: not {'defense', 'ours'} <= {
            roles.role(by_id[pid].binder, info) for pid in flag.page_ids})
        expected = []
        for flag in ordered:
            expected.extend("Model's reading: " + about for about in flag.abouts)
            for pid, cite in zip(flag.page_ids, flag.cites):
                expected.extend([f'From [cite {cite}]:', by_id[pid].text])
        assert text == '\n\n'.join(expected), 'package conflicts must include both full pages in defense-first order'
    elif row == 5:
        text = package()
        assert 'Real names kept in this package (never placeholders): ' + ', '.join(keep) in text
        assert text.split('## Gaps\n\n', 1)[1] == '\n\n'.join(about for flag in cf.flags if flag.kind == 'gap' for about in flag.abouts) + '\n'
    elif row == 6:
        files = {'package.md': package(), 'package.json': value.to_json(), 'check.txt': b'pass\n'}
        target = brief.publish_package(folder / 'outbox', 'abcdef123456', files)
        with pytest.raises(outputs.OutputError, match='already exists'):
            brief.publish_package(folder / 'outbox', 'abcdef123456', {'package.md': 'replacement'})
        assert (target / 'package.md').read_text(encoding='utf-8') == files['package.md']
        assert sorted(p.name for p in target.parent.iterdir()) == ['abcdef123456']
    elif row == 7:
        import os
        previous = os.umask(0o022)
        try:
            target = brief.publish_package(folder / 'outbox', 'abcdef123456',
                                           {'package.md': b'facts', 'package.json': b'{}', 'check.txt': b'pass'})
        finally:
            os.umask(previous)
        assert target.stat().st_mode & 0o777 == 0o700, 'package folder mode must be 700'
        assert all(p.stat().st_mode & 0o777 == 0o600 for p in target.iterdir()), 'package files mode must be 600'
    else:
        assert '\n## Pages\n' not in package(), 'default package unexpectedly includes all pages'
        text = package(pages_all=True, transform=lambda s: s.replace('Patient:', 'Subject:'))
        blocks = '\n\n'.join(f'=== PAGE {p.page_id} BATES {" ".join(p.stamps)} ===\n' +
                               p.text.replace('Patient:', 'Subject:') + '\n' for p in loaded.pages)
        assert text.split('## Pages\n\n', 1)[1] == blocks + '\n', 'all page blocks must preserve page order'
    contract_results['package'][-1] = True


@pytest.mark.parametrize('row', range(10))
def test_draft_contract(tmp_path, row, monkeypatch, contract_results):
    contract_results['draft'].append(False)
    _, _, folder, loaded, _, book = setup_case(tmp_path)
    rules, pages = book.rules, loaded.pages
    path, paragraphs = draft_file(tmp_path, table=True)
    draft = brief.read_draft(path)
    if row == 0:
        assert list(draft) == paragraphs and len(draft) == 12 and draft.skipped == 1
        document = Document(path)
        document.sections[0].header.paragraphs[0].text = 'Header ignored'
        document.sections[0].footer.paragraphs[0].text = 'Footer ignored'
        document.save(path)
        assert brief.read_draft(path).skipped == 3
    elif row == 1:
        units = brief.parse_draft(draft)
        assert sum(bool(u.cites) for u in units) == 10
        assert [u.index for u in units if not u.cites] == [10, 11]
        assert sum(len(u.quotes) for u in units) == 8
    elif row == 2:
        report = check(draft, pages, rules)
        assert report.ok and (report.checks, report.total, report.uncited, report.rule_hits) == (10, 10, (11, 12), 0), 'good draft must pass 10 of 10 including placeholder-form prose numbers'
        assert len(report.paragraphs) == 12
    elif row == 3:
        report = check(['"Pain 7/10 in the low back, no radiation." [cite MED9999]'], pages, rules)
        assert not report.ok and 'MED9999' in report.paragraphs[0].failures[0]
    elif row == 4:
        duplicate = replace(pages[1], stamps=('MED0001',))
        report = check(['"No significant chronic conditions reported." [cite MED0001]'], (*pages, duplicate), rules)
        assert not report.ok and 'ambiguous cite MED0001' in report.paragraphs[0].failures[0]
        # A cite in a quote must not change the paragraph's actual page association.
        embedded = '"[cite MED0001]" [cite MED0003]'
        unit = brief.parse_draft([embedded])[0]
        assert unit.cites == ('MED0003',), 'quoted cite became a paragraph cite'
        assert brief.parse_draft(['"[cite MED0001]"'])[0].cites == (), 'quoted-only cite became a paragraph cite'
    elif row == 5:
        bad_path, _ = draft_file(tmp_path, 'draft-bad-quote')
        report = check(brief.read_draft(bad_path), pages, rules)
        assert (report.checks, report.failed) == (9, 1)
        assert 'medical/001' in report.paragraphs[0].failures[0]
        # This quote is present in the matter, but only on MED0001.
        assert not check(['"Pain 7/10 in the low back, no radiation." [cite MED0003]'], pages, rules).ok
    elif row == 6:
        assert not check(['Pain 999/10 [cite MED0003]'], pages, rules).ok, 'absent page number accepted'
        assert check(['Pain 4/10 [cite MED0003]'], pages, rules).ok, 'digit check must read the page when no quote exists'
        assert check(['At [DATE D+3], pain 4/10 [cite MED0003]'], pages, rules).ok, 'placeholder offset counted as a prose number'
    elif row == 7:
        report = check(['The IME was reviewed [cite MED0003].'], pages, rules)
        text = '\n'.join(saved(tmp_path, report, matterinfo.load(folder), rules))
        assert report.rule_hits == 1 and report.failed == 1
        assert '[withheld: no-dme-ime]' in text and 'The IME' not in text
        # A verified quotation is exempt from output wording rules.
        report = check(['"Independent Medical Evaluation" [cite DEF0001]'], pages, rules)
        assert report.ok and report.rule_hits == 0
        report = check(['The independent [cite MED0003] medical examination was reviewed.'], pages, rules)
        assert report.rule_hits == 1 and not report.ok, 'citation split hid a firm-rule phrase'
    elif row == 8:
        report = check(['Argument'], pages, rules, restore_prose=lambda s: s + '\nnew paragraph')
        assert not report.ok and 'paragraph count changed' in report.paragraphs[0].failures[0]
    else:
        report = check(['"Not on this page." [cite MED0003]'], pages, rules)
        text = saved(tmp_path, report, matterinfo.load(folder), rules)
        assert any(t.startswith('[FACT CHECK FAILED: paragraph 1: quote not on medical/003;') for t in text), 'failed quote check not marked in saved document'
        assert text.index('Fact check') > text.index('Draft')
        assert 'Checks: 0 of 1; failures: 1; uncited: 0; rule hits: 0; skipped: 0' in text
        assert len(report.paragraphs) == 1
    contract_results['draft'][-1] = True


def test_contract_totals(capsys, contract_results):
    with capsys.disabled():
        print(contract_totals(contract_results))


@pytest.mark.parametrize('text', ['“Changed words.” [cite MED0001]',
    '“Unclosed quotation [cite MED0001]', '"Unclosed quotation [cite MED0001]'])
def test_l7_curly_or_unparsable_quote_fails(tmp_path, text):
    _, _, _, loaded, _, book = setup_case(tmp_path)
    report = check([text], loaded.pages, book.rules)
    assert (report.checks, report.total, report.failed) == (0, 1, 1), 'L7: quote verification silently skipped'


def test_l8_totals_measure_executed_cases():
    summary = globals().get('contract_totals')
    assert callable(summary), 'L8: contract totals cannot measure executed cases'
    assert summary({'package': [True, False], 'draft': [True]}) == 'package_ok=1 of 2 draft_ok=1 of 1'
    assert summary({'package': [], 'draft': []}) == 'package_ok=0 of 0 draft_ok=0 of 0'


@pytest.mark.parametrize('command, positional, options', [
    ('brief-out', 'Review injury', ['--pages', 'all']), ('brief-in', 'draft.docx', [])])
def test_brief_cli_arguments(monkeypatch, command, positional, options):
    def stop_after_parse(self, args=None, namespace=None):
        parsed = original(self, args, namespace)
        assert parsed.command == command and parsed.matter == '9001.001'
        assert parsed.config == Path('invented.toml')
        if command == 'brief-out':
            assert parsed.issue == positional and parsed.pages == 'all'
        else:
            assert parsed.file == Path(positional)
        raise SystemExit(0)

    import argparse
    original = argparse.ArgumentParser.parse_args
    monkeypatch.setattr(argparse.ArgumentParser, 'parse_args', stop_after_parse)
    monkeypatch.setattr(sys, 'argv', ['casefile', command, '--matter', '9001.001', positional,
                                    '--config', 'invented.toml', *options])
    with pytest.raises(SystemExit) as caught:
        cli.main()
    assert caught.value.code == 0


def test_package_rollback_owns_paths(tmp_path, monkeypatch):
    folder = tmp_path / 'outbox'
    folder.mkdir()
    sentinel = folder / 'unrelated'
    sentinel.write_text('keep', encoding='utf-8')
    monkeypatch.setattr(brief.os, 'replace', lambda *a: (_ for _ in ()).throw(KeyboardInterrupt()))
    with pytest.raises(KeyboardInterrupt):
        brief.publish_package(folder, 'abcdef123456', {'package.md': b'draft'})
    assert list(folder.iterdir()) == [sentinel]
    (folder / '.abcdef123456.tmp').mkdir()
    (folder / '.abcdef123456.tmp/owned-by-another-run').write_text('keep', encoding='utf-8')
    with pytest.raises(FileExistsError):
        brief.publish_package(folder, 'abcdef123456', {'package.md': b'draft'})
    assert (folder / '.abcdef123456.tmp/owned-by-another-run').read_text(encoding='utf-8') == 'keep'

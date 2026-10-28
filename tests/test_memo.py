"""Read the generated Word paragraphs and runs for the fourteen W5 checks."""

import copy
from pathlib import Path

from docx import Document
import pytest

from casefile import facts, matter, memo, rules


PAGES = (
    matter.Page("medical", 1, 'The IME records a normal gait. No prior injury.', ("MED0001",)),
    matter.Page("medical", 2, 'The IME records a normal gait.', ("MED0002",)),
    matter.Page("probe", 1, "Therapy continued.", ()),
)
QUOTE = 'The IME records a normal gait.'
READINGS = ["The report records a normal gait.", "The gait is recorded as normal."]
GAP_ABOUTS = ["The later record is missing.", "The next visit is not documented."]
OPEN = [memo.OpenItem("scan/003", memo.label(matter.Page("scan", 3, "", ("KMR2344", "TWP7731"))),
                      "empty", "no text after OCR"),
        memo.OpenItem("probe/002", "probe/002", "uncited", "no fact or flag cites probe/002")]


def case_file():
    stored = facts.CaseFile()
    stored.extend([
        facts.Event("fact", {"id": "F1", "cite": "MED0001", "page_id": "medical/001",
                              "quote": QUOTE, "text": READINGS[0], "run_id": 1}),
        facts.Event("cite", {"fact": "F1", "cite": "MED0002", "page_id": "medical/002", "run_id": 2}),
        facts.Event("reading", {"fact": "F1", "text": READINGS[1], "run_id": 2}),
        facts.Event("fact", {"id": "F4", "cite": "probe/001", "page_id": "probe/001",
                              "quote": "Therapy continued.", "text": "Therapy is recorded.", "run_id": 1}),
        facts.Event("fact", {"id": "F5", "cite": "MED0001", "page_id": "medical/001",
                              "quote": "No prior injury.", "text": "The history denies prior injury.", "run_id": 2}),
        facts.Event("flag", {"id": "G6", "kind": "conflict", "cites": ["MED0001", "DEF0002"],
                              "page_ids": ["medical/001", "defense-report/002"],
                              "about": "The histories disagree.", "run_id": 1}),
        facts.Event("flag", {"id": "G7", "kind": "gap", "cites": ["MED0002"],
                              "page_ids": ["medical/002"], "about": GAP_ABOUTS[0], "run_id": 1}),
        facts.Event("about", {"flag": "G7", "about": GAP_ABOUTS[1], "run_id": 2}),
    ])
    # Two runs read pages; a third run contributes no read coverage rows.
    for run_id, read in [(1, ["medical/001", "probe/001"]), (2, ["medical/002"]), (3, [])]:
        stored.extend([facts.Event("run", {"run_id": run_id, "job": "memo", "card": "bench_extract",
                                           "card_version": "invented", "pages": read or ["probe/002"],
                                           "ok": bool(read), "read": read,
                                           "uncited": [], "failed": [] if read else ["probe/002"]})])
    return stored


def build(stored=None, pages=PAGES, opened=OPEN):
    return memo.build(stored if stored is not None else case_file(), pages, opened,
                      matter="9009.001", playbook_version="invented-playbook", model="invented-model")


def section(document, title):
    paragraphs = document.paragraphs
    start = next(index for index, paragraph in enumerate(paragraphs)
                 if paragraph.style.name == "Heading 1" and paragraph.text == title) + 1
    end = next((index for index in range(start, len(paragraphs))
                if paragraphs[index].style.name == "Heading 1"), len(paragraphs))
    return paragraphs[start:end]


def test_memo(tmp_path):
    book_rules = rules.load(Path(__file__).resolve().parents[1] / "playbook" / "rules.toml")
    path = tmp_path / "outputs" / "memo.docx"
    built = build()
    memo.render(built, path, rules=book_rules)
    document = Document(path)
    paragraphs = document.paragraphs
    conflicts = section(document, "Conflicts")
    gaps = section(document, "Gaps and missing records")
    verified = section(document, "Verified facts by record")
    fact = next(paragraph for paragraph in verified if paragraph.text.startswith('"' + QUOTE))
    fact_index = next(index for index, paragraph in enumerate(verified) if paragraph.text == fact.text)
    heading_texts = [paragraph.text for paragraph in verified if paragraph.style.name == "Heading 2"]
    expected_quotes = {
        '"' + QUOTE + '" [cite MED0001]': "medical/001",
        '"No prior injury." [cite MED0001]': "medical/001",
        '"Therapy continued." [page probe/001]': "probe/001",
    }
    seen = {}
    current = None
    page_quotes = {}
    for paragraph in verified:
        if paragraph.style.name == "Heading 2":
            current = next((page.page_id for page in PAGES if paragraph.text.endswith(f'({page.page_id})')
                            or paragraph.text.endswith(f'page {page.page_id}')), None)
            page_quotes[current] = []
        elif paragraph.text.startswith('"'):
            seen[paragraph.text] = current
            page_quotes.setdefault(current, []).append(paragraph.text)
    warning = next(rule.rule for rule in book_rules.rules if rule.id == "bates-cites")
    checks = [
        paragraphs[0].style.name == "Title" and paragraphs[0].text == "Medical memo: matter 9009.001",
        paragraphs[1].text.startswith("built ") and
        f'from case file entry {case_file().last_seq}; playbook invented-playbook; model invented-model' in paragraphs[1].text,
        [p.text for p in section(document, "What changed since the last memo")] == ["First memo for this matter"],
        [p.text for p in section(document, "Outputs out of date")] == ["none"],
        [p.text for p in section(document, "Open items")] == [
            "KMR2344 TWP7731 (scan/003): empty: no text after OCR",
            "probe/002: uncited: no fact or flag cites probe/002"],
        [p.text for p in conflicts[1:]] == ["Conflict between [cite MED0001] and [cite DEF0002].",
                                             "Model's reading: The histories disagree."],
        [p.text for p in gaps] == ["Gap at [cite MED0002].", *[f"Model's reading: {s}" for s in GAP_ABOUTS]],
        fact.text == f'"{QUOTE}" [cite MED0001]' and [run.text for run in fact.runs] == [f'"{QUOTE}"', " [cite MED0001]"],
        verified[fact_index + 1].text == "Also cited: [cite MED0002]",
        [p.text for p in verified[fact_index + 2:fact_index + 4]] == [f"Model's reading: {s}" for s in READINGS],
        heading_texts[0] == "medical page MED0001 (medical/001)",
        verified[0].text == f'{warning} Unstamped pages: probe/001.' and
        '"Therapy continued." [page probe/001]' in [p.text for p in verified],
        heading_texts == ["medical page MED0001 (medical/001)", "medical page MED0002 (medical/002)",
                          "probe page probe/001"] and seen == expected_quotes and page_quotes.get("medical/002") == [],
        conflicts[0].text == "Pages read in 2 runs; conflicts and gaps are checked only between pages sent to the model together.",
    ]
    assert len(checks) == 14
    for number, ok in enumerate(checks, 1):
        assert ok, f'check {number} red'

    refused = 0
    for name in ("title", "binder"):
        bad = copy.deepcopy(built)
        if name == "title":
            bad.sections[0].title = "IME summary"
        else:
            bad = build(pages=(matter.Page("ime", 1, "Invented record.", ("MED0001",)),))
        output = tmp_path / "outputs" / f"{name}.docx"
        try:
            memo.render(bad, output, rules=book_rules)
        except memo.MemoError as exc:
            message = str(exc)
            required = (["IME summary"] if name == "title" else
                        ["binder ime", "no-dme-ime", "remove the binder by hand and intake it under another name (question 34)"])
            refused += int(all(part in message for part in required) and not output.exists())
    assert refused == 2, f'memo_refused={refused} of 2'

    stored = facts.CaseFile()
    stored.extend([facts.Event("output", {"name": "memo", "path": "/invented/outputs/memo.docx",
                                          "built_from_seq": 0, "facts": 0, "conflicts": 0, "gaps": 0, "open": 0}),
                   case_file().events[0]])
    changed_path = tmp_path / "changed.docx"
    memo.render(build(stored), changed_path, rules=book_rules)
    changed_doc = Document(changed_path)
    changed = int([p.text for p in section(changed_doc, "What changed since the last memo")] ==
                  ["1 fact added, 0 cites added, 0 readings added, 0 flags added"] and
                  [p.text for p in section(changed_doc, "Outputs out of date")] == ["memo"])
    assert changed == 1, f'memo_changed={changed}'

    reason = next(rule.rule for rule in book_rules.rules if rule.id == "no-dme-ime")
    reason_path = tmp_path / "reason.docx"
    reason_ok = 0
    try:
        memo.render(build(opened=[memo.OpenItem("probe/002", "probe/002", "failed", reason)]),
                    reason_path, rules=book_rules)
        reason_paragraphs = section(Document(reason_path), "Open items")
        reason_ok = int(len(reason_paragraphs) == 1 and
                        [run.text for run in reason_paragraphs[0].runs] == ["probe/002: failed: ", reason])
    except memo.MemoError:
        pass
    assert reason_ok == 1, f'memo_reason_exempt={reason_ok}'
    print("memo_docx_ok=14 of 14 memo_refused=2 of 2 memo_changed=1 memo_reason_exempt=1")


@pytest.mark.parametrize("count, noun", [(0, "runs"), (1, "run"), (2, "runs")])
def test_conflicts_sentence_counts_read_runs(tmp_path, count, noun):
    stored = facts.CaseFile()
    stored.extend([event for event in case_file().events
                   if event.kind != "run" or event.payload["run_id"] <= count])
    # Repeated coverage from one run and reads by another job do not add memo runs.
    if count:
        stored.extend([facts.Event("run", {"run_id": 1, "job": "memo", "read": ["medical/001"]})])
    stored.extend([facts.Event("run", {"run_id": 99, "job": "bench", "read": ["medical/001"]}),
                   facts.Event("run", {"run_id": 100, "job": "memo", "read": []})])
    path = tmp_path / "memo.docx"
    memo.render(build(stored), path, rules=rules.load(Path(__file__).resolve().parents[1] / "playbook/rules.toml"))
    assert section(Document(path), "Conflicts")[0].text == (
        f"Pages read in {count} {noun}; conflicts and gaps are checked only between pages sent to the model together.")

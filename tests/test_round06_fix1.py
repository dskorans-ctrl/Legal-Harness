"""Isolated reproductions for every ruled round 6 review row."""

from dataclasses import replace
from io import StringIO
import re
from types import SimpleNamespace

from docx import Document
import pytest

from casefile import cards, depo, facts, gates, hearing, matterinfo, memo, outputs, packet, report, roles, status
from test_matterinfo import Terminal
from test_outputs import setup_case


def text(doc):
    return '\n'.join(p.text for p in doc.paragraphs)


def test_h1(tmp_path):
    _, _, _, loaded, cf, book = setup_case(tmp_path)
    reason = 'gate.rules: The IME found a normal gait'
    item = memo.OpenItem('medical/001', 'MED0001', 'failed', reason)
    doc = outputs.Doc('Output', [], [outputs.Section('Failure', [outputs.failure('Primary', reason)]),
                                   outputs.open_items_section([item])])
    path = tmp_path / 'output.docx'
    outputs.render(doc, path, rules=book.rules)
    assert 'IME' not in text(Document(path)), 'H1: rejected model text escaped into output'
    assert '[withheld: no-dme-ime]' in text(Document(path))
    value = memo.build(cf, loaded.pages, [item], matter=loaded.number, playbook_version=book.version, model='invented')
    memo.render(value, path, rules=book.rules)
    assert 'IME' not in text(Document(path)), 'H1: rejected model text escaped into memo'
    assert item.reason == reason, 'H1: changed original reason'
    doc.sections[0].paragraphs = [outputs.failure('Primary', 'gate.quotes: the IME exam of Dr. Marchetti was paid for by the carrier')]
    outputs.render(doc, path, rules=book.rules)
    prose = re.sub(r'\[withheld: [a-z-]+\]', '', text(Document(path)))
    assert 'paid' not in prose and 'IME' not in prose, 'H1: payment excerpt escaped into output'


def test_h2_pages(tmp_path):
    _, _, folder, loaded, cf, book = setup_case(tmp_path)
    value = report.prepare(cf, loaded.pages, book, info=matterinfo.load(folder), number=loaded.number,
                           binder=('defense-report',))
    card = cards.card('05-A', book)
    # Repeated fact blocks force several batches without making any single page too large.
    value = replace(value, facts=value.facts * 30)
    parts = packet.batches(value, card, 22000)
    assert len(parts) > 1
    assert {p.page_id for part in parts for p in part.pages} == {p.page_id for p in value.pages}, 'H2: batch lost factless pages'


def test_h2_elements(tmp_path):
    _, _, _, loaded, _, _ = setup_case(tmp_path)
    a = dict(element='diagnosis', status='STATED', quote='Normal.', cite='MED0001', absent_note='')
    b = dict(a, status='PARTIAL', absent_note='Some findings absent')
    section, _, _ = outputs.elements({'elements': [a, b]}, loaded.pages)
    line = next(''.join(r.text for r in p.runs) for p in section.paragraphs if p.runs[0].text.startswith('Diagnosis:'))
    assert 'STATED' in line, 'H2: later partial overwrote stated element'
    scoped, _, _ = outputs.elements({'elements': [], 'searched_pages': [loaded.pages[0].page_id]}, loaded.pages)
    scope = ''.join(r.text for r in scoped.paragraphs[0].runs)
    assert 'MED0001' in scope and 'MED0002' not in scope, 'H2: negative claims pages not searched'


def test_m1(tmp_path):
    *_, book = setup_case(tmp_path)
    for value in ('What did you bill for this examination, Doctor?', 'What did Dr. Marchetti charge for the examination?',
                  'Doctor, were the bills reasonable?'):
        assert any(f.rule_id == 'no-hired-paid-retained' for f in book.rules.check(value, unit='question')), 'M1: non-bills-card fee question passed'


def test_m2():
    for name in ('IME-report', 'IME', 'DMEreport', 'DME2025', 'D.M.E.', 'Independent_medical', 'imereport', 'D.M.Ereport', 'I-M-E2025', 'D/M/E-report'):
        assert roles.role(name, matterinfo.MatterInfo()) == 'defense', 'M2: examiner binder assigned ours: ' + name
    assert roles.role('time-record', matterinfo.MatterInfo()) == 'ours'
    assert roles.role('IME', matterinfo.MatterInfo(plaintiffs=2)) == 'ROLE NEEDED'


def test_m3(tmp_path):
    _, cfg, folder, loaded, cf, book = setup_case(tmp_path)
    info = matterinfo.load(folder)
    with pytest.raises(matterinfo.MatterInfoError, match='hearing.issues'):
        hearing.prepare(cf, loaded.pages, book, info, number=loaded.number, outline=True, issue=('made up issue',))
    called = []
    def opener(*args, **kwargs):
        called.append(True)
        raise AssertionError('M3: called model before issue refusal')
    out, err = StringIO(), StringIO()
    assert hearing.run(cfg, loaded.number, out=out, err=err, opener=opener, outline=True, issue=('made up issue',)) == 2
    assert not called and 'hearing.issues' in err.getvalue()


def test_m4(tmp_path):
    _, _, _, loaded, cf, _ = setup_case(tmp_path)
    with pytest.raises(outputs.OutputError, match='ambiguous provider.*Marchetti.*Okonkwo-Lind'):
        outputs.provider_pages(cf, 'Dr', pages=loaded.pages)
    assert outputs.provider_pages(cf, 'archett', pages=loaded.pages) == (), 'M4: provider fragment matched'
    assert outputs.provider_pages(cf, 'Marchetti') == ('defense-report/002', 'defense-report/003')
    merged = facts.CaseFile(list(cf.events))
    merged.extend([facts.Event('cite', dict(fact=cf.facts[0].id, cite='MED0002', page_id='medical/002', run_id=2, date='', provider=cf.facts[0].providers[0].upper()))])
    try:
        selected = outputs.provider_pages(merged, cf.facts[0].providers[0])
    except outputs.OutputError:
        pytest.fail('M4: folded provider spellings became ambiguous')
    assert selected


def test_m5(tmp_path):
    _, cfg, folder, loaded, cf, book = setup_case(tmp_path)
    doc = outputs.Doc('Output', [], [outputs.open_items_section(())])
    counts = dict(facts=31, conflicts=5, gaps=4, open=0)
    cf = outputs.publish(folder, 'report-test', doc, book.rules, cf, counts, failed_steps=())
    cf = outputs.publish(folder, 'report-test', doc, book.rules, cf, counts, failed_steps=['elements'])
    state = status.matter_status(cfg, loaded.number)['outputs']
    assert 'report-test:stale' in state and 'report-test.partial.docx' in state, 'M5: failed replacement left prior output current'
    outputs.publish(folder, 'report-test', doc, book.rules, cf, counts, failed_steps=())
    assert not (folder / 'outputs/report-test.partial.docx').exists(), 'M5: success retained partial'


def test_l1():
    missing = depo.missing_questions({'questions': [{'text': 'Were the bills for the treatment sent to you?'}]},
                                     {'deponent': 'our-doctor', 'taken_by': 'us'})
    assert len(missing) == 2, 'L1: compound question satisfied separate topics'


def test_l2(tmp_path):
    *_, loaded, _, book = setup_case(tmp_path)
    page = loaded.pages[0]
    obj = {'undefined_labels': [dict(label='symptom magnification 999999 of 10', quote=page.text.splitlines()[0], cite=page.stamps[0])]}
    detail, _ = gates._references(obj, (page,), book.rules, numbers_in='page')
    assert 'label: number 999999 is not in its page' in detail, 'L2: unsupported label number passed'


def test_l3(tmp_path, monkeypatch):
    _, cfg, folder, loaded, cf, book = setup_case(tmp_path)
    bare = tuple(replace(p, stamps=()) for p in loaded.pages)
    started = SimpleNamespace(casefile=cf, loaded=replace(loaded, pages=bare), book=book,
        info=matterinfo.load(folder), client=None, log=None, counts={'open_items': ()})
    from contextlib import contextmanager
    @contextmanager
    def start(*args, **kwargs):
        yield started
    monkeypatch.setattr(outputs, 'start', start)
    monkeypatch.setattr(outputs, 'ask', lambda *a, **k: SimpleNamespace(ok=True, answer={'elements': []}))
    names = []
    monkeypatch.setattr(outputs, 'finish', lambda s, name, *a, **k: names.append(name) or 0)
    try:
        code = report.run(cfg, loaded.number, ours=True, selected='medical/007', out=StringIO(), err=StringIO())
    except IndexError:
        pytest.fail('L3: unstamped page selection raised IndexError')
    assert code == 0 and names == ['report-medical-007-medical-007'], 'L3: unstable unstamped report name'


def test_l4(tmp_path, monkeypatch):
    _, cfg, folder, loaded, _, _ = setup_case(tmp_path)
    info = replace(matterinfo.load(folder), has_header=False)
    monkeypatch.setattr(matterinfo, 'load', lambda _: info)
    calls = []
    monkeypatch.setattr(matterinfo, 'ask', lambda *a: calls.append('ask') or info)
    def refuse(*args):
        raise outputs.rules.RulesError('incompatible stamp')
    monkeypatch.setattr(outputs.rules, 'check_stamps', refuse)
    with pytest.raises(outputs.rules.RulesError, match='incompatible stamp'):
        with outputs.start(cfg, loaded.number, 'report', stdin=Terminal('')):
            pass
    assert calls == [], 'L4: attorney prompt preceded stamp refusal'


def test_l5(tmp_path):
    info = matterinfo.MatterInfo(folder=tmp_path)
    result = matterinfo.ask(info, Terminal('workers-comp\nInvented Client\n03/11/2025\n\n\n\n2\n1\n'), StringIO())
    assert result.plaintiffs == 2, 'L5: prompt did not record multiple plaintiffs'
    assert roles.role('medical', matterinfo.load(tmp_path)) == 'ROLE NEEDED'


def test_l6(tmp_path):
    _, _, folder, loaded, cf, book = setup_case(tmp_path)
    bare = tuple(replace(p, stamps=()) for p in loaded.pages)
    info = matterinfo.load(folder)
    value = report.prepare(cf, bare, book, number=loaded.number, selected='medical/007')
    documents = [report.render({'elements': {'elements': []}}, value, cf, bare, info, open_items=(), playbook_version=book.version, model='invented')]
    dv = depo.prepare(cf, bare, book, info, number=loaded.number, deponent='client', taken_by='us')
    documents.append(depo.render({}, dv, cf, bare, info, open_items=(), playbook_version=book.version, model='invented'))
    hv = hearing.prepare(cf, bare, book, info, number=loaded.number, outline=True)
    documents.append(hearing.render({}, hv, cf, bare, info, open_items=(), playbook_version=book.version, model='invented'))
    for i, doc in enumerate(documents):
        path = tmp_path / f'warning{i}.docx'
        outputs.render(doc, path, rules=book.rules)
        assert 'These pages carry no Bates stamp and are cited by page id:' in text(Document(path)), 'L6: unstamped output missing warning'


def test_l8(tmp_path):
    _, _, folder, loaded, cf, book = setup_case(tmp_path)
    value = memo.build(cf, loaded.pages, (), matter=loaded.number, playbook_version=book.version, model='invented', info=matterinfo.load(folder))
    section = next(s for s in value.sections if s.title == 'Opinions in other records')
    assert [''.join(r.text for r in p.runs) for p in section.paragraphs] == ['none'], 'L8: ordinary matter demands other role'


def test_l9(tmp_path):
    from test_stage6_exit import _rules
    *_, book = setup_case(tmp_path)
    doc = outputs.Doc('Output', [], [outputs.Section('Failure', [outputs.Paragraph([outputs.Run('The IME found a normal gait', exempt='reason')])])])
    # Build a corrupt Word file directly, bypassing the production renderer and its check.
    saved = Document()
    saved.add_heading(doc.title, 0)
    saved.add_heading('Failure', 1)
    saved.add_paragraph('The IME found a normal gait')
    path = tmp_path / 'corrupt.docx'
    saved.save(path)
    errors = []
    assert _rules(Document(path), doc, book.rules, (), errors) > 0, 'L9: exit oracle trusted production reason exemption'
    from test_stage6_exit import _saved_rules
    from casefile.matter import Page
    saved = Document()
    saved.add_paragraph('"The IME invented a finding" [cite MED0001]')
    saved.save(path)
    source = Page('medical', 1, 'Normal gait.', ('MED0001',))
    assert _saved_rules(Document(path), book.rules, (), (source,)) > 0, 'L9: oracle exempted quote absent from cited page'
    source = replace(source, text='The IME invented a finding')
    assert _saved_rules(Document(path), book.rules, (), (source,)) == 0, 'L9: oracle rejected verified source quote'


def test_l10(tmp_path):
    _, _, _, loaded, cf, book = setup_case(tmp_path)
    subset = loaded.pages[:1]
    value = packet.build(cf, subset, matter=loaded.number, task='report', options={}, playbook_version=book.version)
    assert all(f.page_id in {p.page_id for p in value.pages} for f in value.facts), 'L10: packet listed absent-page fact'


def test_m1_exam(tmp_path):
    *_, book = setup_case(tmp_path)
    for fee in ('bill', 'bills', 'billed', 'charge', 'charges', 'charged', 'fee', 'rate'):
        for event in ('examination', 'report', 'deposition', 'testimony'):
            for value in (f'What {fee} for this {event}?', f'For the {event}, what {fee}?'):
                assert any(f.rule_id == 'no-hired-paid-retained' for f in book.rules.check(
                    value, unit='question', bills_question=True)), 'M1: examination fee passed: ' + value


def test_m1_cards(tmp_path):
    from casefile.model import ModelClient
    from casefile.runlog import open_db
    from fakemodel import FakeModel, FakeTransport
    from fakemodel.bench import recording_for
    _, cfg, _, loaded, cf, book = setup_case(tmp_path)
    value = packet.build(cf, loaded.pages, matter=loaded.number, task='depo', options={}, playbook_version=book.version)
    row = value.facts[0]
    answer = dict(section='primary', questions=[dict(text='Doctor, were the bills reasonable?', quote=row.quote, cite=row.cite, note='')], points_for_our_expert=[])
    with open_db(cfg.paths.state / 'runlog.sqlite') as log:
        for name in ('02-C-OURDOC', '02-C-DEFDOC', '02-C-CLIENT', '02-C-WITNESS', '03-CLIENT'):
            card = cards.card(name, book)
            fake = FakeModel(recordings=[recording_for(card, value.pages, loaded.number, answer, body=value.body())])
            outcome = outputs.ask(value, card, client=ModelClient(cfg.model, FakeTransport(fake)),
                rules=book.rules, log=log, job='depo', playbook_version=book.version)
            assert outcome.ok == (name == '02-C-OURDOC'), 'M1: wrong bills exemption on card ' + name

"""Twenty-two case file checks, including identity and append-only split controls."""

import copy
import hashlib
import fcntl
from io import StringIO
import shutil
from unittest.mock import patch

from docx import Document
import json
from pathlib import Path

from casefile import config, extract, facts, matter, memo_job, outputs, status
from casefile.runlog import open_db
from conftest import ROOT, write_config
from fakemodel import FakeModel, FakeTransport


FIXTURE = Path(__file__).parent / "fixtures" / "matters" / "9001.001"
FIXED_AT = "2026-10-28T00:00:00+00:00"
QUOTE = 'He said "no prior injury".'
PAGES = (
    matter.Page("medical", 1, QUOTE, ("MED0001",)),
    matter.Page("medical", 2, QUOTE, ("MED0002",)),
    matter.Page("probe", 1, QUOTE, ()),
)
ANSWER = {"facts": [{"cite": "MED0001", "quote": QUOTE, "text": "He reported no prior injury."}],
          "flags": [{"kind": "conflict", "cites": ["MED0001", "MED0002"],
                     "about": "The reports disagree."}]}


def fact_event(quote="The record was received."):
    return facts.Event("fact", {"id": "F1", "cite": "MED0001", "page_id": "medical/001",
                                "quote": quote, "text": "The record was received.", "run_id": 1,
                                "date": "", "category": "other", "provider": ""},
                       at=FIXED_AT)


def output_event(seq):
    return facts.Event("output", {"name": "memo", "path": "/invented/outputs/memo.docx",
                                  "built_from_seq": seq, "facts": 1, "conflicts": 0,
                                  "gaps": 0, "open": 0}, at=FIXED_AT)


def three_events(path):
    facts.append(path, [fact_event()])
    prefix = path.read_bytes()
    run = facts.Event("run", {"run_id": 1, "job": "memo", "card": "medical_memo",
                              "card_version": "invented", "pages": ["medical/001"],
                              "ok": True, "read": ["medical/001"], "uncited": [], "failed": []},
                      at=FIXED_AT)
    facts.append(path, [run, output_event(2)])
    return prefix


def refusal(path, line, message):
    try:
        facts.load(path)
    except facts.CaseFileError as exc:
        assert str(exc) == (f"case file chain broken at line {line}; "
                            "the file was edited or truncated"), message
    else:
        raise AssertionError(message)


def round_trip(path):
    assert facts.load(path).last_seq == 0
    prefix = three_events(path)
    loaded = facts.load(path)
    lines = path.read_bytes().splitlines()
    assert len(lines) == 3 and path.read_bytes().startswith(prefix), "W4(d): append overwrote complete events"
    assert [event.seq for event in loaded.events] == [1, 2, 3]
    assert [event.prev for event in loaded.events] == ["", *[hashlib.sha256(line).hexdigest() for line in lines[:2]]]
    assert loaded.facts[0].quote == "The record was received."
    assert loaded.runs[0]["job"] == "memo" and loaded.outputs[0]["built_from_seq"] == 2


def edited_line(path):
    three_events(path)
    raw = path.read_bytes()
    assert raw.count(b"The record was received.") == 2
    path.write_bytes(raw.replace(b"The record was received.", b"The record was altered.", 1))
    refusal(path, 2, "W4(a): edited line accepted")


def removed_middle(path):
    three_events(path)
    lines = path.read_bytes().splitlines(keepends=True)
    path.write_bytes(lines[0] + lines[2])
    refusal(path, 2, "W4: removed middle line accepted")


def cut_middle(path):
    three_events(path)
    lines = path.read_bytes().splitlines(keepends=True)
    path.write_bytes(lines[0] + lines[1][:20] + b"\n" + lines[2])
    refusal(path, 2, "W4(h): cut middle line accepted")


def torn_tail(path):
    three_events(path)
    lines = path.read_bytes().splitlines(keepends=True)
    tail = lines[2][:25]
    prefix = b"".join(lines[:2])
    path.write_bytes(prefix + tail)
    before = path.read_bytes()
    try:
        loaded = facts.load(path)
    except facts.CaseFileError as exc:
        raise AssertionError("W4(f): torn tail refused") from exc
    assert loaded.torn == len(tail) and loaded.last_seq == 2, "W4(f): torn tail counted as an event"
    assert path.read_bytes() == before, "load rewrote the torn file"
    try:
        facts.append(path, [output_event(2)])
        recovered = facts.load(path)
    except facts.CaseFileError as exc:
        raise AssertionError("W4(g): append did not discard the torn tail") from exc
    assert recovered.last_seq == 3 and recovered.torn == 0, "W4(g): append did not discard the torn tail"
    assert path.read_bytes().startswith(prefix)
    assert len(path.read_bytes().splitlines()) == 3


def missing_newline(path):
    three_events(path)
    path.write_bytes(path.read_bytes()[:-1])
    prefix = path.read_bytes()
    loaded = facts.load(path)
    assert loaded.last_seq == 3 and loaded.torn == 0, "complete last object treated as torn"
    facts.append(path, [output_event(3)])
    loaded = facts.load(path)
    assert loaded.last_seq == 4 and path.read_bytes().startswith(prefix + b"\n")


def ideal_merge(path):
    pages = matter.load(FIXTURE).pages
    answer = json.loads((FIXTURE / "ideal.json").read_text(encoding="utf-8"))
    events = facts.merge(facts.CaseFile(), answer, pages, 7)
    assert sum(event.kind == "fact" for event in events) == 38
    assert sum(event.kind == "flag" for event in events) == 9
    stored = facts.append(path, events)
    assert len(stored.facts) == 38 and len(stored.flags) == 9 and stored.last_seq == 47
    assert stored.flags[0].page_ids == ["medical/001", "employer/001"], "G1 pages were not resolved"
    known = {page.page_id for page in pages}
    assert all(set(record.page_ids) <= known for record in [*stored.facts, *stored.flags])
    assert all(event.payload["run_id"] == 7 for event in stored.events)
    assert [event.payload["id"] for event in events if event.kind in ("fact", "flag")] == [
        *[f"F{seq}" for seq in range(1, 39)], *[f"G{seq}" for seq in range(39, 48)]]


def repeat_merge(path):
    pages = matter.load(FIXTURE).pages
    answer = json.loads((FIXTURE / "ideal.json").read_text(encoding="utf-8"))
    stored = facts.append(path, facts.merge(facts.CaseFile(), answer, pages, 1))
    assert facts.merge(stored, answer, pages, 2) == [], "repeat ideal answer added events"
    stored = facts.CaseFile()
    stored.extend(facts.merge(stored, ANSWER, PAGES, 1))
    variant = copy.deepcopy(ANSWER)
    variant["facts"][0]["quote"] = "He said “no prior injury”."
    events = facts.merge(stored, variant, PAGES, 2)
    assert len(events) == 0, f"W4(b): folded repeat added {len(events)} event(s)"
    assert len(stored.events) == 2, "merge mutated the case file"


def new_cite(path):
    stored = facts.append(path, facts.merge(facts.CaseFile(), ANSWER, PAGES, 1))
    answer = copy.deepcopy(ANSWER)
    answer["facts"][0]["cite"] = "MED 0002"
    events = facts.merge(stored, answer, PAGES, 2)
    assert [event.kind for event in events] == ["cite", "reading"], "new page lost its cite or reading source"
    assert events[0].payload == {"fact": "F1", "cite": "MED0002", "page_id": "medical/002", "run_id": 2, "date": "", "provider": ""}
    assert events[1].payload == {"fact": "F1", "text": answer["facts"][0]["text"],
                                 "page_id": "medical/002", "run_id": 2}
    stored = facts.append(path, events)
    assert stored.facts[0].page_ids == ["medical/001", "medical/002"]
    assert stored.facts[0].cites == ["MED0001", "MED0002"]
    assert facts.merge(stored, answer, PAGES, 3) == []
    answer["facts"][0]["cite"] = "probe/001"
    assert facts.merge(stored, answer, PAGES, 3)[0].payload["page_id"] == "probe/001"


def new_reading(path):
    stored = facts.append(path, facts.merge(facts.CaseFile(), ANSWER, PAGES, 1))
    answer = copy.deepcopy(ANSWER)
    answer["facts"][0]["text"] = "The history records a denial of prior injury."
    events = facts.merge(stored, answer, PAGES, 2)
    assert len(events) == 1 and events[0].kind == "reading", "new reading was lost or duplicated"
    assert events[0].payload["fact"] == "F1" and events[0].payload["run_id"] == 2
    stored = facts.append(path, events)
    assert stored.facts[0].readings == [ANSWER["facts"][0]["text"], answer["facts"][0]["text"]]
    assert facts.merge(stored, answer, PAGES, 3) == []


def new_about(path):
    stored = facts.append(path, facts.merge(facts.CaseFile(), ANSWER, PAGES, 1))
    answer = copy.deepcopy(ANSWER)
    answer["flags"][0]["about"] = "The two records give different histories."
    events = facts.merge(stored, answer, PAGES, 2)
    assert len(events) == 1 and events[0].kind == "about", "W4(c): known flag became a new flag"
    assert events[0].payload["flag"] == "G2" and events[0].payload["run_id"] == 2
    stored = facts.append(path, events)
    assert stored.flags[0].abouts == [ANSWER["flags"][0]["about"], answer["flags"][0]["about"]]
    assert facts.merge(stored, answer, PAGES, 3) == []


def reordered_cites(path):
    stored = facts.append(path, facts.merge(facts.CaseFile(), ANSWER, PAGES, 1))
    answer = copy.deepcopy(ANSWER)
    answer["flags"][0]["cites"] = ["medical/002", "MED-0001"]
    assert facts.merge(stored, answer, PAGES, 2) == [], "W4(c): reordered flag cites added events"


def changed_since(path):
    stored = facts.append(path, [fact_event(), output_event(1)])
    before = stored.last_seq
    answer = copy.deepcopy(ANSWER)
    answer["flags"][0]["kind"] = "gap"
    stored = facts.append(path, facts.merge(stored, answer, PAGES, 2))
    answer["facts"][0]["cite"] = "MED0002"
    answer["facts"][0]["text"] = "A second reading."
    answer["flags"][0]["about"] = "A second gap description."
    stored = facts.append(path, facts.merge(stored, answer, PAGES, 3))
    assert stored.changed_since(before) == {"fact": 1, "cite": 1, "reading": 1, "flag": 1, "about": 1, "split": 0}
    assert stored.changed_since(stored.last_seq) == dict.fromkeys(facts.CHANGE_KINDS, 0)


def stale_outputs(path):
    stored = facts.append(path, [fact_event(), output_event(1)])
    assert stored.stale_outputs() == [], "fresh memo marked stale"
    event = fact_event("Another received record.")
    event = facts.Event("fact", {**event.payload, "id": "F3"}, at=FIXED_AT)
    stored = facts.append(path, [event])
    assert {event.at for event in stored.events} == {FIXED_AT}
    assert stored.stale_outputs() == ["memo"], "W4(e): new fact with the same timestamp left memo current"
    stored = facts.append(path, [output_event(stored.last_seq)])
    assert stored.stale_outputs() == [], "rebuilt memo still stale"


CASES = (round_trip, edited_line, removed_middle, cut_middle, torn_tail, missing_newline,
         ideal_merge, repeat_merge, new_cite, new_reading, new_about, reordered_cites,
         changed_since, stale_outputs)


def test_casefile(tmp_path):
    assert len(CASES) == 22
    completed = 0
    for case in CASES:
        case(tmp_path / f"{case.__name__}.jsonl")
        completed += 1
    print(f"casefile_ok={completed} of {len(CASES)}")


def legacy_bytes(events):
    result=[];previous=''
    for index,event in enumerate(events,1):
        data={k:v for k,v in event.payload.items() if k in facts.FIELDS_V1[event.kind]}
        obj=dict(seq=index,at=FIXED_AT,prev=previous,kind=event.kind,payload=data)
        raw=json.dumps(obj,ensure_ascii=False,separators=(',',':')).encode('utf-8')
        result.append(raw+b'\n');previous=hashlib.sha256(raw).hexdigest()
    return b''.join(result)


def upgrade(path):
    raw=legacy_bytes([fact_event(),facts.Event('cite',dict(fact='F1',cite='MED0002',page_id='medical/002',run_id=2))])
    path.write_bytes(raw)
    loaded=facts.load(path)
    assert loaded.upgraded == 2 and path.read_bytes() == raw
    assert loaded.split_needed == 0 and facts.upgrade_events(loaded) == []
    fact=loaded.facts[0]
    assert fact.dates == ['',''] and fact.category == 'other' and fact.providers == ['','']
    event=fact_event('A new record.');event=facts.Event('fact',{**event.payload,'id':'F3'})
    updated=facts.append(path,[event])
    assert updated.upgraded == 2 and updated.last_seq == 3 and path.read_bytes().startswith(raw)
    assert set(updated.events[-1].payload) == facts.FIELDS['fact']
    path.write_bytes(raw+b'{"seq":3')
    assert facts.load(path).torn and facts.load(path).upgraded == 2
    assert facts.append(path,[event]).last_seq == 3


def legacy_corrupt(path):
    # Rehash the next line to isolate the shape guard from the chain guard.
    events=[fact_event(),output_event(1)]
    raw=legacy_bytes(events);lines=raw.splitlines();obj=json.loads(lines[0]);del obj['payload']['text']
    lines[0]=json.dumps(obj,separators=(',',':')).encode()
    second=json.loads(lines[1]);second['prev']=hashlib.sha256(lines[0]).hexdigest()
    lines[1]=json.dumps(second,separators=(',',':')).encode()
    path.write_bytes(b'\n'.join(lines)+b'\n')
    refusal(path,1,'W1(f2): old key subset accepted')


def cite_dates(path):
    answer=copy.deepcopy(ANSWER);answer['facts'][0].update(date='03/11/2025',category='patient_statement',provider='Invented Clinic')
    stored=facts.append(path,facts.merge(facts.CaseFile(),answer,PAGES,1))
    answer['facts'][0].update(cite='MED0002',date='03/12/2025',provider='Other Invented Clinic')
    stored=facts.append(path,facts.merge(stored,answer,PAGES,2))
    assert [fact.dates for fact in stored.facts] == [['03/11/2025'], ['03/12/2025']]
    assert [fact.providers for fact in stored.facts] == [['Invented Clinic'], ['Other Invented Clinic']]
    assert stored.facts[0].category == 'patient_statement'


CASES = (*CASES,upgrade,legacy_corrupt,cite_dates)



def merged_pair_events():
    """The invented F05/F34 pair as stored by round 6, with both reading sources."""
    answer = json.loads((FIXTURE / "ideal.json").read_text(encoding="utf-8"))
    pair = [item for item in answer["facts"]
            if item["quote"] == "Straight leg raise negative bilaterally."]
    assert [item["cite"] for item in pair] == ["MED0002", "DEF0002"]
    first, second = pair
    return [facts.Event("fact", dict(id="F1", page_id="medical/002", run_id=1, **first)),
            facts.Event("cite", dict(fact="F1", page_id="defense-report/002", run_id=1,
                cite=second["cite"], date=second["date"], provider=second["provider"])),
            facts.Event("reading", dict(fact="F1", page_id="defense-report/002", run_id=1,
                text=second["text"])),
            facts.Event("reading", dict(fact="F1", page_id="defense-report/002", run_id=2,
                text="The later examination records the negative finding."))]


def split_on_append(path):
    stored = facts.append(path, merged_pair_events())
    before = path.read_bytes()
    assert stored.split_needed == 1 and facts.load(path).split_needed == 1
    assert path.read_bytes() == before, "W1: load rewrote the case file"
    events = facts.upgrade_events(stored)
    assert len(events) == 1 and events[0].kind == "split"
    assert events[0].payload == dict(fact="F1", page_id="defense-report/002", to="F5",
        cite="DEF0002", date="06/20/2025", provider="Dr. Theodore Marchetti",
        text=stored.facts[0].readings[1])
    updated = facts.append(path, events)
    assert len(updated.facts) == 2, "split_on_append: split view did not create two facts"
    first, second = updated.facts
    assert first.page_ids == ["medical/002"] and first.cites == ["MED0002"]
    assert first.dates == ["03/11/2025"] and len(first.providers) == 1
    assert first.reading_pages == ["medical/002"]
    assert first.readings == stored.facts[0].readings[:1]
    assert second.id == "F5" and second.page_ids == ["defense-report/002"]
    assert second.cites == ["DEF0002"] and second.dates == ["06/20/2025"]
    assert second.providers == ["Dr. Theodore Marchetti"]
    assert second.reading_pages == ["defense-report/002"] * 2
    assert second.readings == stored.facts[0].readings[1:]
    assert second.quote == first.quote and second.category == first.category
    assert updated.split_needed == 0 and facts.upgrade_events(updated) == []
    assert path.read_bytes().startswith(before), "W1: split rewrote complete lines"
    extra = facts.Event("reading", dict(fact=second.id, page_id="defense-report/002",
        text="Another later reading.", run_id=3))
    assert facts.append(path, [extra]).facts[1].readings[-1] == extra.payload["text"]
    # Multiple later cites split in cite order, leaving the first identity group intact.
    grouped = path.with_name("split-groups.jsonl")
    source = merged_pair_events()
    source.extend([
        facts.Event("cite", dict(fact="F1", page_id="medical/003", run_id=2,
            cite="MED0003", date="3/11/2025", provider=source[0].payload["provider"])),
        facts.Event("cite", dict(fact="F1", page_id="defense-report/003", run_id=2,
            cite="DEF0003", date="06/20/2025", provider="Dr. Theodore Marchetti")),
        facts.Event("reading", dict(fact="F1", page_id="defense-report/003", run_id=2,
            text="Another page of the later examination."))])
    old = facts.append(grouped, source)
    splits = facts.upgrade_events(old)
    assert len(splits) == 1 and splits[0].payload['page_id'] == 'defense-report/002'
    assert [move['page_id'] for move in splits[0].payload['moves']] == ['defense-report/003']
    new = facts.append(grouped, splits)
    assert new.split_needed == 0 and len(new.facts) == 2, 'M4: one identity became two facts'
    assert new.facts[0].page_ids == ["medical/002", "medical/003"]
    assert new.facts[1].page_ids == ['defense-report/002', 'defense-report/003']
    assert new.facts[1].readings == old.facts[0].readings[1:]
    assert new.facts[1].reading_pages == ['defense-report/002'] * 2 + ['defense-report/003']
    assert facts.upgrade_events(new) == []


def same_identity_still_cites(path):
    answer = copy.deepcopy(ANSWER)
    answer["flags"] = []
    answer["facts"][0].update(date="03/11/2025", provider="Invented Clinic")
    stored = facts.append(path, facts.merge(facts.CaseFile(), answer, PAGES, 1))
    answer["facts"][0].update(cite="MED0002", date="3/11/2025", provider="INVENTED CLINIC")
    events = facts.merge(stored, answer, PAGES, 2)
    assert [event.kind for event in events] == ["cite", "reading"], \
        "same_identity_still_cites: parsed date or folded provider failed to merge"
    updated = facts.append(path, events)
    assert len(updated.facts) == 1 and updated.split_needed == 0
    assert facts.upgrade_events(updated) == []


def different_date_new_fact(path):
    pages = matter.load(FIXTURE).pages
    pair = merged_pair_events()
    first, second = pair[0].payload, pair[1].payload
    answer = {"facts": [{key: value for key, value in first.items()
                         if key not in ("id", "run_id", "page_id")},
                        dict(cite=second["cite"], quote=first["quote"], text=pair[2].payload["text"],
                             date=second["date"], provider=second["provider"], category=first["category"])],
              "flags": []}
    events = facts.merge(facts.CaseFile(), answer, pages, 1)
    assert [event.kind for event in events] == ["fact", "fact"], \
        "different_date_new_fact: different dates merged"
    assert len(facts.append(path, events).facts) == 2
    # Isolate date and provider: either dimension alone must create a fact.
    answer["facts"][1]["provider"] = first["provider"]
    assert [event.kind for event in facts.merge(facts.CaseFile(), answer, pages, 1)] == ["fact", "fact"], \
        "different_date_new_fact: date-only difference merged"
    answer["facts"][1].update(date=first["date"], provider="Other Invented Clinic")
    assert [event.kind for event in facts.merge(facts.CaseFile(), answer, pages, 1)] == ["fact", "fact"], \
        "different_date_new_fact: provider-only difference merged"


def split_refused(path):
    source = facts.CaseFile()
    source.extend(merged_pair_events())
    event = facts.upgrade_events(source)[0]
    for key, value, message in (
            ("fact", "UNKNOWN", "split_refused: unknown origin accepted"),
            ("to", "F1", "split_refused: reused to id accepted")):
        bad = facts.CaseFile(list(source.events))
        bad.extend([facts.Event("split", {**event.payload, key: value})])
        # Encode the wire independently; the shape and chain are otherwise valid.
        path.write_bytes(b"".join(json.dumps(dict(seq=e.seq, at=e.at, prev=e.prev,
            kind=e.kind, payload=e.payload), ensure_ascii=False,
            separators=(",", ":")).encode("utf-8") + b"\n" for e in bad.events))
        refusal(path, 5, message)
    facts.append(path.with_name("split-valid.jsonl"), [*merged_pair_events(), event])


def split_through_start(path):
    for job in ("report", "memo"):
        root = path.parent / job
        root.mkdir()
        settings = config.load(write_config(root, playbook=ROOT / "playbook"))
        folder = settings.paths.matters / "9001.001"
        shutil.copytree(FIXTURE, folder)
        target = folder / "casefile.jsonl"
        loaded = matter.load(folder)
        card = extract.card()
        with open_db(settings.paths.state / "runlog.sqlite") as log:
            run_id = log.start_run(job="memo", matter="9001.001", card=card.name,
                card_version=card.version, playbook_version="invented", model="invented",
                pages_in=len(loaded.pages), prompt_sha256="invented")
            for page in loaded.pages:
                log.record_coverage(run_id, matter="9001.001", page_id=page.page_id,
                                    card=card.name, status="read", reason="")
        run = facts.Event("run", dict(run_id=run_id, job="memo", card=card.name,
            card_version=card.version, pages=[p.page_id for p in loaded.pages], ok=True,
            read=[p.page_id for p in loaded.pages], uncited=[], failed=[]))
        old = facts.append(target, [*merged_pair_events(), run, output_event(5)])
        prefix = target.read_bytes()
        assert facts.load(target).split_needed == 1
        out, err = StringIO(), StringIO()
        assert status.run(settings, "9001.001", out=out, err=err) == 0
        assert "gaps=0 split_needed=1 outputs=memo:current" in out.getvalue()
        assert target.read_bytes() == prefix, "W1: status wrote an upgrade"
        out, err = StringIO(), StringIO()
        transport = FakeTransport(FakeModel(recordings=[]))
        append = facts.append
        def locked_append(*args, **kwargs):
            with (folder / "casefile.jsonl.lock").open("a") as competitor:
                try:
                    fcntl.flock(competitor, fcntl.LOCK_SH | fcntl.LOCK_NB)
                except BlockingIOError:
                    pass
                else:
                    raise AssertionError("W1: split append outside the matter lock")
            return append(*args, **kwargs)
        with patch.object(facts, "append", locked_append):
            if job == "report":
                with outputs.start(settings, "9001.001", job, opener=transport, out=out,
                                   stdin=StringIO()) as started:
                    assert started.casefile.split_needed == 0, \
                        "split_through_start: split_needed=1 after outputs.start"
                    assert [event.kind for event in started.casefile.events[len(old.events):]] == ["split"]
                    assert started.casefile.stale_outputs() == ["memo"], \
                        "split_through_start: split left memo current"
                    assert started.casefile.changed_since(old.last_seq).get("split") == 1
            else:
                code = memo_job.run(settings, "9001.001", out=out, err=err, opener=transport)
                assert code == 0 and not err.getvalue()
                assert "chunks=0 runs=+0" in out.getvalue() and "split=1" in out.getvalue(), \
                    "split_through_start: memo upgrade missing"
                doc = Document(folder / "outputs/memo.docx")
                index = next(i for i, p in enumerate(doc.paragraphs)
                             if p.text == "What changed since the last memo")
                assert doc.paragraphs[index + 1].text == (
                    "1 fact added, 0 cites added, 0 readings added, 0 flags added, 1 fact split"), \
                    "split_through_start: memo omitted split change"
        updated = facts.load(target)
        assert updated.split_needed == 0 and len(updated.facts) == 2
        assert sum(event.kind == "split" for event in updated.events) == 1
        assert target.read_bytes().startswith(prefix) and not transport.requests
        assert out.getvalue().count("split=1") == 1
        with outputs.start(settings, "9001.001", job, opener=transport, out=StringIO(),
                           stdin=StringIO()) as started:
            assert started.casefile.last_seq == updated.last_seq, "W1: upgrade repeated"


CASES = (*CASES, split_on_append, same_identity_still_cites, different_date_new_fact,
         split_refused, split_through_start)

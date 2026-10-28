"""Memo and status integration uses invented benchmark pages and in-process replies."""

from io import StringIO
import json
import os
import shutil
import sys

from docx import Document
import pytest

from casefile import cli, config, extract, facts, matter, memo, memo_job, model, pages, pipeline, playbook, status
from casefile.runlog import open_db
from conftest import split_card, ROOT, write_config
from fakemodel import FakeModel, FakeTransport
from fakemodel.bench import recording_for, recordings_for


FIXTURE = ROOT / "tests/fixtures/matters/9001.001"


@pytest.fixture(scope="module")
def totals():
    checked = {"status": 0, "binder": 0}
    yield checked
    print(f'status_ok={checked["status"]} of 6 binder_shown={checked["binder"]}', flush=True)


def setup_matter(tmp_path):
    path = write_config(tmp_path)
    settings = config.load(path)
    shutil.copytree(ROOT / "playbook", settings.paths.playbook, dirs_exist_ok=True)
    folder = settings.paths.matters / "9001.001"
    shutil.copytree(FIXTURE, folder)
    return path, settings, folder


def invoke(settings, transport):
    out, err = StringIO(), StringIO()
    code = memo_job.run(settings, "9001.001", out=out, err=err, opener=transport)
    return code, out.getvalue(), err.getvalue()


def test_status_six_cases(tmp_path, monkeypatch, capsys, totals):
    config_path, settings, folder = setup_matter(tmp_path)
    transport = FakeTransport(FakeModel(recordings=recordings_for(FIXTURE)))
    with open_db(settings.paths.state / "runlog.sqlite") as log:
        result = pipeline.bench_run(matter.load(folder), settings, log,
                                    playbook.load(settings.paths.playbook).version, opener=transport)
        assert result.pages_failed == 0
    before = status.matter_status(settings, "9001.001")
    assert {key: value for key, value in before.items() if key != "open_items"} == dict(
        pages=14, read=0, failed=0, uncited=0, empty=0, open=14, facts=0,
        conflicts=0, gaps=0, outputs="none"), "status before memo mixed coverage jobs"
    assert all(item.status == "unread" and item.reason == "no run has read this page"
               for item in before["open_items"])
    code, line, error = invoke(settings, transport)
    assert code == 1, "memo with open items must exit=1"
    assert error == ""
    assert line.splitlines()[0] == (
        "memo 9001.001 chunks=1 runs=+1 read=13 open=1 facts=31 conflicts=5 gaps=4 "
        f"new_facts=31 new_flags=9 memo={folder / 'outputs/memo.docx'}")
    current = status.matter_status(settings, "9001.001")
    current_ok = {key: value for key, value in current.items() if key != "open_items"} == dict(
        pages=14, read=13, failed=0, uncited=1, empty=0, open=1, facts=31,
        conflicts=5, gaps=4, outputs="memo:current")
    opened = "  open: DEF0001 (defense-report/001): uncited: no fact or flag cites defense-report/001"
    assert line.splitlines()[1:] == [opened], "open label lost its stamp or page ID"
    paragraphs = [p.text for p in Document(folder / "outputs/memo.docx").paragraphs]
    assert "Medical memo: matter 9001.001" in paragraphs
    assert "DEF0001 (defense-report/001): uncited: no fact or flag cites defense-report/001" in paragraphs

    # A failed re-send may record its run, but may never contribute a fact.
    page = matter.load(folder).pages[11]
    answer = {"facts": [{"cite": "DEF0001", "quote": "The examinee reports lifting a tote of about 60 pounds.",
                         "text": "The examinee reports lifting a tote of about 60 pounds."}], "flags": []}
    bad = FakeTransport(FakeModel(recordings=[recording_for(extract.card(), (page,), "9001.001", answer)],
                                  mode="bad_schema"))
    failed_code, failed_line, failed_error = invoke(settings, bad)
    assert failed_code == 1 and failed_error == ""
    assert len(bad.requests) == 2
    stored = facts.load(folder / "casefile.jsonl")
    assert len(stored.facts) == 31, "failed outcome added fact events"
    assert stored.runs[-1]["ok"] is False and stored.runs[-1]["failed"] == ["defense-report/001"]
    assert "new_facts=0 new_flags=0" in failed_line
    output = folder / "outputs/memo.docx"
    modified = output.stat().st_mtime_ns
    facts.append(folder / "casefile.jsonl", [facts.Event("fact", dict(id=f"F{stored.last_seq + 1}",
        cite="DEF0001", page_id=page.page_id, quote=answer["facts"][0]["quote"],
        text=answer["facts"][0]["text"], run_id=stored.runs[-1]["run_id"]))])
    os.utime(folder / "casefile.jsonl", ns=(modified, modified))
    assert status.matter_status(settings, "9001.001")["outputs"] == "memo:stale", "stale output hidden by mtime"
    assert current_ok, "status after memo counts differ"

    monkeypatch.setattr(sys, "argv", ["casefile", "status", "--matter", "9001.001", "--config", str(config_path)])
    assert cli.main() == 0, "status is a view and must exit=0 with open items"
    printed = capsys.readouterr()
    assert printed.err == ""
    assert printed.out.splitlines()[0] == (
        "status 9001.001 pages=14 read=13 failed=1 uncited=0 empty=0 open=1 facts=32 "
        "conflicts=5 gaps=4 outputs=memo:stale")
    assert printed.out.splitlines()[1].startswith("  open: DEF0001 (defense-report/001): failed: gate.schema:")
    monkeypatch.setattr(sys, "argv", ["casefile", "status", "--matter", "9009.001", "--config", str(config_path)])
    assert cli.main() == 2
    missing = capsys.readouterr()
    assert str(settings.paths.matters / "9009.001") in missing.err
    assert not (settings.paths.matters / "9009.001").exists()
    totals["status"] = 6


def test_binder_shown(tmp_path, totals):
    from fakemodel.bench import recording_for

    _, settings, folder = setup_matter(tmp_path)
    metadata = folder / "matter.toml"
    metadata.write_text(metadata.read_text(encoding="utf-8").replace('"medical"', '"ime"'), encoding="utf-8")
    (folder / "records/medical").rename(folder / "records/ime")
    loaded = matter.load(folder)
    answer = json.loads((FIXTURE / "ideal.json").read_text(encoding="utf-8"))
    transport = FakeTransport(FakeModel(recordings=[recording_for(extract.card(), loaded.pages, loaded.number, answer)]))
    code, _, error = invoke(settings, transport)
    assert code == 1 and not error, f"binder_shown=0: {error}"
    paragraphs = Document(folder / "outputs/memo.docx").paragraphs
    heading = next(p for p in paragraphs if p.text == 'Page MED0001 (ime/001), binder ime')
    assert heading.runs[-1].text == 'ime', 'binder_shown=0'
    assert len(transport.requests) == 1
    totals["binder"] = 1


def test_publish_failure_preserves_existing_memo(tmp_path, monkeypatch):
    _, settings, folder = setup_matter(tmp_path)
    transport = FakeTransport(FakeModel(recordings=recordings_for(FIXTURE)))
    assert invoke(settings, transport)[0] == 1
    path = folder / "outputs/memo.docx"
    original = path.read_bytes()
    previous = facts.load(folder / "casefile.jsonl").outputs
    build = memo.build

    def unsafe(*args, **kwargs):
        document = build(*args, **kwargs)
        document.title = "IME summary"
        return document

    monkeypatch.setattr(memo, "build", unsafe)
    code, _, error = invoke(settings, FakeTransport(FakeModel(recordings=[])))
    assert code == 2 and "no-dme-ime" in error
    assert path.read_bytes() == original, "memo published before check"
    assert facts.load(folder / "casefile.jsonl").outputs == previous
    assert not (folder / "outputs/memo.docx.tmp").exists()

    def partial(document, temporary, *, rules):
        temporary.write_bytes(b"invented incomplete document")
        raise memo.MemoError("invented render failure")

    monkeypatch.setattr(memo, "render", partial)
    code, _, error = invoke(settings, FakeTransport(FakeModel(recordings=[])))
    assert code == 2 and "invented render failure" in error
    assert path.read_bytes() == original, "failed render replaced the current memo"
    assert facts.load(folder / "casefile.jsonl").outputs == previous
    assert not (folder / "outputs/memo.docx.tmp").exists(), "failed render left a partial file"


def test_no_unread_pages_calls_no_model_and_exits_zero(tmp_path):
    _, settings, folder = setup_matter(tmp_path)
    loaded = matter.load(folder)
    answer = json.loads((FIXTURE / "ideal.json").read_text(encoding="utf-8"))
    answer["facts"].append(dict(cite="DEF0001", quote="The examinee reports lifting a tote of about 60 pounds.",
                                text="The examinee reports lifting a tote of about 60 pounds.", date="06/20/2025", category="patient_statement", provider="Dr. Theodore Marchetti"))
    transport = FakeTransport(FakeModel(recordings=[recording_for(extract.card(), loaded.pages, "9001.001", answer)]))
    assert invoke(settings, transport)[0] == 0
    before = len(transport.requests)
    code, line, error = invoke(settings, transport)
    assert code == 0 and error == ""
    assert len(transport.requests) == before, "read pages sent again"
    assert "chunks=0 runs=+0 read=14 open=0 facts=32" in line
    assert "new_facts=0 new_flags=0" in line
    assert len(facts.load(folder / "casefile.jsonl").outputs) == 2


def test_later_no_reply_preserves_passing_chunk(tmp_path, monkeypatch):
    _, settings, folder = setup_matter(tmp_path)
    card = split_card(extract.card())
    monkeypatch.setattr(extract, "card", lambda book=None: card)
    loaded = matter.load(folder)
    chunks, _ = pipeline.chunk(card, loaded.pages, settings.model.num_ctx, loaded.number)
    ideal = json.loads((FIXTURE / "ideal.json").read_text(encoding="utf-8"))
    stamps = {stamp for page in chunks[0] for stamp in page.stamps}
    answer = dict(facts=[fact for fact in ideal["facts"] if fact["cite"] in stamps],
                  flags=[flag for flag in ideal["flags"] if set(flag["cites"]) <= stamps])
    transport = FakeTransport(FakeModel(recordings=[recording_for(card, chunks[0], loaded.number, answer)]))
    original_open = transport.open

    def open_request(request, timeout=None):
        if transport.requests:
            raise model.ModelError("invented no reply", received=False)
        return original_open(request, timeout=timeout)

    transport.open = open_request
    code, line, error = invoke(settings, transport)
    assert code == 2 and line == ""
    assert "invented no reply; the pages of this run stay open until a later run reads them" in error
    stored = facts.load(folder / "casefile.jsonl")
    assert len(stored.facts) == len(answer["facts"]), "passing chunk lost on later no reply"
    assert len(stored.runs) == 1 and stored.runs[0]["ok"]
    assert stored.outputs == [] and not (folder / "outputs").exists()
    with open_db(settings.paths.state / "runlog.sqlite") as log:
        assert log.connection.execute("SELECT COUNT(*) FROM runs").fetchone()[0] == 2
        assert log.connection.execute("SELECT COUNT(*) FROM answers").fetchone()[0] == 1


def test_empty_binder_keeps_intake_order():
    loaded = pages.PageSet("9009.001", (matter.Page("zulu", 1, "Invented record.", ("MED0001",)),),
                           (("alpha/001", (), "empty", "no text after OCR"),))
    result = status.snapshot(loaded, facts.CaseFile(), {})
    assert [item.page_id for item in result["open_items"]] == ["alpha/001", "zulu/001"], "empty binder out of order"

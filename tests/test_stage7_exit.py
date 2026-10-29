"""Stage 7 exit proof over invented matters and one socket stand-in."""

from dataclasses import asdict, replace
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
from threading import Thread

from docx import Document
import pytest

from casefile import anon, brief_in, brief_out, cards, config, dates, extract, facts, matterinfo, packet, pages, playbook
from casefile.runlog import open_db
from conftest import ROOT, write_config
from fakemodel import FakeModel, FakeTransport
from fakemodel.bench import recording_for
from fakemodel.server import serve
from test_outputs import setup_case

NUMBER = '9001.001'
FIXTURE = ROOT / 'tests/fixtures/matters' / NUMBER
ANSWERS = ROOT / 'tests/fixtures/anonymizer' / NUMBER
SURNAMES = ('Ashby-Vale', 'Quist-Benning', 'Castellanos-Reyes', 'Ferro-Lindqvist', 'Velloso-Krantz',
            'Okonkwo-Lind', 'Mbeki-Farrow', 'Varga-Ellison', 'Vane-Harrow', 'Marchetti')
EXPECTED_QUOTES = (
    'Date of Service: 03/11/2025', 'Date of Service: 03/14/2025',
    'Date of Service: 03/24/2025', 'Referral from Dr. Ines Castellanos-Reyes dated 03/14/2025.',
    'Next appointment scheduled for 04/07/2025.', 'Date of Service: 04/18/2025',
    'Physical examination was performed today, 05/09/2025.', 'Date of Examination: 06/20/2025')


def _answer(name):
    return json.loads((ANSWERS / (name + '.json')).read_text(encoding='utf-8'))


def _draft(tmp_path, name):
    document = Document()
    for value in (ANSWERS / (name + '.txt')).read_text(encoding='utf-8').splitlines():
        document.add_paragraph(value)
    path = tmp_path / (name + '.docx')
    document.save(path)
    return path


def _private_values(info):
    return (*SURNAMES, info.anon_employer, *info.anon_ids, info.anon_dob, info.client_name,
            info.date_of_accident, info.treating_physician.removeprefix('Dr. '),
            info.iab_hearing_number, info.hearing_date)


def _identifiers(text, info):
    # This oracle is deliberately independent of anon.variants and its date readers.
    text = re.sub(r'\[(?!cite |page )[^\]\n]+\]', '', text)
    patterns = {re.escape(value) for value in _private_values(info) if value}
    for full in info.anon_people:
        words = full.removeprefix('Dr. ').split()
        given, surname = words[0], words[-1]
        patterns.update((re.escape(given), re.escape(surname), re.escape(full),
                         re.escape(surname) + r',\s*' + re.escape(given),
                         re.escape(given) + r'\s+[A-Z]\.?\s+' + re.escape(surname)))
    for organization in (info.anon_employer, *(n for n, k in info.anon_organizations.items() if k in ('employer', 'other'))):
        words = organization.replace('.', '').replace(',', '').split()
        if words and words[-1].casefold() in ('llc', 'inc', 'co', 'corp'):
            words.pop()
        if words:
            patterns.add(r'[ .,]+'.join(re.escape(word) for word in words))
    patterns.update(re.escape(alias) for alias, target in info.anon_aliases.items()
                    if target not in info.anon_keep and info.anon_people.get(target) != 'adjuster'
                    and info.anon_organizations.get(target) not in ('provider', 'insurer'))
    names = sum(len(re.findall(r'(?<!\w)' + pattern + r'(?!\w)', text, re.I)) for pattern in patterns)
    calendar = (r'\b\d{1,4}[/-]\d{1,2}[/-]\d{2,4}\b|'
                r'\b(?:Jan\w*|Feb\w*|Mar\w*|Apr\w*|May|Jun\w*|Jul\w*|Aug\w*|Sep\w*|Oct\w*|Nov\w*|Dec\w*)'
                r'\.?\s+\d{1,2}(?:st|nd|rd|th)?(?:,?\s+\d{2,4})?|'
                r'\b\d{1,2}\s+[A-Za-z]{3,9}\.?\s+\d{4}\b')
    return names + len(re.findall(calendar, text, re.I))


def _read_package(target, *, transform=lambda s: s):
    return transform((target / 'package.md').read_text(encoding='utf-8')), (target / 'package.json').read_text(encoding='utf-8')


@pytest.mark.parametrize('value', ['Morgan', 'Ashby-Vale', 'ASHBY-VALE, MORGAN',
    'Morgan A. Ashby-Vale', 'Morgan A.', 'Quarry Ridge Logistics, LLC', 'Quarry Ridge Logistics',
    'MAV', '3/11/25', 'March 11th, 2025', '11 March 2025', 'Mar. 11 2025'])
def test_identifier_oracle_detects_each_variant(value):
    info = replace(matterinfo.load(FIXTURE), anon_aliases={'MAV': 'Morgan Ashby-Vale'})
    assert _identifiers(value, info) > 0, 'exit oracle missed a planted identifier'


def _summary(result):
    match = re.search(r'checks=(\d+) of (\d+) failed=(\d+) uncited=(\d+) rule_hits=(\d+)', result.stdout)
    assert match, result.stdout + result.stderr
    return tuple(map(int, match.groups()))


def _draft_paragraphs(path):
    document = Document(path)
    values = [p.text for p in document.paragraphs]
    return document, values[values.index('Draft') + 1:values.index('Fact check')]


@pytest.mark.socket
def test_stage7_exit(tmp_path, capsys):
    fake = FakeModel()
    server = serve(port=0, model=fake)
    thread = Thread(target=server.serve_forever)
    thread.start()
    lines = []
    packages = identifiers = answers = drafts_ok = blocked_ok = 0
    try:
        path = write_config(tmp_path, host=f'http://127.0.0.1:{server.server_port}', playbook=ROOT / 'playbook')
        settings = config.load(path)
        book = playbook.load(settings.paths.playbook)
        folders, metadata = {}, {}
        ideal = json.loads((FIXTURE / 'ideal.json').read_text(encoding='utf-8'))
        for number in (NUMBER, '9001.002'):
            folder = settings.paths.matters / number
            shutil.copytree(FIXTURE, folder)
            text = (folder / 'matter.toml').read_text(encoding='utf-8').replace(NUMBER, number)
            if number != NUMBER:
                text = text.replace('"Rosalind Mbeki-Farrow" = "person"\n', '')
            (folder / 'matter.toml').write_text(text, encoding='utf-8')
            metadata[number] = (folder / 'matter.toml').read_bytes()
            folders[number] = folder
            loaded = pages.load(settings.paths.matters, number, settings.ocr.stamp_pattern)
            fake.record(recording_for(extract.card(book), loaded.pages, number, ideal))
        env = dict(os.environ, CASEFILE_CONFIG=str(path))

        def cli(number, *args):
            return subprocess.run(['scripts/run.sh', 'casefile', *args, '--matter', number], cwd=ROOT,
                                  env=env, input='', capture_output=True, text=True, timeout=60)

        for number, folder in folders.items():
            result = cli(number, 'memo')
            assert result.returncode == 1 and not result.stderr, result.stdout + result.stderr
            assert 'facts=38 conflicts=5 gaps=4' in result.stdout
            loaded = pages.load(settings.paths.matters, number, settings.ocr.stamp_pattern)
            info = matterinfo.load(folder)
            cf = facts.load(folder / 'casefile.jsonl')
            value = packet.build(cf, loaded.pages, matter=number, task='brief', options={'issue': 'Review the records'},
                                 playbook_version=book.version)
            value = replace(value, pages=tuple(loaded.pages))
            first = _answer('entities')
            if number != NUMBER:
                first['people'] = [h for h in first['people'] if h['name'] != 'Rosalind Mbeki-Farrow']
            fake.record(recording_for(cards.card('07-ANON'), value.pages, number, first))
            key = anon.Key(anchor='2025-03-11')
            anon.mapping(anon.names(info), [], value.pages, key)
            anonymous = anon.packet(value, key)
            second = _answer('entities-empty' if number == NUMBER else 'entities-second-hit')
            second_pages = anonymous.pages
            if number != NUMBER:
                used = {pid for flag in anonymous.flags if flag['kind'] == 'conflict' for pid in flag['page_ids']}
                used.add('defense-report/001')
                second_pages = tuple(p for p in anonymous.pages if p.page_id in used)
                second['pages_seen'] = [p.page_id for p in second_pages]
            fake.record(recording_for(cards.card('07-ANON'), second_pages, number, second))

        good = cli(NUMBER, 'brief-out', 'Review the records', '--pages', 'all')
        targets = list((folders[NUMBER] / 'outbox').glob('[a-f0-9]' * 12))
        packages = len(targets)
        assert good.returncode == 0 and not good.stderr and packages == 1, (
            f'packages={packages} of 1; ' + good.stdout + good.stderr)
        target = targets[0]
        markdown, serialized = _read_package(target)
        info = matterinfo.load(folders[NUMBER])
        identifiers = _identifiers(markdown + serialized, info)
        assert identifiers == 0, f'identifiers_in_package={identifiers}'
        key = anon.load_key(folders[NUMBER])
        loaded = pages.load(settings.paths.matters, NUMBER, settings.ocr.stamp_pattern)
        value = packet.build(facts.load(folders[NUMBER] / 'casefile.jsonl'), loaded.pages, matter=NUMBER,
                             task='brief', options={'issue': 'Review the records'}, playbook_version=book.version)
        anonymous = anon.packet(replace(value, pages=loaded.pages), key)
        payload = json.loads(serialized)
        assert payload['pages'] == json.loads(anonymous.to_json())['pages']
        assert payload['facts'] == [dict(quote=f.quote, cite=f.cite) for f in anonymous.facts]
        text = '\n'.join(page.text for page in anonymous.pages)
        review = re.search(r'review=(\w+)', good.stdout)[1]
        entities = len(key.entities) + len(anon.names(info).kept)
        anchor = (dates.parse('03/11/2025') - dates.parse(key.anchor)).days
        people = sum(text.count(tag) for tag in key.entities if key.entities[tag]['kind'] in
                     ('claimant', 'doctor', 'attorney', 'person'))
        employer = text.count('[EMPLOYER]')
        ids = text.count('[ID-1]')
        dob = text.count('[DOB]')
        date_count = len(re.findall(r'\[DATE D[+-]\d+\]', text))
        years = len(re.findall(r'\[YEAR Y[+-]\d+\]', text))
        kept = sum(name in text for name in anon.names(info).kept)
        pages_in = markdown.split('## Pages\n\n', 1)[1].count('=== PAGE ')
        mode = (folders[NUMBER] / 'anon/key.json').stat().st_mode & 0o777
        assert (len(anonymous.pages), pages_in, entities, anchor, people, employer, ids, dob,
                date_count, years, kept, mode) == (14, 14, 17, 0, 46, 8, 3, 14, 61, 2, 6, 0o600)
        assert (len(anonymous.facts), sum(f['kind'] == 'conflict' for f in anonymous.flags)) == (38, 5)
        assert target.name == hashlib.sha256((target / 'package.md').read_bytes()).hexdigest()[:12]
        assert 'sha256=' + hashlib.sha256((target / 'package.md').read_bytes()).hexdigest() in good.stdout
        assert target.stat().st_mode & 0o777 == 0o700
        assert set(p.name for p in target.iterdir()) == {'package.md', 'package.json', 'check.txt'}
        assert all(p.stat().st_mode & 0o777 == 0o600 for p in target.iterdir())
        assert f'check=pass leaks=0 entities=17 pages=14 review={review}\n' in (target / 'check.txt').read_text(encoding='utf-8')
        review_words = (target / 'check.txt').read_text(encoding='utf-8').split('Review words: ', 1)[1].strip()
        if review == 'unavailable':
            assert review_words == 'unavailable'
        else:
            assert len(review_words.split(', ')) == int(review) and all(
                word in markdown for word in review_words.split(', ')), 'review words missing'
        assert review == ('24' if Path('/usr/share/dict/words').is_file() else 'unavailable'), f'review={review}'
        assert 'open_items=1' in good.stdout
        lines.append(f'stage7 package matter={NUMBER} pages=14 pages_in_package=14 entities=17 of 17 anchor=D+0 '
                     f'people=46 employer=8 ids=3 dob=14 dates=61 years=2 kept_orgs=6 leaks=0 review={review} '
                     'key_mode=600 facts=38 conflicts=5 exit=0')

        # A new, uncited page exercises the default packet boundary without moving the 14-page baselines.
        extra = folders[NUMBER] / 'records/medical/009.txt'
        extra.write_bytes((ANSWERS.parent / 'variant-page.txt').read_bytes())
        full = pages.load(settings.paths.matters, NUMBER, settings.ocr.stamp_pattern)
        first = _answer('entities')
        first['pages_seen'] = [p.page_id for p in full.pages]
        fake.record(recording_for(cards.card('07-ANON'), full.pages, NUMBER, first))
        listed = anon.names(info)
        anon.mapping(listed, [], full.pages, key)
        converted = anon.packet(replace(anonymous, pages=full.pages), key)
        used = {pid for flag in anonymous.flags if flag['kind'] == 'conflict' for pid in flag['page_ids']}
        used.update(('defense-report/001', 'medical/009'))
        selected = tuple(p for p in converted.pages if p.page_id in used)
        for part in (converted.pages, selected):
            second = _answer('entities-empty')
            second['pages_seen'] = [p.page_id for p in part]
            fake.record(recording_for(cards.card('07-ANON'), part, NUMBER, second))
        default = cli(NUMBER, 'brief-out', 'Review selected records')
        assert default.returncode == 0 and not default.stderr, default.stdout + default.stderr
        default_id = re.search(r'packet=([a-f0-9]+)', default.stdout)[1]
        default_md, default_json = _read_package(folders[NUMBER] / 'outbox' / default_id)
        default_packet = json.loads(default_json)
        assert default_packet['pages'] == json.loads(json.dumps([asdict(p) for p in selected])), 'M1: default JSON carries unused pages'
        assert all(p['text'] in default_md for p in default_packet['pages']), 'M1: JSON page not displayed in Markdown'
        assert 'open_items=2' in default.stdout
        assert '## Pages\n' not in default_md
        assert _identifiers(default_md + default_json, info) == 0
        # The new page itself must round-trip and pass the independent leak oracle.
        variant = converted.pages[-1] if converted.pages[-1].stamps == ('VAR0001',) else next(p for p in converted.pages if p.stamps == ('VAR0001',))
        source = next(p for p in full.pages if p.stamps == ('VAR0001',))
        assert _identifiers(variant.text, info) == 0, 'H1/M3: variant fixture leaks'
        assert anon.restore_quote(variant.text, key, source) == source.text
        assert variant.text.count('\n') == source.text.count('\n')
        identifiers += _identifiers(default_md + default_json + variant.text, info)
        packages += 1
        lines.append(f'stage7 default pages={len(full.pages)} pages_in_package={len(default_packet["pages"])} identifiers_in_package=0 exit=0')

        blocked = cli('9001.002', 'brief-out', 'Review the records')
        folder = folders['9001.002']
        written = int((folder / 'outbox').exists())
        key_written = int((folder / 'anon/key.json').is_file())
        key_leak = _identifiers(blocked.stdout + blocked.stderr, matterinfo.load(folder))
        assert key_leak == 0, f'key_leak={key_leak}'
        assert written == 0, f'package_written={written}'
        assert blocked.returncode == 1 and key_written == 1, blocked.stdout + blocked.stderr
        assert 'matter.toml [anon.people]' in blocked.stdout + blocked.stderr and '[anon.aliases]' in blocked.stdout + blocked.stderr, 'L6: block message lacks name-list remedy'
        assert re.fullmatch(r'brief-out: blocked at step 4 \(person, offset \d+\); correct matter.toml \[anon.people\], '
                           r'\[anon.organizations\] or \[anon.aliases\] and rerun; nothing was written\n', blocked.stdout + blocked.stderr)
        assert blocked.stderr == ''
        assert (folder / 'anon/key.json').stat().st_mode & 0o777 == 0o600
        blocked_ok = 1
        lines.append('stage7 blocked matter=9001.002 step=4 kind=person package_written=0 key_written=1 exit=1')

        draft_path = _draft(tmp_path, 'draft')
        result = cli(NUMBER, 'brief-in', str(draft_path))
        assert result.returncode == 0 and not result.stderr, result.stdout + result.stderr
        checks, total, failed, uncited, rule_hits = _summary(result)
        assert (checks, total, failed, uncited, rule_hits) == (44, 44, 0, 2, 0), f'checks={checks} of {total} failed={failed}'
        document, paragraphs = _draft_paragraphs(folders[NUMBER] / 'outputs/brief-in-draft.docx')
        assert len(paragraphs) == 12
        assert [re.search(r'"([^\"]*)"', p)[1] for p in paragraphs[:8]] == list(EXPECTED_QUOTES)
        assert paragraphs[8:] == ['At 03/14/2025, pain was 4/10 [cite MED0003].',
            'The job required lifting up to 50 pounds [page employer/002].',
            'Morgan Ashby-Vale asks that the records be considered together.',
            'The competing accounts warrant careful review.']
        original = [p.text for p in Document(draft_path).paragraphs]
        placeholders = re.findall(r'\[(?!cite |page )[^\]]+\]', '\n'.join(original))
        restored = len(placeholders) - len(re.findall(r'\[(?!cite |page )[^\]]+\]', '\n'.join(paragraphs)))
        assert restored == 11, f'placeholders_restored={restored}'
        assert 'Checks: 44 of 44; failures: 0; uncited: 2; rule hits: 0; skipped: 0' in [p.text for p in document.paragraphs]
        drafts_ok += 1
        lines.append(f'stage7 draft=draft paragraphs=12 checks=44 of 44 uncited=2 placeholders_restored={restored} rule_hits=0 exit=0')

        bad_path = _draft(tmp_path, 'draft-bad-quote')
        result = cli(NUMBER, 'brief-in', str(bad_path))
        assert result.returncode == 1 and not result.stderr, result.stdout + result.stderr
        assert _summary(result) == (43, 44, 1, 2, 0)
        partial = folders[NUMBER] / 'outputs/brief-in-draft-bad-quote.partial.docx'
        document, paragraphs = _draft_paragraphs(partial)
        marked = sum(p.startswith('[FACT CHECK FAILED: paragraph 1: quote not on MED0001;') for p in paragraphs)
        assert marked == 1 and len(paragraphs) == 13
        assert not partial.with_name('brief-in-draft-bad-quote.docx').exists()
        drafts_ok += 1
        lines.append('stage7 draft=draft-bad-quote checks=43 of 44 failed=1 marked=1 partial=1 exit=1')
        result = cli(NUMBER, 'status')
        assert result.returncode == 0 and not result.stderr
        output_status = result.stdout.splitlines()[0].split(' outputs=', 1)[1]
        assert output_status == ('memo:current,brief-out:current,brief-in-draft:current;'
                                 'partial=brief-in-draft-bad-quote.partial.docx'), output_status
        lines.append('stage7 status outputs=' + output_status)
        with open_db(settings.paths.state / 'runlog.sqlite') as log:
            answers = log.connection.execute('SELECT COUNT(*) FROM answers').fetchone()[0]
            exported = '\n'.join(row[0] for table in ('failures', 'gates')
                                 for row in log.connection.execute(f'SELECT message FROM {table}'))
        exported += ''.join(p.read_text(encoding='utf-8') for p in target.iterdir())
        assert _identifiers(exported + good.stdout + good.stderr, info) == 0, 'key_leak=1'
        assert answers == 8 and sum(r['path'] == '/api/chat' for r in fake.requests) == answers
        assert all((folder / 'matter.toml').read_bytes() == metadata[number] for number, folder in folders.items())
    finally:
        server.shutdown()
        thread.join(timeout=5)
        stopped = not thread.is_alive()
        server.server_close()
    failed = bool(packages != 2 or blocked_ok != 1 or drafts_ok != 2 or answers != 8 or identifiers or not stopped)
    lines.append(f'stage7 exit={int(failed)} packages={packages} of 2 blocked_seeded={blocked_ok} of 1 '
                 f'drafts={drafts_ok} of 2 answers={answers:+d} identifiers_in_package={identifiers} server_stopped={int(stopped)}')
    with capsys.disabled():
        print('\n'.join(lines), flush=True)
    assert not failed, '\n'.join(lines)


@pytest.mark.parametrize('bad', ['[DR-999]', '[DR 3]', '[Dr-3]', '[DATE]', '[DR-n]', '[DATE D+99999]',
                               '[CLAIMANT', 'Ashby-Vale', 'March 11, 2025'])
def test_draft_blocked_before_publication(tmp_path, bad):
    _, cfg, folder, loaded, cf, _ = setup_case(tmp_path)
    key = anon.Key(anchor='2025-03-11')
    anon.mapping(anon.names(matterinfo.load(folder)), [], loaded.pages, key)
    anon.save_key(folder, key)
    document = Document()
    document.add_paragraph(bad)
    path = tmp_path / 'blocked.docx'
    document.save(path)
    from io import StringIO
    out, err = StringIO(), StringIO()
    before = (folder / 'casefile.jsonl').read_bytes()
    code = brief_in.run(cfg, NUMBER, path, out=out, err=err)
    assert code == 2 and not out.getvalue() and 'blocked at step' in err.getvalue()
    assert bad not in err.getvalue(), 'blocked draft value escaped'
    assert (folder / 'casefile.jsonl').read_bytes() == before
    assert not (folder / 'outputs').exists()


def test_draft_unknown_cite_is_partial(tmp_path):
    from io import StringIO

    _, cfg, folder, loaded, _, _ = setup_case(tmp_path)
    key = anon.Key(anchor='2025-03-11')
    anon.mapping(anon.names(matterinfo.load(folder)), [], loaded.pages, key)
    anon.save_key(folder, key)
    document = Document()
    document.add_paragraph('Pain 4/10 [cite MED9999].')
    path = tmp_path / 'unknown-cite.docx'
    document.save(path)
    out, err = StringIO(), StringIO()
    assert brief_in.run(cfg, NUMBER, path, out=out, err=err) == 1
    assert not err.getvalue()
    partial = folder / 'outputs/brief-in-unknown-cite.partial.docx'
    assert any('unknown or ambiguous cite MED9999' in p.text for p in Document(partial).paragraphs)
    assert not facts.load(folder / 'casefile.jsonl').outputs


@pytest.mark.parametrize('missing', ['name', 'anchor'])
def test_out_prerequisite_exit_two(tmp_path, missing):
    from io import StringIO

    _, cfg, folder, loaded, _, book = setup_case(tmp_path)
    info = matterinfo.load(folder)
    first = _answer('entities')
    if missing == 'name':
        people = dict(info.anon_people)
        people.pop('Rosalind Mbeki-Farrow')
        info = replace(info, anon_people=people)
    else:
        info = replace(info, date_of_accident='')
        first['anchor_dates'] = []
    matterinfo.write(folder, info)
    before = (folder / 'matter.toml').read_bytes()
    opener = FakeTransport(FakeModel(recordings=[recording_for(cards.card('07-ANON'), loaded.pages, NUMBER, first)]))
    out, err = StringIO(), StringIO()
    code = brief_out.run(cfg, NUMBER, 'Review the records', out=out, err=err, opener=opener, stdin=StringIO())
    assert code == 2, f'prerequisite exit={code}'
    assert ('NAME LIST NEEDED' if missing == 'name' else 'ANCHOR NEEDED') in out.getvalue() + err.getvalue()
    assert _identifiers(out.getvalue() + err.getvalue(), info) == 0, 'key_leak=1'
    assert len(opener.requests) == 1
    assert (folder / 'matter.toml').read_bytes() == before
    assert not (folder / 'anon').exists() and not (folder / 'outbox').exists()


@pytest.mark.parametrize('row', ['pages', 'facts', 'rules', 'open-items', 'review'])
def test_fix2_package_surfaces(tmp_path, monkeypatch, row):
    from io import StringIO

    _, cfg, folder, loaded, cf, book = setup_case(tmp_path)
    from casefile.runlog import open_db
    with open_db(cfg.paths.state / 'runlog.sqlite') as log:
        log.record_coverage(1, matter=NUMBER, page_id='defense-report/001', card='medical_memo', status='uncited', reason='uncited')
    value = packet.build(cf, loaded.pages, matter=NUMBER, task='brief', options={'issue': 'Review'}, playbook_version=book.version)
    key = anon.mapping(anon.names(matterinfo.load(folder)), [], loaded.pages, anon.Key(anchor='2025-03-11'))
    anonymous = anon.packet(replace(value, pages=loaded.pages), key)
    expected = {'medical/001', 'medical/003', 'medical/006', 'medical/007', 'medical/008',
                'defense-report/001', 'defense-report/002', 'defense-report/003', 'employer/001'}
    selected = tuple(p for p in anonymous.pages if p.page_id in expected)
    fake = FakeModel()
    fake.record(recording_for(cards.card('07-ANON'), loaded.pages, NUMBER, _answer('entities')))
    old_ids = {p.page_id for p in value.pages} | {pid for f in value.flags for pid in f['page_ids']}
    old_selected = tuple(p for p in anonymous.pages if p.page_id in old_ids)
    for part in (anonymous.pages, old_selected, selected):
        second = _answer('entities-empty')
        second['pages_seen'] = [p.page_id for p in part]
        fake.record(recording_for(cards.card('07-ANON'), part, NUMBER, second))
    out, err = StringIO(), StringIO()
    assert brief_out.run(cfg, NUMBER, 'Review', out=out, err=err, opener=FakeTransport(fake), stdin=StringIO()) == 0, out.getvalue() + err.getvalue()
    target = next((folder / 'outbox').iterdir())
    md, serialized = _read_package(target)
    payload = json.loads(serialized)
    if row == 'pages':
        assert {p['binder'] + '/' + str(p['number']).zfill(3) for p in payload['pages']} == expected, 'M1: JSON includes pages whose full text is not displayed'
        assert all(p['text'] in md for p in payload['pages']), 'M1: JSON page text absent from Markdown'
    elif row == 'facts':
        assert all(set(f) == {'quote', 'cite'} for f in payload['facts']), 'M1: package facts carry more than quote and cite'
        assert len(payload['facts']) == 38
    elif row == 'rules':
        rules_text = md.split('## Facts', 1)[0]
        for rule in book.rules.rules:
            assert rule.id in rules_text and rule.rule in rules_text, 'M2: package omits a rule id or its source text'
        assert 'header-block' in rules_text and 'No em or en dashes' in rules_text
        assert 'footnote' in rules_text and 'straight double quotes' in rules_text
    elif row == 'open-items':
        assert '## Pages not yet read\n' in md, 'M4: package omits unread-page section'
        section = md.split('## Pages not yet read\n', 1)[1]
        assert 'DEF0001' in section and next(p.text for p in anonymous.pages if p.page_id == 'defense-report/001') in section
        assert 'open_items=1' in out.getvalue(), 'M4: status omits open-item count'
    else:
        # An empty dictionary makes exclusion of template vocabulary observable on every machine.
        original = Path.read_text
        def read_text(path, *args, **kwargs):
            return '' if path == Path('/usr/share/dict/words') else original(path, *args, **kwargs)
        monkeypatch.setattr(Path, 'read_text', read_text)
        original_is_file = Path.is_file
        monkeypatch.setattr(Path, 'is_file', lambda p: True if p == Path('/usr/share/dict/words') else original_is_file(p))
        words = brief_out._review_words(md, anon.names(matterinfo.load(folder)), anonymous.pages)
        assert not {'Conflicts', 'Facts', 'Gaps', 'Pages', 'Dr', 'Drafter', 'Matter', 'Issue'} & set(words.split(', ')), 'L6: template vocabulary fills review list'


def test_fix2_page_aware_name_restore(tmp_path):
    from io import StringIO

    _, cfg, folder, loaded, _, _ = setup_case(tmp_path)
    key = anon.mapping(anon.names(matterinfo.load(folder)), [], loaded.pages, anon.Key(anchor='2025-03-11'))
    anon.save_key(folder, key)
    doc = Document()
    doc.add_paragraph('The report identifies "Mr. [CLAIMANT]" [cite DEF0001].')
    path = tmp_path / 'source-name.docx'
    doc.save(path)
    out, err = StringIO(), StringIO()
    assert brief_in.run(cfg, NUMBER, path, out=out, err=err) == 0, 'W5c: quote restoration lost the page spelling'
    assert any('"Mr. Ashby-Vale"' in p.text for p in Document(folder / 'outputs/brief-in-source-name.docx').paragraphs)

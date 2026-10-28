"""Stage 6 exit proof through six CLI jobs and their reopened Word outputs."""

import json
import os
import re
import shutil
import subprocess
from threading import Thread

from docx import Document
import pytest

from casefile import cards, config, depo, facts, hearing, matterinfo, outputs, packet, pages, playbook, quotes, report, status
from casefile.runlog import open_db
from conftest import ROOT, write_config
from fakemodel import FakeModel
from fakemodel.bench import recording_for, recordings_for
from fakemodel.server import serve
from test_stage5_exit import _read_memo

NUMBER = '9001.001'
MATTER = ROOT / 'tests/fixtures/matters' / NUMBER
ANSWERS = ROOT / 'tests/fixtures/outputs' / NUMBER
EXPECTED = (
    'stage6 memo sections=15 chronology=31 dated=31 of 31 providers=6 conflicts=5 conflict_quotes_ok=1 defense_first=1 header_ok=1 rule_hits=0 exit=1',
    'stage6 depo=defense-doctor-us outline_our=25 outline_defense=7 questions=8 rests_on_page=8 of 8 danger=1 points=1 forbidden_topics=0 rule_hits=0 header_ok=1 exit=1',
    'stage6 depo=our-doctor-us outline_our=25 outline_defense=7 questions=4 rests_on_page=4 of 4 danger=1 points=0 treatment_q=1 bills_q=1 compound=0 providers_named=6 of 6 rule_hits=0 header_ok=1 exit=1',
    'stage6 hearing checklist_first=1 starred=1 issues=2 standard_needed=2 client_direct=2 client_cross=2 elements=14 not_found=2 closing=2 relief_needed=2 rule_hits=0 header_ok=1 exit=1',
    'stage6 report=defense-report concessions=1 labels=1 danger=1 conflicts=3 conflict_quotes_ok=1 audit_limit=1 rule_hits=0 header_ok=1 exit=1',
    'stage6 report=okonkwo-lind elements=14 stated=10 partial=2 not_found=2 gaps_separate=1 rule_hits=0 header_ok=1 exit=1',
    'stage6 status outputs=memo:current,depo-defense-doctor-us:current,depo-our-doctor-us:current,hearing:current,report-defense-report:current,report-okonkwo-lind:current partial=0',
    'stage6 exit=0 docs=6 of 6 answers=+16 rule_hits=0 server_stopped=1',
)
HEADER = ['Attorney Work Product | Privileged & Confidential', 'Matter type: workers-comp',
          'Client: Morgan Ashby-Vale', 'Date of accident: 03/11/2025',
          'Treating physician: Dr. Priya Okonkwo-Lind', 'IAB hearing number: 1234567',
          'Hearing date: 01/15/2026 9:00 AM']
ELEMENTS = ('treating_relationship', 'mechanism', 'history_prior_conditions', 'exam_findings',
            'imaging_reviewed', 'diagnosis', 'causation', 'treatment_reasonable_necessary_related',
            'bills_reasonable_related', 'work_status_restrictions', 'mmi', 'permanency_rating',
            'future_treatment', 'records_reviewed')


def _answer(name):
    return json.loads((ANSWERS / (name + '.json')).read_text(encoding='utf-8'))


def _section(doc, title):
    paragraphs = doc.paragraphs
    start = next((i + 1 for i, p in enumerate(paragraphs)
                  if p.style.name == 'Heading 1' and p.text == title), len(paragraphs))
    end = next((i for i in range(start, len(paragraphs))
                if paragraphs[i].style.name == 'Heading 1'), len(paragraphs))
    return paragraphs[start:end]


def _texts(paragraphs):
    return [p.text for p in paragraphs]


def _header_ok(doc, cf, book, model, *, who=''):
    expected = [*HEADER]
    if who:
        expected.append('Deposition of Dr. ' + who + ', NEEDED')
    expected.append('Matter ' + NUMBER)
    actual = _texts(doc.paragraphs)
    index = 1 + len(expected) + int(bool(who))
    return int(actual[1:1 + len(expected)] == expected and len(actual) > index and
        (not who or actual[index - 1] == 'Purpose: discovery') and
        re.fullmatch(r'Built \d{4}-\d{2}-\d{2} from case file entry ' + str(cf.last_seq) +
                     re.escape(f'; playbook {book.version}; model {model}'), actual[index]) is not None)


def _saved_rules(doc, rules, providers, source_pages, *, our_doctor=False):
    """Check saved Word text against source pages without renderer exemption metadata."""
    hits = 0
    for paragraph in doc.paragraphs:
        value = paragraph.text
        cited = []
        def citation(match):
            label = match[1]
            matches = [page for page in source_pages if label == page.page_id or label in page.stamps]
            if len(matches) == 1:
                cited.append(matches[0])
                return '\n'
            return match[0]
        value = re.sub(r'\[(?:cite|page) ([^\]]+)\]', citation, value)
        quoted = re.search(r'"(.*)"', value)
        if quoted and any(quotes.match(quoted[1], page.text) is not None for page in cited):
            value = value[:quoted.start()] + '\n' + value[quoted.end():]
        question = re.match(r'^\d+\. ', paragraph.text) is not None
        hits += sum(hit.severity == 'error' for hit in rules.check(
            value, unit='question' if question else 'text', providers=providers,
            bills_question=our_doctor and question))
    return hits


def _rules(doc, expected, rules, providers, errors, source_pages=(), *, our_doctor=False):
    # Structure is a separate check; the safety oracle uses saved text and source pages.
    templates = [(outputs.p(expected.title), 'Title')]
    templates.extend((p, None) for p in expected.header)
    for section in expected.sections:
        templates.append((outputs.p(section.title), 'Heading 1'))
        templates.extend((p, None) for p in section.paragraphs or [outputs.p('none')])
    if len(doc.paragraphs) != len(templates):
        errors.append('reopened document paragraph count differs from checked runs')
        return 0
    hits = _saved_rules(doc, rules, providers, source_pages, our_doctor=our_doctor)
    for actual, (template, style) in zip(doc.paragraphs, templates):
        actual_runs = [r.text for r in actual.runs]
        expected_runs = [r.text for r in template.runs]
        if template.page_break:
            actual_runs = actual_runs[:-1] if actual_runs and actual_runs[-1] == '' else actual_runs
        if actual_runs != expected_runs or style and actual.style.name != style:
            errors.append('reopened document runs differ: ' + actual.text)
            continue
    return hits


def _open_output(path, errors):
    partial = path.with_name(path.stem + '.partial.docx')
    if partial.exists():
        doc = Document(partial)
        errors.extend(p.text for p in doc.paragraphs if "the model's answer failed" in p.text)
        return doc
    return Document(path) if path.exists() else Document()


def _questions(doc, title):
    return [p for p in _section(doc, title) if re.match(r'^\d+\. ', p.text)]


def _rests(doc, loaded):
    checked = 0
    for paragraph in doc.paragraphs:
        if not paragraph.text.startswith('Rests on: '):
            continue
        match = re.fullmatch(r'Rests on: "(.*)" \[cite ([A-Z]+\d+)\]', paragraph.text)
        if match:
            quote, cite = match.groups()
            pages = [p for p in loaded.pages if cite in p.stamps]
            checked += int(len(pages) == 1 and quote in pages[0].text)
    return checked


def _outline(doc, title):
    return sum(any(r.text.startswith('"') for r in p.runs) for p in _section(doc, title))


def _conflicts_ok(paragraphs, cf):
    actual = _texts(paragraphs)
    blocks = []
    for i, text in enumerate(actual):
        match = re.fullmatch(r'From \[cite ([A-Z]+\d+)\]:', text)
        if match:
            end = next((j for j in range(i + 1, len(actual)) if not actual[j].startswith('"')), len(actual))
            expected = [f'"{fact.quote}" [cite {match[1]}]' for fact in cf.facts if match[1] in fact.cites]
            blocks.append(actual[i + 1:end] == expected)
    return int(bool(blocks) and all(blocks))


def _depo_line(doc, who, code, cf, loaded, book, header_ok, rule_hits, errors):
    data = {part: _answer(f'questions-{who}-{part}') for part in ('primary', 'cross', 'redirect')}
    questions = [p for title in ('Questions: primary', 'Questions: anticipated cross', 'Questions: redirect')
                 for p in _questions(doc, title)]
    total = sum(len(answer['questions']) for answer in data.values())
    rests = _rests(doc, loaded)
    for part, title in zip(data, ('Questions: primary', 'Questions: anticipated cross', 'Questions: redirect')):
        actual = _texts(_questions(doc, title))
        expected = [f'{n}. {item["text"]}' for n, item in enumerate(data[part]['questions'], 1)]
        if actual != expected:
            errors.append(f'{who} {part}: question texts differ')
    breaks = sum(len(r._r.xpath('.//w:br[@w:type="page"]')) for p in doc.paragraphs for r in p.runs)
    if breaks != 2:
        errors.append(f'{who}: page_breaks={breaks}, expected 2')
    providers = outputs.providers(cf)
    findings = [hit for p in questions for hit in book.rules.check(p.text, unit='question', providers=providers)
                if hit.severity == 'error']
    danger = _outline(doc, 'Danger zone')
    points = _outline(doc, 'Points for our expert')
    line = (f'stage6 depo={who}-us outline_our={_outline(doc, "Case outline: our view")} '
        f'outline_defense={_outline(doc, "Case outline: the defense view")} questions={len(questions)} '
        f'rests_on_page={rests} of {total} danger={danger} points={points} ')
    if who == 'our-doctor':
        primary = _questions(doc, 'Questions: primary')
        treatment = sum(bool(re.search(r'\btreatment\b', p.text, re.I)) for p in primary)
        bills = [p for p in primary if re.search(r'\bbills?\b', p.text, re.I)]
        named = sum(any(provider in p.text for p in bills) for provider in providers)
        compound = sum(hit.rule_id == 'treatment-and-bills-apart' for hit in findings)
        line += f'treatment_q={treatment} bills_q={len(bills)} compound={compound} providers_named={named} of {len(providers)} '
    else:
        line += f'forbidden_topics={len(findings)} '
    return line + f'rule_hits={rule_hits} header_ok={header_ok} exit={code}'


def _hearing_line(doc, code, header_ok, rule_hits):
    titles = [p.text for p in doc.paragraphs if p.style.name == 'Heading 1']
    checklist = _texts(_section(doc, 'Pre-hearing checklist'))
    outline = _texts(_section(doc, 'Hearing outline'))
    client = _texts(_section(doc, 'Client preparation'))
    doctor = _texts(_section(doc, 'Doctor preparation'))
    closing = _texts(_section(doc, 'Closing outline'))
    direct = client[client.index('Direct examination') + 1:client.index('Likely cross')] if 'Likely cross' in client else []
    cross = client[client.index('Likely cross') + 1:client.index('Consistency flags')] if 'Consistency flags' in client else []
    elements = doctor[doctor.index('Opinions the case needs') + 1:doctor.index('Gaps')] if 'Gaps' in doctor else []
    return (f'stage6 hearing checklist_first={int(bool(titles) and titles[0] == "Pre-hearing checklist")} '
        f'starred={int(bool(checklist) and checklist[0] == hearing.STARRED)} '
        f'issues={sum(p.startswith("Issue: ") for p in outline)} standard_needed={outline.count("Standard: STANDARD NEEDED")} '
        f'client_direct={sum(bool(re.match(r"^\d+\. ", p)) for p in direct)} '
        f'client_cross={sum(bool(re.match(r"^\d+\. ", p)) for p in cross)} '
        f'elements={sum(any(p.startswith(e + ":") for e in ELEMENTS) for p in elements)} '
        f'not_found={sum(": NOT FOUND" in p for p in elements)} '
        f'closing={sum(p.startswith("Issue: ") for p in closing)} relief_needed={closing.count("Relief requested: RELIEF NEEDED")} '
        f'rule_hits={rule_hits} header_ok={header_ok} exit={code}')


def _report_line(doc, defense, code, cf, header_ok, rule_hits):
    if defense:
        conflicts = _section(doc, 'Conflicts with other records')
        topics = _texts(_section(doc, 'Candidate cross topics (document-based)'))
        limit = ("Records review audit: not built this round; the report's records-referenced statements "
                 'are listed under What the report states for your own comparison.')
        return (f'stage6 report=defense-report concessions={_outline(doc, "Concessions (the report\'s own words)")} '
            f'labels={_outline(doc, "Undefined labels")} danger={_outline(doc, "Danger zone")} '
            f'conflicts={sum(p.text.startswith("Conflict between ") for p in conflicts)} '
            f'conflict_quotes_ok={_conflicts_ok(conflicts, cf)} audit_limit={int(limit in topics)} '
            f'rule_hits={rule_hits} header_ok={header_ok} exit={code}')
    elements = _texts(_section(doc, 'Elements'))
    gaps = _texts(_section(doc, 'Gaps'))
    actual = [p.split(':', 1)[0] for p in elements if any(p.startswith(e + ':') for e in ELEMENTS)]
    separate = int(actual == list(ELEMENTS) and all(any(p.startswith(e + ': NOT FOUND') for p in gaps)
                   for e in ('treatment_reasonable_necessary_related', 'bills_reasonable_related')))
    return (f'stage6 report=okonkwo-lind elements={len(actual)} stated={sum(": STATED:" in p for p in elements)} '
        f'partial={sum(": PARTIAL:" in p for p in elements)} not_found={sum(": NOT FOUND" in p for p in elements)} '
        f'gaps_separate={separate} rule_hits={rule_hits} header_ok={header_ok} exit={code}')


def _register_depo(fake, value, book, settings, who):
    data = {'factsheet': _answer('factsheet-' + who)}
    card = cards.card('02-A', book)
    for part in packet.batches(value, card, settings.model.num_ctx):
        fake.record(recording_for(card, part.pages, NUMBER, data['factsheet'], body=part.body()))
    card = cards.card(depo.CARD_IDS[who], book)
    for section in ('primary', 'cross', 'redirect'):
        key = 'questions:' + section
        data[key] = _answer(f'questions-{who}-{section}')
        part = depo.question_packet(value, data['factsheet'], section, card, settings.model.num_ctx)
        fake.record(recording_for(card, part.pages, NUMBER, data[key], body=part.body()))
    return data


@pytest.mark.socket
def test_stage6_exit(tmp_path, capsys):
    fake = FakeModel(recordings=recordings_for(MATTER))
    server = serve(port=0, model=fake)
    thread = Thread(target=server.serve_forever)
    thread.start()
    errors, lines, rule_hits = [], [], 0
    try:
        path = write_config(tmp_path, host=f'http://127.0.0.1:{server.server_port}', playbook=ROOT / 'playbook')
        settings = config.load(path)
        folder = settings.paths.matters / NUMBER
        shutil.copytree(MATTER, folder)
        metadata = (folder / 'matter.toml').read_bytes()
        loaded = pages.load(settings.paths.matters, NUMBER, settings.ocr.stamp_pattern)
        info = matterinfo.load(folder)
        book = playbook.load(settings.paths.playbook)
        env = dict(os.environ, CASEFILE_CONFIG=str(path))

        def cli(*args):
            return subprocess.run(['scripts/run.sh', 'casefile', *args, '--matter', NUMBER],
                cwd=ROOT, env=env, input='', capture_output=True, text=True, timeout=60)

        result = cli('memo')
        cf = facts.load(folder / 'casefile.jsonl')
        with open_db(settings.paths.state / 'runlog.sqlite') as log:
            coverage = log.coverage(NUMBER, job='memo')
        memo_line = _read_memo(1, folder / 'outputs/memo.docx', cf, loaded.pages, coverage, book, errors)
        metrics = memo_line.split(' sections=', 1)[1].split(' cites=', 1)[0]
        memo_hits = _saved_rules(Document(folder / 'outputs/memo.docx'), book.rules, outputs.providers(cf), loaded.pages)
        rule_hits += memo_hits
        lines.append('stage6 memo sections=' + metrics + f' rule_hits={memo_hits} exit=' + str(result.returncode))
        if result.stderr or 'chunks=1 runs=+1 read=13 open=1 facts=31 conflicts=5 gaps=4' not in result.stdout:
            errors.append('memo CLI differs: ' + result.stdout + result.stderr)
        if result.returncode != 1:
            errors.append(f'memo exit={result.returncode}, expected 1')
        counted = status.snapshot(loaded, cf, coverage, book)
        for who, provider, name in (('defense-doctor', 'Marchetti', 'Marchetti'),
                                     ('our-doctor', 'Okonkwo-Lind', 'Okonkwo-Lind')):
            cf = facts.load(folder / 'casefile.jsonl')
            value = depo.prepare(cf, loaded.pages, book, info, number=NUMBER,
                                 deponent=who, taken_by='us', provider=provider)
            if who == 'defense-doctor' and set(value.options['deponent_pages']) != {
                    'defense-report/001', 'defense-report/002', 'defense-report/003'}:
                errors.append('defense provider selection did not widen to the report binder')
            if who == 'our-doctor' and value.options['deponent_pages'] != ['medical/007', 'medical/008']:
                errors.append('our provider selection widened beyond the matching facts')
            data = _register_depo(fake, value, book, settings, who)
            result = cli('depo', '--deponent', who, '--taken-by', 'us', '--provider', provider, '--name', 'Dr. ' + name)
            value.options.update(name='Dr. ' + name, when='')
            expected = depo.render(data, value, cf, loaded.pages, info, open_items=counted['open_items'],
                                   playbook_version=book.version, model=settings.model.name)
            output = folder / f'outputs/depo-{who}-us.docx'
            doc = _open_output(output, errors)
            hits = _rules(doc, expected, book.rules, outputs.providers(cf), errors, loaded.pages,
                          our_doctor=who == 'our-doctor')
            rule_hits += hits
            lines.append(_depo_line(doc, who, result.returncode, cf, loaded, book,
                _header_ok(doc, cf, book, settings.model.name, who=name), hits, errors))
            if result.stderr:
                errors.append(f'{who}: ' + result.stderr.strip())

        cf = facts.load(folder / 'casefile.jsonl')
        values = hearing.prepare(cf, loaded.pages, book, info, number=NUMBER, provider='Okonkwo-Lind')
        data = {}
        for key, value in values.items():
            if key == 'base':
                continue
            data[key] = _answer('factsheet-issue-' + key.rsplit(':', 1)[1] if key.startswith('factsheet:') else
                'client-' + key.split(':')[1] if key.startswith('questions:') else 'elements-okonkwo-lind')
            card = hearing.step_card(key, book)
            parts = [packet.fit(value, card, settings.model.num_ctx)] if key.startswith('questions:') else packet.batches(value, card, settings.model.num_ctx)
            for part in parts:
                fake.record(recording_for(card, part.pages, NUMBER, data[key], body=part.body()))
        result = cli('hearing', '--provider', 'Okonkwo-Lind')
        expected = hearing.render(data, values, cf, loaded.pages, info, open_items=counted['open_items'],
                                  playbook_version=book.version, model=settings.model.name)
        output = folder / 'outputs/hearing.docx'
        doc = _open_output(output, errors)
        hits = _rules(doc, expected, book.rules, outputs.providers(cf), errors, loaded.pages)
        rule_hits += hits
        lines.append(_hearing_line(doc, result.returncode, _header_ok(doc, cf, book, settings.model.name), hits))
        if result.stderr:
            errors.append('hearing: ' + result.stderr.strip())
        for defense in (True, False):
            cf = facts.load(folder / 'casefile.jsonl')
            kwargs = dict(binder=('defense-report',), defense=True) if defense else dict(provider='Okonkwo-Lind')
            value = report.prepare(cf, loaded.pages, book, info=info, number=NUMBER, **kwargs)
            key, card_id, filename = ('concessions', '04-G', 'concessions-defense-report') if defense else ('elements', '05-A', 'elements-okonkwo-lind')
            data = {key: _answer(filename)}
            card = cards.card(card_id, book)
            for part in packet.batches(value, card, settings.model.num_ctx):
                fake.record(recording_for(card, part.pages, NUMBER, data[key], body=part.body()))
            if not defense and value != values['elements']:
                errors.append('hearing and ours report elements packets differ')
            result = cli('report', *(['--binder', 'defense-report', '--defense'] if defense else ['--provider', 'Okonkwo-Lind', '--ours']))
            expected = report.render(data, value, cf, loaded.pages, info, open_items=counted['open_items'],
                                     playbook_version=book.version, model=settings.model.name)
            slug = 'defense-report' if defense else 'okonkwo-lind'
            output = folder / f'outputs/report-{slug}.docx'
            doc = _open_output(output, errors)
            hits = _rules(doc, expected, book.rules, outputs.providers(cf), errors, loaded.pages)
            rule_hits += hits
            lines.append(_report_line(doc, defense, result.returncode, cf, _header_ok(doc, cf, book, settings.model.name), hits))
            if result.stderr:
                errors.append('report: ' + result.stderr.strip())
        result = cli('status')
        fields = dict(re.findall(r'(\w+)=([^ ]+)', result.stdout.splitlines()[0]))
        partial = len(list((folder / 'outputs').glob('*.partial.docx')))
        lines.append(f'stage6 status outputs={fields["outputs"]} partial={partial}')
        if result.returncode or result.stderr or fields['open'] != '1' or fields['read'] != '13':
            errors.append('status CLI differs: ' + result.stdout + result.stderr)
        if (folder / 'matter.toml').read_bytes() != metadata:
            errors.append('non-interactive task wrote matter.toml')
        with open_db(settings.paths.state / 'runlog.sqlite') as log:
            answers = log.connection.execute('SELECT COUNT(*) FROM answers').fetchone()[0]
        docs = len(list((folder / 'outputs').glob('*.docx'))) - partial
        chats = sum(request['path'] == '/api/chat' for request in fake.requests)
        if chats != 16 or answers != 16:
            errors.append(f'calls={chats} answers={answers}, expected 16')
    finally:
        server.shutdown()
        thread.join(timeout=5)
        stopped = not thread.is_alive()
        server.server_close()
    failed = bool(errors or lines != list(EXPECTED[:7]) or docs != 6 or not stopped or rule_hits)
    lines.append(f'stage6 exit={int(failed)} docs={docs} of 6 answers={answers:+d} rule_hits={rule_hits} server_stopped={int(stopped)}')
    with capsys.disabled():
        print('\n'.join(lines), flush=True)
    mismatches = [f'expected: {expected}\nactual:   {actual}' for actual, expected in zip(lines, EXPECTED) if actual != expected]
    assert lines == list(EXPECTED) and not errors, 'stage6 exit check failed:\n' + '\n'.join(errors + mismatches)

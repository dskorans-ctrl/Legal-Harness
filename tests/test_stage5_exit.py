"""Stage 5 exit proof: four CLI jobs, durable coverage and reopened Word memos."""

import json
import os
import re
import shutil
import subprocess
from threading import Thread
from urllib.request import ProxyHandler, Request, build_opener

from docx import Document
import pytest

from casefile import config, extract, facts, matter, playbook
from casefile.runlog import open_db
from conftest import ROOT, write_config
from fakemodel import FakeModel
from fakemodel.bench import recording_for, recordings_for
from fakemodel.server import serve


MATTER = ROOT / "tests/fixtures/matters/9001.001"
NUMBER = "9001.001"
ONE_PAGE = {"facts": [{
    "cite": "DEF0001", "quote": "The examinee reports lifting a tote of about 60 pounds.",
    "text": "The examinee reported lifting a tote of about 60 pounds.",
}], "flags": []}
EXPECTED = (
    "stage5 run=1 mode=replay exit=1 chunks=1 runs=+1 read=13 open=1 facts=31 conflicts=5 gaps=4 new_facts=31 new_flags=9",
    "stage5 run=2 mode=bad_schema exit=1 chunks=1 runs=+1 read=13 open=1 facts=31 conflicts=5 gaps=4 new_facts=0 new_flags=0",
    "stage5 run=3 mode=replay exit=0 chunks=1 runs=+1 read=14 open=0 facts=32 conflicts=5 gaps=4 new_facts=1 new_flags=0",
    "stage5 run=4 mode=replay exit=0 chunks=0 runs=+0 read=14 open=0 facts=32 conflicts=5 gaps=4 new_facts=0 new_flags=0",
    "stage5 memo=1 quotes=31 of 31 by_page=1 cites=45 readings=31 conflicts=5 gaps=4 runs_seen=1 open=1 open_status=uncited rule_hits=0 changed=first",
    "stage5 memo=2 quotes=31 of 31 by_page=1 cites=45 readings=31 conflicts=5 gaps=4 runs_seen=1 open=1 open_status=failed rule_hits=0 changed=0",
    "stage5 memo=3 quotes=32 of 32 by_page=1 cites=46 readings=32 conflicts=5 gaps=4 runs_seen=2 open=0 open_status=none rule_hits=0 changed=1",
    "stage5 status read=13 then 14 open=1 then 0 outputs=memo:current",
    "stage5 exit=0 runs=4 of 4 memos=3 of 3 answers=+4 casefile_lines=48 chain_ok=1 server_stopped=1",
)
BAD_REASON = ("gate.schema: the answer does not match the schema: facts[0].text: missing required key; "
              "the page stays open until a later run reads it")
SECTIONS = ("What changed since the last memo", "Outputs out of date", "Open items", "Conflicts",
            "Gaps and missing records", "Verified facts by record")


def _section(document, title):
    paragraphs = document.paragraphs
    start = next(i for i, p in enumerate(paragraphs) if p.style.name == "Heading 1" and p.text == title) + 1
    end = next((i for i in range(start, len(paragraphs)) if paragraphs[i].style.name == "Heading 1"), len(paragraphs))
    return paragraphs[start:end]


def _cite(cite, page_id):
    return f"[page {page_id}]" if cite == page_id else f"[cite {cite}]"


def _read_memo(number, path, stored, pages, coverage, book, errors):
    rules = book.rules
    document = Document(path)
    paragraphs = document.paragraphs
    opened = _section(document, "Open items")
    verified = _section(document, "Verified facts by record")
    conflicts = _section(document, "Conflicts")
    gaps = _section(document, "Gaps and missing records")
    quote_texts = [f'"{fact.quote}"' for fact in stored.facts]
    quote_runs = [run.text for p in paragraphs for run in p.runs if run.text in quote_texts]
    quotes = sum(quote_runs.count(text) == 1 for text in quote_texts)
    expected_page_texts = []
    for page in pages:
        label = f'{" ".join(page.stamps)} ({page.page_id})' if page.stamps else page.page_id
        expected_page_texts.append(("Heading 2", f"{page.binder} page {label}"))
        for fact in stored.facts:
            if fact.page_ids[0] != page.page_id:
                continue
            expected_page_texts.append(("Normal", f'"{fact.quote}" {_cite(fact.cites[0], fact.page_ids[0])}'))
            if len(fact.cites) > 1:
                expected_page_texts.append(("Normal", "Also cited: " + " and ".join(
                    _cite(cite, page_id) for cite, page_id in zip(fact.cites[1:], fact.page_ids[1:]))))
            expected_page_texts.extend(("Normal", f"Model's reading: {reading}") for reading in fact.readings)
    by_page = int([(p.style.name, p.text) for p in verified] == expected_page_texts)
    readings = sum(p.text.startswith("Model's reading: ") for p in verified)
    expected_flags = {}
    for kind in ("conflict", "gap"):
        texts = []
        for flag in stored.flags:
            if flag.kind != kind:
                continue
            prefix = "Conflict between " if kind == "conflict" else "Gap at "
            texts.append(prefix + " and ".join(_cite(cite, page_id)
                         for cite, page_id in zip(flag.cites, flag.page_ids)) + ".")
            texts.extend(f"Model's reading: {about}" for about in flag.abouts)
        expected_flags[kind] = texts
    for kind, section in (("conflict", conflicts[1:]), ("gap", gaps)):
        if [p.text for p in section] != expected_flags[kind]:
            errors.append(f"memo={number}: {kind} entries or model readings differ")
    if [p.text for p in paragraphs if p.style.name == "Heading 1"] != list(SECTIONS):
        errors.append(f"memo={number}: section order differs")
    if paragraphs[0].text != f"Medical memo: matter {NUMBER}" or paragraphs[0].style.name != "Title":
        errors.append(f"memo={number}: title differs")
    output = stored.outputs[-1]
    if not re.fullmatch(r"built \d{4}-\d{2}-\d{2} from case file entry " + str(output["built_from_seq"]) +
                        re.escape(f"; playbook {book.version}; model qwen3.8:27b"), paragraphs[1].text):
        errors.append(f"memo={number}: build line differs: {paragraphs[1].text}")
    expected_open = []
    reasons = []
    for page in pages:
        state, reason, _ = coverage[(page.page_id, extract.card().name)]
        if state != "read":
            label = f'{" ".join(page.stamps)} ({page.page_id})' if page.stamps else page.page_id
            expected_open.append(f"{label}: {state}: {reason}")
            reasons.append(reason)
    if [p.text for p in opened] != expected_open:
        errors.append(f"memo={number}: open items differ: {[p.text for p in opened]}")
    exempt_reasons = set()
    for paragraph, reason in zip(opened, reasons):
        if len(paragraph.runs) == 2 and paragraph.runs[1].text == reason:
            exempt_reasons.add(paragraph._p)
        else:
            errors.append(f"memo={number}: reason is not a separate run")
    rule_hits = 0
    for paragraph in paragraphs:
        own_words = "".join("\n" if run.text in quote_texts or
                            (paragraph._p in exempt_reasons and i == 1) else run.text
                            for i, run in enumerate(paragraph.runs))
        rule_hits += sum(hit.severity == "error" for hit in rules.check(own_words, unit="text", stamped=True))
    sentence = re.fullmatch(r"Pages read in (\d+) runs?; conflicts and gaps are checked only between "
                            r"pages sent to the model together\.", conflicts[0].text)
    runs_seen = int(sentence[1]) if sentence else -1
    read_runs = {row[2] for row in coverage.values() if row[0] == "read"}
    if runs_seen != len(read_runs):
        errors.append(f"memo={number}: runs_seen={runs_seen} coverage_read_runs={len(read_runs)}")
    changed_text = [p.text for p in _section(document, SECTIONS[0])]
    expected_change = ("First memo for this matter" if number == 1 else
                       f"{int(number == 3)} fact added, 0 cites added, 0 readings added, 0 flags added")
    if changed_text != [expected_change]:
        errors.append(f"memo={number}: changes differ: {changed_text}")
    changed = "first" if changed_text == ["First memo for this matter"] else (
        changed_text[0].split(" ", 1)[0] if changed_text else "missing")
    if [p.text for p in _section(document, SECTIONS[1])] != ["memo" if number == 3 else "none"]:
        errors.append(f"memo={number}: outputs out of date differ")
    cites = sum(p.text.count("[cite ") for p in paragraphs)
    expected_cites = sum(len(f.cites) for f in stored.facts) + sum(len(f.cites) for f in stored.flags)
    if cites != expected_cites:
        errors.append(f"memo={number}: cites={cites} casefile_cites={expected_cites}")
    open_status = opened[0].text.split(": ")[1] if opened else "none"
    return (f"stage5 memo={number} quotes={quotes} of {len(stored.facts)} by_page={by_page} cites={cites} "
            f"readings={readings} conflicts={sum(p.text.startswith('Conflict between ') for p in conflicts)} "
            f"gaps={sum(p.text.startswith('Gap at ') for p in gaps)} runs_seen={runs_seen} "
            f"open={len(opened)} open_status={open_status} rule_hits={rule_hits} changed={changed}")


@pytest.mark.socket
def test_stage5_exit(tmp_path, capsys):
    loaded = matter.load(MATTER)
    page = next(page for page in loaded.pages if page.stamps == ("DEF0001",))
    assert ONE_PAGE["facts"][0]["quote"] in page.text
    fake = FakeModel(recordings=[*recordings_for(MATTER), recording_for(extract.card(), (page,), NUMBER, ONE_PAGE)])
    server = serve(port=0, model=fake)
    thread = Thread(target=server.serve_forever)
    thread.start()
    run_lines, memo_lines, statuses, errors = [], [], [], []
    try:
        path = write_config(tmp_path, host=f"http://127.0.0.1:{server.server_port}", playbook=ROOT / "playbook")
        settings = config.load(path)
        folder = settings.paths.matters / NUMBER
        shutil.copytree(MATTER, folder)
        env = dict(os.environ, CASEFILE_CONFIG=str(path))
        opener = build_opener(ProxyHandler({}))
        book = playbook.load(settings.paths.playbook)
        previous_seq = previous_runs = previous_answers = 0
        with open_db(settings.paths.state / "runlog.sqlite") as log:
            for number, mode in enumerate(("replay", "bad_schema", "replay", "replay"), 1):
                request = Request(settings.model.host + "/_fake/mode", method="POST",
                                  data=json.dumps({"mode": mode}).encode("utf-8"),
                                  headers={"Content-Type": "application/json"})
                with opener.open(request, timeout=5) as response:
                    assert response.status == 200 and json.load(response) == {"mode": mode}
                requests0 = len(fake.requests)
                result = subprocess.run(["scripts/run.sh", "casefile", "memo", "--matter", NUMBER],
                                        cwd=ROOT, env=env, capture_output=True, text=True, timeout=60)
                stored = facts.load(folder / "casefile.jsonl")
                coverage = log.coverage(NUMBER, job="memo")
                runs = log.connection.execute("SELECT COUNT(*) FROM runs").fetchone()[0]
                answers = log.connection.execute("SELECT COUNT(*) FROM answers").fetchone()[0]
                counts = {state: sum(row[0] == state for row in coverage.values())
                          for state in ("read", "uncited", "failed")}
                opened = counts["uncited"] + counts["failed"]
                changes = stored.changed_since(previous_seq)
                chats = [r for r in fake.requests[requests0:] if r["path"] == "/api/chat"]
                line = (f"stage5 run={number} mode={mode} exit={result.returncode} chunks={runs - previous_runs} "
                        f"runs={runs - previous_runs:+d} read={counts['read']} open={opened} "
                        f"facts={len(stored.facts)} conflicts={sum(f.kind == 'conflict' for f in stored.flags)} "
                        f"gaps={sum(f.kind == 'gap' for f in stored.flags)} new_facts={changes['fact']} "
                        f"new_flags={changes['flag']}")
                run_lines.append(line)
                # CLI counts must agree with the independently read case file and coverage.
                expected_stdout = (f"memo {NUMBER} chunks=" + line.split(" chunks=", 1)[1] +
                                   f" memo={folder / 'outputs/memo.docx'}")
                expected_open = ([] if not opened else [
                    "  open: DEF0001 (defense-report/001): " +
                    ("uncited: no fact or flag cites defense-report/001" if number == 1 else "failed: " + BAD_REASON)])
                if result.stdout.splitlines() != [expected_stdout, *expected_open] or result.stderr:
                    errors.append(f"run={number}: CLI differs: {result.stdout!r} {result.stderr!r}")
                expected_chats = (1, 2, 1, 0)[number - 1]
                if len(chats) != expected_chats or answers - previous_answers != expected_chats:
                    errors.append(f"run={number}: chats={len(chats)} answers={answers - previous_answers}")
                expected_pages = [p.page_id for p in loaded.pages] if number == 1 else [page.page_id]
                for chat in chats:
                    body = json.loads(chat["body"])
                    sent = re.findall(r"^=== PAGE (\S+) BATES ", body["messages"][1]["content"], re.MULTILINE)
                    if sent != expected_pages:
                        errors.append(f"run={number}: sent pages differ: {sent}")
                expected_line_count = (42, 44, 47, 48)[number - 1]
                if len(stored.events) != expected_line_count:
                    errors.append(f"run={number}: casefile_lines={len(stored.events)} expected={expected_line_count}")
                if number <= 3:
                    memo_lines.append(_read_memo(number, folder / "outputs/memo.docx", stored,
                                                 loaded.pages, coverage, book, errors))
                if number in (1, 3):
                    status = subprocess.run(["scripts/run.sh", "casefile", "status", "--matter", NUMBER],
                                            cwd=ROOT, env=env, capture_output=True, text=True, timeout=60)
                    fields = dict(re.findall(r"(\w+)=([^ ]+)", status.stdout.splitlines()[0]))
                    statuses.append(fields)
                    expected_status = (f"status {NUMBER} pages=14 read={counts['read']} failed=0 "
                                       f"uncited={opened} empty=0 open={opened} facts={len(stored.facts)} "
                                       "conflicts=5 gaps=4 outputs=memo:current")
                    if status.returncode != 0 or status.stderr or status.stdout.splitlines() != [expected_status, *expected_open]:
                        errors.append(f"run={number}: status differs: {status.stdout!r} {status.stderr!r}")
                previous_seq, previous_runs, previous_answers = stored.last_seq, runs, answers
    finally:
        server.shutdown()
        thread.join(timeout=5)
        stopped = not thread.is_alive()
        server.server_close()
    status_line = (f"stage5 status read={statuses[0]['read']} then {statuses[1]['read']} "
                   f"open={statuses[0]['open']} then {statuses[1]['open']} outputs={statuses[1]['outputs']}")
    chain_ok = int(facts.load(folder / "casefile.jsonl").torn == 0)
    complete_lines = len((folder / "casefile.jsonl").read_bytes().splitlines())
    lines = [*run_lines, *memo_lines, status_line]
    good_runs = sum(actual == expected for actual, expected in zip(run_lines, EXPECTED[:4]))
    good_memos = sum(actual == expected for actual, expected in zip(memo_lines, EXPECTED[4:7]))
    failed = bool(errors or lines != list(EXPECTED[:8]) or answers != 4 or complete_lines != 48
                  or not chain_ok or not stopped)
    lines.append(f"stage5 exit={int(failed)} runs={good_runs} of 4 memos={good_memos} of 3 answers={answers:+d} "
                 f"casefile_lines={complete_lines} chain_ok={chain_ok} server_stopped={int(stopped)}")
    with capsys.disabled():
        print("\n".join(lines), flush=True)
    mismatches = [f"expected: {expected}\nactual:   {actual}"
                  for actual, expected in zip(lines, EXPECTED) if actual != expected]
    assert lines == list(EXPECTED) and not errors, "stage5 exit check failed:\n" + "\n".join(errors + mismatches)

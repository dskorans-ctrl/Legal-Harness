"""Exercise checked publication, preflight order and the shared document contract."""

from dataclasses import replace
from io import StringIO
import json
import os
from pathlib import Path
import shutil
import subprocess

from docx import Document
import pytest

from casefile import cards, config, facts, matter, matterinfo, outputs, packet, playbook, status
from conftest import ROOT, write_config
from fakemodel import FakeModel, FakeTransport
from fakemodel.bench import recording_for

ANSWERS = ROOT / 'tests/fixtures/outputs/9001.001'


def answer(name):
    return json.loads((ANSWERS / (name + '.json')).read_text(encoding='utf-8'))


def setup_case(tmp_path):
    path = write_config(tmp_path)
    cfg = config.load(path)
    shutil.copytree(ROOT / 'playbook', cfg.paths.playbook, dirs_exist_ok=True)
    folder = cfg.paths.matters / '9001.001'
    shutil.copytree(ROOT / 'tests/fixtures/matters/9001.001', folder)
    loaded = matter.load(folder)
    book = playbook.load(cfg.paths.playbook)
    ideal = json.loads((folder / 'ideal.json').read_text(encoding='utf-8'))
    cf = facts.CaseFile()
    cf = facts.append(folder / 'casefile.jsonl', facts.merge(cf, ideal, loaded.pages, 1))
    # Tests of outputs begin from a fully read memo, not the shared fixture folder.
    cf = facts.append(folder / 'casefile.jsonl', [facts.Event('run', dict(run_id=1, job='memo',
        card='medical_memo', card_version='invented', pages=[p.page_id for p in loaded.pages], ok=True,
        read=[p.page_id for p in loaded.pages], uncited=[], failed=[]))])
    from casefile.runlog import open_db
    with open_db(cfg.paths.state / 'runlog.sqlite') as log:
        run = log.start_run(job='memo', matter=loaded.number, card='medical_memo', card_version='invented',
                            playbook_version=book.version, model=cfg.model.name, pages_in=14, prompt_sha256='invented')
        assert run == 1
        for page in loaded.pages:
            log.record_coverage(run, matter=loaded.number, page_id=page.page_id, card='medical_memo', status='read', reason='')
    return path, cfg, folder, loaded, cf, book


def flat(doc):
    return [paragraph.text for paragraph in doc.paragraphs]


def section(doc, title):
    paragraphs = doc.paragraphs
    start = next(i for i,p in enumerate(paragraphs) if p.style.name == 'Heading 1' and p.text == title) + 1
    end = next((i for i in range(start,len(paragraphs)) if paragraphs[i].style.name == 'Heading 1'), len(paragraphs))
    return [p.text for p in paragraphs[start:end]]


def document(tmp_path, doc, rules, name='render.docx', providers=()):
    path = tmp_path / name
    outputs.render(doc, path, rules=rules, providers=providers)
    return Document(path)


@pytest.mark.parametrize('row', range(13))
def test_outputs_contract(tmp_path, monkeypatch, row, capsys):
    path, cfg, folder, loaded, cf, book = setup_case(tmp_path)
    info = matterinfo.load(folder)
    doc = outputs.Doc('Checked output', outputs.header(info, loaded.number, 'report'),
                      [outputs.Section('Evidence', [outputs.p('Checked.')]), outputs.open_items_section(())])
    counts = dict(facts=len(cf.facts), conflicts=5, gaps=4, open=0)
    if row == 0:
        lines = flat(document(tmp_path,doc,book.rules))
        assert lines[1:9] == ['Attorney Work Product | Privileged & Confidential', 'Matter type: workers-comp',
            'Client: Morgan Ashby-Vale', 'Date of accident: 03/11/2025', 'Treating physician: Dr. Priya Okonkwo-Lind',
            'IAB hearing number: 1234567', 'Hearing date: 01/15/2026 9:00 AM', 'Matter 9001.001'], 'W5 header order differs'
    elif row == 1:
        doc.header = outputs.header(replace(info, hearing_date=''), loaded.number, 'report')
        assert 'Hearing date: NEEDED' in flat(document(tmp_path,doc,book.rules)), 'W5 blank field did not print NEEDED'
    elif row == 2:
        # Use the intake layout so missing metadata is valid. The child exits through finish.
        (folder / 'matter.toml').unlink()
        target = folder / 'pages'
        for page in loaded.pages:
            (target / page.binder).mkdir(parents=True,exist_ok=True)
            data = dict(matter=loaded.number,binder=page.binder,page=page.number,id=page.page_id,
                        text=page.text,stamps=list(page.stamps),status='ok',reason='')
            # Intake stamps carry raw and normalized forms.
            data['stamps'] = [dict(raw=s, stamp=s) for s in page.stamps]
            (target / page.binder / f'{page.number:03d}.json').write_text(json.dumps(data),encoding='utf-8')
        for binder in {page.binder for page in loaded.pages}:
            (target / binder / 'binder.json').write_text(json.dumps({'pages': sum(page.binder == binder for page in loaded.pages)}), encoding='utf-8')
        code = '''from casefile import config, outputs
import sys
from casefile import matterinfo
original_ask = matterinfo.ask
def guarded_ask(info, stdin, stdout):
 assert stdin.isatty(), 'W5 non-tty start called ask'
 return original_ask(info, stdin, stdout)
matterinfo.ask = guarded_ask
with outputs.start(config.load(sys.argv[1]), '9001.001', 'report') as s:
 d=outputs.Doc('Report',outputs.header(s.info,'9001.001','report'),[outputs.open_items_section((),outputs.needed(s.info,s.loaded.pages))])
 sys.exit(outputs.finish(s,'missing-header',d,failed_steps=(),out=sys.stdout))'''
        result = subprocess.run(['scripts/run.sh','python','-c',code,str(path)],input='',text=True,
                                capture_output=True,timeout=20,env=os.environ)
        assert result.returncode == 1, 'W5 non-tty run blocked, wrote metadata, or failed to disclose missing header: ' + result.stderr
        assert not result.stderr, result.stderr
        assert not (folder/'matter.toml').exists(), 'W5 non-tty run wrote matter.toml'
        lines = flat(Document(folder/'outputs/missing-header.docx'))
        assert all(label + ': NEEDED' in lines for label in ('Matter type','Client','Date of accident',
                       'Treating physician','IAB hearing number','Hearing date'))
    elif row == 3:
        doc.sections[0].paragraphs = [outputs.p('What treatment did you receive and what bills followed?',unit='question')]
        hits = outputs.check(doc, book.rules, providers=outputs.providers(cf))
        assert any(f.rule_id=='treatment-and-bills-apart' for f in hits), 'W5 question was checked as text'
        doc.sections[0].paragraphs = [outputs.p('Doctor, as to the bills themselves, the charges from Ridgeline Physical Therapy, were those bills reasonable?',unit='question',bills_question=True)]
        assert not [f for f in outputs.check(doc,book.rules,providers=outputs.providers(cf)) if f.severity=='error']
    elif row == 4:
        doc.sections[0].paragraphs = [outputs.Paragraph([outputs.Run('IME',exempt='binder'),outputs.cite('ime/001')])]
        assert not outputs.check(doc,book.rules), 'W5 binder or cite was checked as prose'
    elif row == 5:
        seen=[]
        original = outputs.os.replace
        def rename(source,dest):
            seen.append((Path(source).name,Path(dest).name))
            return original(source,dest)
        monkeypatch.setattr(outputs.os,'replace',rename)
        new = outputs.publish(folder,'report-test',doc,book.rules,cf,counts,failed_steps=())
        assert seen == [('report-test.docx.tmp','report-test.docx')], 'W5 final path written without temporary publication'
        assert len(new.outputs)==1 and new.outputs[-1]['name']=='report-test'
        assert Document(folder/'outputs/report-test.docx')
    elif row == 6:
        new = outputs.publish(folder,'report-test',doc,book.rules,cf,counts,failed_steps=())
        item = dict(cf.events[0].payload,id='F999')
        new = facts.append(folder/'casefile.jsonl',[facts.Event('fact',item)])
        assert new.stale_outputs()==['report-test']
    elif row == 7:
        before = (folder/'casefile.jsonl').read_bytes()
        new = outputs.publish(folder,'report-test',doc,book.rules,cf,counts,failed_steps=['elements'])
        assert (folder/'outputs/report-test.partial.docx').is_file()
        assert not new.outputs and (folder/'casefile.jsonl').read_bytes()==before, 'W5 partial output was recorded as current'
        assert status.matter_status(cfg,loaded.number)['outputs']=='none;partial=report-test.partial.docx'
    elif row in (8,9,10):
        (cfg.paths.state/'runlog.sqlite').unlink()
        if row==8:
            (folder/'casefile.jsonl').unlink()
            expected='run casefile memo first'
        elif row==9:
            with (folder/'matter.toml').open('a',encoding='utf-8') as stream: stream.write('\n[unknown]\nx=1\n')
            expected='unknown'
        else:
            monkeypatch.setattr(outputs.rules,'check_stamps',lambda *args: (_ for _ in ()).throw(outputs.rules.RulesError('incompatible stamp')))
            expected='incompatible stamp'
        with pytest.raises(outputs.ERRORS,match=expected):
            with outputs.start(cfg,loaded.number,'report'): pass
        assert not (cfg.paths.state/'runlog.sqlite').exists(), 'W5 start opened run log before stamp or matter refusal'
    elif row==11:
        doc.sections[0].paragraphs = [outputs.failure('Elements','gate.quotes: quote not on MED0007')]
        doc.sections[-1] = outputs.open_items_section((),['Elements: failed section'])
        with outputs.start(cfg,loaded.number,'report') as started:
            assert outputs.finish(started,'failed',doc,failed_steps=['elements'],out=StringIO())==1
        lines = section(Document(folder/'outputs/failed.partial.docx'),'Evidence')
        assert lines[0]=="Elements: the model's answer failed the gate.quotes check after one retry: quote not on MED0007", 'W5 failed section was silently omitted'
    else:
        modified = facts.CaseFile(list(cf.events))
        modified.events[0] = replace(modified.events[0],payload={**modified.events[0].payload,'text':'nobody'})
        assert outputs.provider_pages(modified,'nobody')==(), 'W5 provider selection searched the fact text'
        assert outputs.provider_pages(cf,'okonkwo-lind')==('medical/007','medical/008')
        widened = outputs.provider_pages(cf,'Marchetti',pages=loaded.pages,info=info)
        assert set(widened)=={'defense-report/001','defense-report/002','defense-report/003'}, 'W5 defense provider binder was not widened to DEF0001 to DEF0003'
        assert outputs.provider_pages(cf,'Okonkwo-Lind',pages=loaded.pages,info=info)==('medical/007','medical/008'), 'W5 ours provider widened beyond its pages'
        value = packet.build(cf,loaded.pages,matter=loaded.number,task='report',options={},playbook_version=book.version)
        card = cards.card('02-A',book)
        data = {'facts':[dict(cite=value.facts[0].cite,quote=value.facts[0].quote,use='both',note='')]}
        fake=FakeModel(recordings=[recording_for(card,value.pages,value.matter,data,body=value.body())])
        with outputs.start(cfg,loaded.number,'report',opener=FakeTransport(fake)) as started:
            result=outputs.ask(value,card,client=started.client,rules=book.rules,log=started.log,
                               job='report',playbook_version=book.version)
            assert result.ok, 'W5 ask omitted packet body or lookup pages'
    if row==12:
        with capsys.disabled(): print('outputs_ok=13 of 13')


def test_evidence_closed():
    for key in ('quote', 'cite'):
        item = {'text': 'Invented evidence item', 'quote': 'Invented quote', 'cite': 'MED0001'}
        del item[key]
        with pytest.raises(outputs.OutputError) as error:
            outputs.evidence(item)
        assert str(error.value) == f'evidence item {item!r}: missing {key}'
    assert outputs.question_paragraphs(None, 'Questions', reason='gate.schema: missing quote')[0].runs
    print('evidence_closed=1')


def test_roles_through_start(tmp_path):
    from test_matterinfo import Terminal

    _, cfg, folder, loaded, _, _ = setup_case(tmp_path)
    info = replace(matterinfo.load(folder), plaintiffs=2)
    matterinfo.write(folder, info)
    before = (folder / 'matter.toml').read_bytes()
    non_tty = StringIO('ours\nother\n')
    with outputs.start(cfg, loaded.number, 'report', stdin=non_tty, out=StringIO()) as started:
        assert outputs.needed(started.info, loaded.pages) == ['ROLE NEEDED: medical', 'ROLE NEEDED: employer']
    assert non_tty.tell() == 0 and (folder / 'matter.toml').read_bytes() == before
    out = StringIO()
    with outputs.start(cfg, loaded.number, 'report', stdin=Terminal('ours\nother\n'), out=out) as started:
        assert started.info.has_header
        assert started.info.roles == {'defense-report': 'defense', 'medical': 'ours', 'employer': 'other'}
        assert not outputs.needed(started.info, loaded.pages)
    assert out.getvalue() == 'Binder medical: defense, ours or other: Binder employer: defense, ours or other: '
    saved = (folder / 'matter.toml').read_bytes()
    inp = Terminal('defense\n')
    with outputs.start(cfg, loaded.number, 'report', stdin=inp, out=StringIO()):
        pass
    assert inp.tell() == 0 and (folder / 'matter.toml').read_bytes() == saved

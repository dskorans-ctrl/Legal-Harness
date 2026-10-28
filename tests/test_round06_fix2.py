"""Document regressions for the round 6 output critic and append recovery."""

from dataclasses import replace
from io import StringIO
import re

from docx import Document
import pytest

from casefile import depo, facts, gates, hearing, matterinfo, memo, memo_job, outputs, report, status
from test_depo import answers_for
from test_hearing import hearing_answers
from test_outputs import answer, document, flat, section, setup_case


@pytest.fixture
def case(tmp_path):
    _, cfg, folder, loaded, cf, book = setup_case(tmp_path)
    return cfg, folder, loaded, cf, book, matterinfo.load(folder)


def built(tmp_path, case, kind, *, data=None, cf=None):
    cfg, folder, loaded, original, book, info = case
    cf = original if cf is None else cf
    kwargs = dict(open_items=(), playbook_version=book.version, model=cfg.model.name)
    if kind.startswith('depo-'):
        who = kind.removeprefix('depo-')
        value = depo.prepare(cf, loaded.pages, book, info, number=loaded.number,
            deponent=who, taken_by='us', provider='Marchetti' if who == 'defense-doctor' else 'Okonkwo-Lind')
        doc = depo.render(answers_for(who) if data is None else data, value, cf, loaded.pages, info, **kwargs)
    elif kind == 'hearing':
        values = hearing.prepare(cf, loaded.pages, book, info, number=loaded.number, provider='Okonkwo-Lind')
        doc = hearing.render(hearing_answers(values) if data is None else data,
            values, cf, loaded.pages, info, **kwargs)
    else:
        defense = kind == 'report-defense'
        value = report.prepare(cf, loaded.pages, book, info=info, number=loaded.number,
            binder=('defense-report',) if defense else (), provider='' if defense else 'Okonkwo-Lind', defense=defense)
        data = data if data is not None else {
            'concessions' if defense else 'elements': answer('concessions-defense-report' if defense else 'elements-okonkwo-lind')}
        doc = report.render(data, value, cf, loaded.pages, info, **kwargs)
    return document(tmp_path, doc, book.rules, kind + '.docx', providers=outputs.providers(cf))


def built_memo(tmp_path, case, cf=None):
    cfg, _, loaded, original, book, info = case
    doc = memo.build(original if cf is None else cf, loaded.pages, (), matter=loaded.number,
                     playbook_version=book.version, model=cfg.model.name, info=info)
    path = tmp_path / 'memo.docx'
    memo.render(doc, path, rules=book.rules)
    return Document(path)


def test_h4_empty_conflict(tmp_path, case):
    cf = case[3]
    # G6 has no verified fact on DEF0002 in the old key. Keep that case even after the fixture grows.
    events = [e for e in cf.events if not (e.kind in ('fact', 'cite', 'reading') and e.payload['page_id'] == 'defense-report/002')]
    cf = facts.CaseFile(events)
    for doc in (built(tmp_path, case, 'report-defense', cf=cf), built_memo(tmp_path, case, cf)):
        lines = flat(doc)
        places = [i for i, text in enumerate(lines) if text == 'From [cite DEF0002]:']
        assert places and all(lines[i + 1] == 'No verified fact quoted from DEF0002; read the page.' for i in places), 'H4: empty conflict block hides missing evidence'


def test_m5_conflict_categories(tmp_path, case):
    cf = case[3]
    for doc in (built(tmp_path, case, 'report-defense'), built_memo(tmp_path, case)):
        lines = section(doc, 'Conflicts with other records' if doc.paragraphs[0].text == 'Report analysis' else 'Conflicts')
        # Both documents include G6: the matching examination category excludes diagnosis.
        start = lines.index('From [cite MED0007]:') + 1
        end = lines.index('From [cite DEF0002]:', start)
        assert not any('Impression:' in text for text in lines[start:end]), 'M5: conflict includes an unmatched diagnosis category'
    doc = built_memo(tmp_path, case)
    lines = section(doc, 'Conflicts')
    start = lines.index('From [cite MED0001]:') + 1
    end = next(i for i in range(start, len(lines)) if lines[i].startswith('From '))
    assert not any('Date of service' in line or 'Pain 7/10' in line for line in lines[start:end]), 'M5: weight conflict dumps unrelated facts'


def test_m5_fallback(tmp_path, case):
    cf = case[3]
    flag = next(f for f in cf.flags if set(f.cites) == {'MED0006', 'DEF0002'})
    # Remove the matching imaging category from one side to force explicit fallback.
    events = [replace(e, payload={**e.payload, 'category': 'other'}) if e.kind == 'fact' and
              e.payload['page_id'] == 'defense-report/002' else e for e in cf.events]
    cf = facts.CaseFile(events)
    paragraphs = outputs.conflicts(cf, case[2].pages, case[5], predicate=lambda f: f.id == flag.id)
    doc = document(tmp_path, outputs.Doc('Conflict', [], [outputs.Section('Conflict', paragraphs)]), case[4].rules)
    assert '(all facts on this page)' in flat(doc), 'M5: fallback facts are not labelled'


def test_h2_bills_providers(tmp_path, case):
    # Provider selection is code, independent of what the recording chooses to say.
    cfg, _, loaded, cf, book, info = case
    value = depo.prepare(cf, loaded.pages, book, info, number=loaded.number,
                         deponent='our-doctor', taken_by='us', provider='Okonkwo-Lind')
    assert value.options.get('bills_providers') == [p for p in outputs.providers(cf) if 'Marchetti' not in p], 'H2: bills packet includes defense-only providers or has no provider list'
    # A provider appearing on both sides must remain eligible.
    first = next(e for e in cf.events if e.kind == 'fact' and e.payload['provider'])
    mixed = facts.CaseFile([*cf.events, facts.Event('fact', {**first.payload, 'id': 'F999',
        'page_id': 'medical/001', 'cite': 'MED0001', 'provider': 'Dr. Theodore Marchetti'})])
    value = depo.prepare(mixed, loaded.pages, book, info, number=loaded.number,
                         deponent='our-doctor', taken_by='us', provider='Okonkwo-Lind')
    assert 'Dr. Theodore Marchetti' in value.options['bills_providers'], 'H2: mixed-role provider was excluded'


def test_m1_strongest(tmp_path, case):
    values = hearing.prepare(case[3], case[2].pages, case[4], case[5], number=case[2].number, provider='Okonkwo-Lind')
    data = hearing_answers(values)
    # 'both' indicates relevance, not a helpfulness ranking, even when the fact is neutral.
    for key in data:
        if key.startswith('factsheet:'):
            next(item for item in data[key]['facts'] if 'Date of service' in item['quote'])['use'] = 'both'
    doc = built(tmp_path, case, 'hearing', data=data)
    lines = section(doc, 'Closing outline')
    start, end = lines.index('Our strongest facts') + 1, lines.index('Defense says')
    assert not any('40 pounds' in x or 'denies any prior' in x or 'Date of service' in x or 'negative bilaterally' in x for x in lines[start:end]), 'M1: strongest facts include harmful or unranked facts'
    for key in data:
        if key.startswith('factsheet:'):
            data[key] = {'facts': []}
    lines = section(built(tmp_path, case, 'hearing', data=data), 'Closing outline')
    assert lines.count('Facts (not ranked)') == 2 and 'Our strongest facts' not in lines, 'M1: unlabelled facts claim a ranking'


@pytest.mark.parametrize('kind', ['depo-defense-doctor', 'depo-our-doctor', 'hearing'])
def test_m2_defense_roles(tmp_path, case, kind):
    doc = built(tmp_path, case, kind)
    lines = section(doc, 'Hearing outline' if kind == 'hearing' else 'Case outline: the defense view')
    if kind == 'hearing':
        lines = lines[lines.index('Defense position') + 1:lines.index('Conflicts')]
    assert any('[cite DEF' in x for x in lines) and not any('[cite MED' in x or '[cite EMP' in x for x in lines), 'M2: defense view includes our records'


def test_m2_closing_roles(tmp_path, case):
    lines = section(built(tmp_path, case, 'hearing'), 'Closing outline')
    start, end = lines.index('Defense says') + 1, lines.index('Our answer')
    assert any('[cite DEF' in x for x in lines[start:end]) and not any('[cite MED' in x for x in lines[start:end]), 'M2: closing defense position includes our records'


def test_m3_headings(tmp_path, case):
    doc = built(tmp_path, case, 'hearing')
    inside = False
    for p in doc.paragraphs:
        if p.style.name == 'Heading 1':
            inside = p.text == 'Hearing outline'
        if inside and p.text in outputs.GROUPS:
            assert p.style.name == 'Heading 4', 'M3: category heading sits outside its issue subsection'
    paragraphs = doc.paragraphs
    i = next(i for i,p in enumerate(paragraphs) if p.text == "What the doctor's records say")
    assert paragraphs[i + 1].style.name == 'Heading 3', 'M3: doctor-record heading has no nested content'


def test_m4_conflicts_once(tmp_path, case):
    lines = section(built(tmp_path, case, 'hearing'), 'Doctor preparation')
    headers = [x for x in lines if x.startswith('Conflict between ')]
    assert headers and len(headers) == len(set(headers)), 'M4: doctor section repeats a conflict block'


def test_h3_m9_honest_absence(tmp_path, case):
    data = {'concessions': answer('concessions-defense-report')}
    data['concessions']['undefined_labels'][0]['supporting_findings'] = []
    doc = built(tmp_path, case, 'report-defense', data=data)
    text = '\n'.join(flat(doc))
    assert 'model returned no item' in text and 'DEF0002' in text and 'none stated on the pages sent' not in text, 'H3: model omission is presented as absence on the page'
    statements = section(doc, 'What the report states')
    assert any('model returned no item' in x and 'DEF0001' in x for x in statements), 'M9: report omits unrepresented pages without disclosure'
    assert any('Records referenced' in x for x in statements) or any('No records-referenced item' in x for x in flat(doc)), 'M9: cross-reference points to absent records list'


@pytest.mark.parametrize('field,title', [('concessions', "Concessions (the report's own words)"),
    ('undefined_labels','Undefined labels'), ('danger_zone','Danger zone')])
def test_h3_empty_report_sections(tmp_path, case, field, title):
    data = answer('concessions-defense-report')
    data[field] = []
    doc = built(tmp_path, case, 'report-defense', data={'concessions':data})
    lines = section(doc, title)
    assert any('model returned no item' in x and 'DEF0001' in x and 'DEF0002' in x and 'DEF0003' in x for x in lines), 'H3: empty report section omits model limit and searched pages'


def test_l7_multiple_page_dates(case):
    from casefile.matter import Page
    page = Page('medical', 1, 'Date of service: 03/11/2025\nDate of service: 03/14/2025\nNormal gait.', ('ABC0001',))
    detail, _ = gates._references({'facts':[dict(cite='ABC0001',quote='Normal gait.',date='03/11/2025')]}, (page,), case[4].rules)
    assert 'faults=' not in detail, 'L7: a parseable date written on the cited page was refused'


@pytest.mark.parametrize('card,text', [
    ('01-MEMO', 'copied exactly as written on the page'),
    ('02-A', 'Do not label a harmful fact helpful'),
    ('02-C-DEFDOC', 'Do not turn an adverse conclusion'),
    ('02-C-OURDOC', 'Anchor treatment to a treatment quote'),
    ('04-G', 'Read the entire examination'),
    ('05-A', 'Do not label a permanency rating present from an MMI statement'),
])
def test_practice_cards(card, text):
    from casefile import cards
    assert text in cards.card(card).text, 'Practice card omits ruled instruction: ' + card


def test_l1_jargon(tmp_path, case):
    docs = [built(tmp_path, case, k) for k in ('hearing','report-defense','report-ours')]
    docs.append(built_memo(tmp_path, case))
    assert not re.search(r'question \d|stage \d|this round|card runs', '\n'.join(x for d in docs for x in flat(d))), 'L1: attorney document contains build jargon'


def test_l2_human_labels(tmp_path, case):
    docs = [built(tmp_path, case, k) for k in ('hearing','report-defense','report-ours')]
    docs.append(built_memo(tmp_path, case))
    assert not re.search(r'\b[a-z]+_[a-z_]+:', '\n'.join(x for d in docs for x in flat(d))), 'L2: attorney document exposes raw field labels'


def test_l3_empty_headings(tmp_path, case):
    doc = built_memo(tmp_path, case)
    lines = section(doc, 'Facts without a date, by page')
    paragraphs = doc.paragraphs
    for i,p in enumerate(paragraphs):
        if p.style.name == 'Heading 2' and p.text.startswith('Page '):
            assert i + 1 < len(paragraphs) and not paragraphs[i + 1].style.name.startswith('Heading'), 'L3: undated page heading has no content'
    assert lines, 'L3: empty section has no none marker'


def test_l6_deponent_placeholder(tmp_path, case):
    doc = document(tmp_path, outputs.Doc('Deposition', outputs.header(case[5], '9001.001','depo'), []), case[4].rules)
    assert 'Deposition of: NEEDED (doctor\'s name)' in flat(doc) and 'Deposition date and time: NEEDED' in flat(doc), 'L6: deposition placeholder is unlabelled'


@pytest.mark.parametrize('date,good,message', [
    ('03/11/2025', True, 'L7: record date on the cited page but outside the quote was refused'),
    ('03/12/2025', False, 'L7: date written on no page was accepted'),
    ('03/14/2025', False, 'L7: date written only on another page was accepted'),
    ('02/30/2025', False, 'L7: unparseable date written on the cited page was accepted'),
])
def test_l7_date_rule(case, date, good, message):
    from casefile.matter import Page
    base = dict(cite='ABC0001', quote='Date of hire: 08/02/2021', date=date)
    page = Page('medical', 1, 'Date of Report: 03/11/2025\nDate of hire: 08/02/2021\nInvalid date: 02/30/2025', ('ABC0001',))
    other = Page('medical', 2, 'Date of Report: 03/14/2025', ('ABC0002',))
    detail, _ = gates._references({'facts': [base]}, (page, other), case[4].rules)
    assert ('faults=' not in detail) == good, message
    if not good:
        assert '.date:' in detail, message


def test_chronology_record_date_label(tmp_path, case):
    doc = built_memo(tmp_path, case)
    first = section(doc, 'Treatment chronology')[0]
    assert first == 'Entries run by the date of the record.', 'Chronology: record-date convention missing from first paragraph'
    assert not case[4].rules.check(first, unit='text', stamped=True), 'Chronology: record-date convention violates firm rules'


@pytest.mark.parametrize('kind', ['report','memo'])
def test_a1_append_failure(tmp_path, case, monkeypatch, kind):
    cfg, folder, loaded, cf, book, info = case
    counts = dict(facts=len(cf.facts), conflicts=5,gaps=4,open=0)
    doc = outputs.Doc('Old', [], [outputs.open_items_section(())])
    name = 'memo' if kind == 'memo' else 'report-test'
    cf = outputs.publish(folder, name, doc, book.rules, cf, counts, failed_steps=())
    target = folder / 'outputs' / (name + '.docx')
    old = target.read_bytes()
    append = facts.append
    def fail_output(path, events):
        if any(e.kind == 'output' for e in events):
            raise OSError('seeded output append failure')
        return append(path, events)
    monkeypatch.setattr(facts, 'append', fail_output)
    if kind == 'report':
        with pytest.raises(OSError, match='seeded output append failure'):
            outputs.publish(folder, name, replace(doc, title='New'), book.rules, cf, counts, failed_steps=())
    else:
        assert memo_job.run(cfg, loaded.number, out=StringIO(), err=StringIO()) == 2
    assert target.read_bytes() == old, 'A1: append failure replaced the published Word file'
    assert name + ':stale' in status.matter_status(cfg, loaded.number)['outputs'], 'A1: append failure leaves status current'


def test_practice_defense(tmp_path, case):
    doc = built(tmp_path, case, 'depo-defense-doctor')
    primary = [x for x in section(doc, 'Questions: primary') if re.match(r'^\d+\. ', x)]
    assert not any(re.search(r'negative bilaterally|degenerative and pre-existing|maximum medical improvement|No further treatment|no work restrictions', x, re.I) and 'correct?' in x for x in primary), 'H1: primary locks in adverse conclusions'
    assert any('specific findings' in x for x in primary) and any('MRI' in x for x in primary), 'M8: primary omits document-based cross tools'
    danger = '\n'.join(section(doc, 'Danger zone'))
    for text in ('lumbar strain in 2019','negative bilaterally','mild degenerative change','denies any prior','40 pounds','60 pounds'):
        assert text in danger, 'M7: danger zone omits ' + text
    cross = '\n'.join(section(doc,'Questions: anticipated cross'))
    assert 'positive straight leg raise' not in cross.lower(), 'L5: cross opens a finding that cuts both ways'
    redirect = [x.partition('. ')[2] for x in section(doc,'Questions: redirect') if re.match(r'^\d+\. ', x)]
    assert not set(redirect) & {x.partition('. ')[2] for x in primary}, 'L5: redirect repeats primary'


def test_practice_ours(tmp_path, case):
    doc = built(tmp_path, case, 'depo-our-doctor')
    primary = section(doc,'Questions: primary')
    bills = next(x for x in primary if 'as to the bills themselves' in x)
    assert 'Marchetti' not in bills, 'H2: bills question names the defense examiner'
    i = next(i for i,x in enumerate(primary) if re.match(r'^\d+\. ', x) and 'treatment' in x)
    assert 'Plan: left L4-L5 transforaminal epidural steroid injection.' in primary[i+1], 'M6: treatment question rests on a causation opinion'
    assert any('not establish' in x.lower() for x in primary[i+1:i+3]), 'M6: treatment premise overstates record support'


def test_practice_report(tmp_path, case):
    defense = built(tmp_path, case, 'report-defense')
    labels = '\n'.join(section(defense,'Undefined labels'))
    for text in ('5/5 bilaterally','2+ and symmetric','Sensation is intact','Straight leg raise negative'):
        assert text in labels, 'H3: report omits documented supporting findings'
    ours = built(tmp_path, case, 'report-ours')
    elements = section(ours,'Elements')
    for name in ('Prior conditions and history', 'Permanency rating', 'Records reviewed'):
        assert any(x.startswith(name + ': NOT FOUND') for x in elements), 'L4: unsupported element labelled present: ' + name


def test_fixture_def0002_and_dates(tmp_path, case):
    doc = built_memo(tmp_path, case)
    lines = section(doc,'Conflicts')
    indexes = [i for i,x in enumerate(lines) if x == 'From [cite DEF0002]:']
    assert len(indexes) == 2 and all(lines[i+1].startswith('"') for i in indexes), 'H4 fixture: DEF0002 has no quote beneath its conflicts'
    chronology = section(doc,'Treatment chronology')
    assert any(x.startswith('03/11/2025,') and 'Date of hire' in x for x in chronology), 'L7 fixture: hire fact is not under the record date'
    assert any('5/5 bilaterally' in x and x.startswith('06/20/2025,') for x in chronology), 'H3 fixture: examination findings absent from chronology'


# These expected values are pinned independently of the runtime fixture loader.
DEF0002_QUOTES = ['Range of motion (ROM) in the lumbar spine is limited in flexion.', 'Strength testing of lower extremities is 5/5 bilaterally.', 'Reflexes are 2+ and symmetric in all tested segments.', 'Sensation is intact to light touch throughout the dermatomes.', 'Straight leg raise negative bilaterally.', "No objective findings of radiculopathy on today's examination.", 'The L4-L5 disc changes are described as degenerative.']
DATE_MOVES = [('Restrictions: no lifting over 10 pounds and no repetitive bending through 04/11/2025.', '03/14/2025'), ('Patient reports new numbness into the left calf since approximately 04/10/2025.', '04/14/2025'), ('In my opinion the L4-L5 disc protrusion is causally related to the lifting injury of 03/11/2025 described by the patient, within a reasonable degree of medical probability.', '05/09/2025'), ('Employee reported lifting a tote weighing approximately 40 pounds from a floor pallet.', '03/11/2025'), ('Average weekly wage: $912.40', '03/15/2025'), ('The L4-L5 findings are degenerative and pre-existing, not causally related to the work incident of 03/11/2025.', '06/20/2025'), ('Restrictions: sedentary duty, no lifting over 15 pounds, until re-evaluation on 06/20/2025.', '05/09/2025'), ('Date of hire: 08/02/2021', '03/11/2025'), ('Lift up to 50 pounds repeatedly during the shift.', '03/11/2025')]

def test_fixture_each_new_fact(tmp_path, case):
    doc = built_memo(tmp_path, case)
    chronology = section(doc, 'Treatment chronology')
    conflicts = section(doc, 'Conflicts')
    for quote in DEF0002_QUOTES:
        evidence = f'"{quote}" [cite DEF0002]'
        if quote == 'Straight leg raise negative bilaterally.':
            assert evidence in conflicts, 'Fixture: new DEF0002 quote missing from built memo: ' + quote
        else:
            assert any(x.startswith('06/20/2025, Dr. Theodore Marchetti, ') and x.endswith(evidence) for x in chronology), 'Fixture: new DEF0002 quote missing from built memo: ' + quote
    for quote, date in DATE_MOVES:
        if date:
            assert any(x.startswith(date + ', ') and f'"{quote}"' in x for x in chronology), 'Fixture: date move missing from built memo: ' + quote
        else:
            assert any(x.startswith(f'"{quote}" ') for x in section(doc, 'Facts without a date, by page')), 'Fixture: undated move missing from built memo: ' + quote
